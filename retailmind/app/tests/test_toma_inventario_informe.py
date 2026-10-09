"""
Toma completa de una tienda contada ANOCHE y cargada HOY (caso real PAO4 Matta
2458, 07/08-10-2026), de punta a punta hasta mover el stock, y el informe final
con el formato antiguo (por marca + diferencias por SKU).

Casos que trae la pistola real de PAO4:
- SKU vendido hoy en la mañana que anoche estaba (stock actual 0): debe estar en
  la toma con el stock AL CORTE y quedar en 0 después de aplicar.
- SKUs operativos (VISA/DIFER VISA, BOLSA CALZADOS/PAPEL) que nadie cuenta: se
  excluyen, nunca se llevan a 0.
- SKUs del sistema que no aparecieron en la pistola: faltante (se cuentan en 0);
  los de 2+ unidades piden reconteo.
- Códigos de la pistola que no existen en la sucursal (EAN del proveedor o SKU de
  otra tienda): quedan en «No cargados» con su cantidad.
- Stock negativo: el inventario completo lo deja en 0.

Correr (sin tocar la BD del .env):
    DATABASE_URL='sqlite://:memory:' python manage.py test app.tests.test_toma_inventario_informe
"""
import io
import json
from datetime import timedelta
from decimal import Decimal

import openpyxl
from django.test import SimpleTestCase
from django.urls import reverse
from django.utils import timezone

from app.models import (
    AtributoOpcion, Categoria, Movimientos_Producto, Producto, Producto_Talla, Productos_Atributos,
    TomaInventario,
)
from app.services import informe_toma_inventario as informe
from app.views_gestion_inventarios import _ejecutar_ajustes_background, _iniciar_tarea_ajustes
from .factories import crear_sucursal
from .test_toma_inventario import BaseTomaInventarioTest


class BaseTomaContadaAnoche(BaseTomaInventarioTest):
    """Marcas NIKE/PAOLA, corte anoche y venta de hoy; helpers de productos y POST."""

    def setUp(self):
        super().setUp()
        atributo_marca = Productos_Atributos.objects.create(nombre='Marca', descripcion='Marca')
        self.nike = AtributoOpcion.objects.create(atributo=atributo_marca, valor='NIKE')
        self.paola = AtributoOpcion.objects.create(atributo=atributo_marca, valor='PAOLA')
        self.ahora = timezone.localtime()
        self.corte = self.ahora - timedelta(hours=14)   # anoche, al terminar de contar
        self.venta_hoy = self.ahora - timedelta(hours=1)

    def _pt(self, articulo, sku, stock, marca=None, descripcion=None, costo=10000, sobreprecio=1500,
            precio=19990, sucursal=None):
        producto = Producto.objects.create(
            articulo=articulo, descripcion=descripcion or articulo,
            sucursal=sucursal or self.sucursal, costo=costo, sobreprecio=sobreprecio,
            precioventa=precio, categoria=self.categoria, atributo1=marca or self.nike,
        )
        return Producto_Talla.objects.create(producto=producto, sku=sku, stock=stock, talla='40')

    def _post(self, nombre_url, toma_id, payload=None):
        return self.client.post(
            reverse(nombre_url, args=[toma_id]),
            data=json.dumps(payload or {}), content_type='application/json',
        ).json()


class TomaCompletaContadaAnocheTest(BaseTomaContadaAnoche):

    def test_ciclo_completo_y_informe(self):
        pt = {
            'exacto': self._pt('EXACTO', 9100001, 5),
            'sobra': self._pt('SOBRA', 9100002, 2),
            'falta': self._pt('FALTA', 9100003, 4),
            'cero_contado': self._pt('STOCK CERO', 9100004, 0),
            'vendido_hoy': self._pt('VENDIDO HOY', 9100005, 0),   # anoche 1, hoy se vendió
            'no_aparecio': self._pt('NO APARECIO', 9100006, 1),
            'no_aparecio_4': self._pt('NO APARECIO 4', 9100007, 4),
            'negativo': self._pt('NEGATIVO', 9100008, -1),
            'visa': self._pt('VISA', 123456789, 9914, marca=self.paola, descripcion='DIFER VISA', costo=0, sobreprecio=0, precio=0),
            'bolsa': self._pt('BOLSA CALZADOS', 9100009, 1138, marca=self.paola, descripcion='PAPEL', costo=143, sobreprecio=0, precio=490),
        }
        Movimientos_Producto.objects.create(
            ProductoTalla=pt['vendido_hoy'], cantidad=-1, concepto='VENTA_PUBLICO',
            sucursal_origen=self.sucursal, responsable='POS',
            fecha=self.venta_hoy.date(), hora=self.venta_hoy.time(),
        )
        otra = crear_sucursal(self.empresa, alias='OTRA')
        self._pt('DE OTRA TIENDA', 9100010, 1, sucursal=otra)

        # 1) Crear la toma con el corte de anoche y «conté con la tienda cerrada»
        data = self.client.post(reverse('api_crear_inventario'), data=json.dumps({
            'nombre': 'Inventario completo', 'tipo_inventario': 'COMPLETO',
            'fecha_corte': self.corte.strftime('%Y-%m-%dT%H:%M'), 'conteo_tienda_cerrada': True,
            'filtros': {'solo_con_stock': True},
        }), content_type='application/json').json()
        self.assertTrue(data['success'], data.get('error'))
        toma = TomaInventario.objects.get(id=data['inventario_id'])
        skus_toma = set(toma.detalles.values_list('sku', flat=True))
        # Lo vendido hoy estaba anoche: entra a la toma con el stock AL CORTE
        self.assertIn('9100005', skus_toma)
        self.assertEqual(toma.detalles.get(sku='9100005').stock_sistema, 1)
        # El negativo también entra (el inventario completo lo corrige); el stock 0 no
        self.assertIn('9100008', skus_toma)
        self.assertNotIn('9100004', skus_toma)

        # 2) Pistola (filas repetidas se suman; dos códigos que no existen en la tienda)
        pistola = (
            'sku,stock\n9100001,5\n9100002,2\n9100002,1\n9100003,3\n9100004,1\n9100005,1\n'
            '7900204305426,1\n9100010,2\n'
        )
        resp = self._importar_pistola(toma.id, pistola)
        self.assertTrue(resp['success'], resp)
        self.assertEqual(resp['unidades_leidas'], 16)
        self.assertEqual(sorted(resp['no_encontrados']), ['7900204305426', '9100010'])
        self.assertEqual(resp['no_encontrados_unidades'], 3)
        diferencias = dict(toma.detalles.filter(contado=True).values_list('sku', 'diferencia'))
        self.assertEqual(diferencias, {
            '9100001': 0, '9100002': 1, '9100003': -1, '9100004': 1, '9100005': 0,
        })

        # 3) Lo no contado como faltante: operativos excluidos, nunca en 0
        previa = self._post('api_resolver_no_contados', toma.id, {'accion': 'faltante', 'previsualizar': True})
        self.assertTrue(previa['success'], previa)
        self.assertEqual({o['sku'] for o in previa['resultado']['operativos']}, {'123456789', '9100009'})
        self.assertEqual(previa['resultado']['operativos_unidades'], 9914 + 1138)
        self.assertEqual(previa['resultado']['lineas'], 3)  # no_aparecio, no_aparecio_4, negativo
        self.assertFalse(toma.detalles.get(sku='9100006').contado)  # previsualizar no escribe

        resp = self._post('api_resolver_no_contados', toma.id, {'accion': 'faltante'})
        self.assertTrue(resp['success'], resp)
        visa = toma.detalles.get(sku='123456789')
        self.assertTrue(visa.excluir_de_analisis)
        self.assertFalse(visa.contado)
        self.assertIn('SKU operativo', visa.observaciones)
        d4 = toma.detalles.get(sku='9100007')
        self.assertEqual((d4.stock_fisico, d4.diferencia, d4.reconteo_requerido), (0, -4, True))
        self.assertEqual(toma.detalles.get(sku='9100006').diferencia, -1)
        self.assertEqual(toma.detalles.get(sku='9100008').diferencia, 1)  # −1 → 0

        # 4) Finalizar: queda en revisión por el reconteo; se reconfirma 0
        resp = self._post('api_finalizar_conteo', toma.id)
        self.assertTrue(resp['success'], resp)
        self.assertEqual(resp['estado'], 'EN_REVISION')
        resp = self._post('api_registrar_reconteo', toma.id, {
            'reconteos': [{'detalle_id': d4.id, 'stock_reconteo': 0, 'observaciones': 'no está'}],
        })
        self.assertEqual(resp['reconteos_realizados'], 1)
        self.assertTrue(self._post('api_enviar_aprobacion', toma.id)['success'])
        self.assertTrue(self._post('api_aprobar_inventario', toma.id)['success'])

        # 5) Informe ANTES de aplicar (es la vista previa de lo que se ajustará)
        toma.refresh_from_db()
        analisis = informe.analizar(informe.filas_desde_toma(toma))
        total = analisis['total']
        # antiguo = 5+2+4+1(vendido hoy, al corte)+1+4+(−1) ; nuevo = 5+3+3+1+1
        self.assertEqual((total['ant_stock'], total['nue_stock'], total['dif']), (16, 13, -3))
        # P COSTO = unidades × (costo + sobreprecio) = 11.500 por unidad
        self.assertEqual(total['ant_costo'], 16 * 11500)
        self.assertEqual(total['nue_venta'], 13 * 19990)
        self.assertEqual([m['marca'] for m in analisis['marcas']], ['NIKE'])  # PAOLA (operativos) excluida
        self.assertEqual(analisis['resumen']['excluidos'], 2)
        self.assertEqual(
            {f['sku']: f['final'] for f in analisis['diferencias']},
            {'9100002': 1, '9100003': -1, '9100004': 1, '9100006': -1, '9100007': -4, '9100008': 1},
        )
        no_cargados = {n['sku']: n for n in informe.no_cargados_desde_logs(toma)}
        self.assertEqual(no_cargados['7900204305426']['cantidad'], 1)
        self.assertIn('EAN', no_cargados['7900204305426']['sugerencia'])
        self.assertEqual(no_cargados['9100010']['cantidad'], 2)
        self.assertEqual(no_cargados['9100010']['existe_en'], 'OTRA (1)')

        # 6) Aplicar (síncrono, como el command): stock final = contado − vendido después
        tarea, iniciada = _iniciar_tarea_ajustes(toma, self.user)
        self.assertTrue(iniciada)
        _ejecutar_ajustes_background(toma.id, self.user.id, cerrar_conexion=False)
        toma.refresh_from_db()
        self.assertEqual(toma.estado, 'COMPLETADO')
        esperado = {
            'exacto': 5, 'sobra': 3, 'falta': 3, 'cero_contado': 1, 'vendido_hoy': 0,
            'no_aparecio': 0, 'no_aparecio_4': 0, 'negativo': 0, 'visa': 9914, 'bolsa': 1138,
        }
        for clave, stock in esperado.items():
            pt[clave].refresh_from_db()
            self.assertEqual(pt[clave].stock, stock, clave)
        ajustes = Movimientos_Producto.objects.filter(referencia_externa=toma.numero_inventario)
        self.assertEqual(ajustes.filter(concepto='AJUSTE_INVENTARIO_ENTRADA').count(), 3)
        self.assertEqual(ajustes.filter(concepto='AJUSTE_INVENTARIO_SALIDA').count(), 3)
        self.assertFalse(ajustes.filter(ProductoTalla__in=[pt['visa'], pt['bolsa']]).exists())

        # 7) Endpoints del informe
        resp = self.client.get(reverse('api_informe_marcas_inventario', args=[toma.id])).json()
        self.assertTrue(resp['success'], resp)
        self.assertEqual(resp['total']['dif'], -3)
        self.assertEqual(len(resp['no_cargados']), 2)
        excel = self.client.get(reverse('api_informe_final_inventario', args=[toma.id]))
        self.assertEqual(excel.status_code, 200)
        self.assertIn('spreadsheetml', excel['Content-Type'])
        wb = openpyxl.load_workbook(io.BytesIO(excel.content))
        self.assertEqual(wb.sheetnames, ['Por marca', 'Diferencias', 'Mantenidos', 'No cargados', 'Excluidos', 'Resumen'])
        hoja = wb['Por marca']
        self.assertIn('Inventario General', hoja['A1'].value)
        fila_total = next(r for r in hoja.iter_rows(min_row=5, values_only=True) if r[0] == 'Total general')
        self.assertEqual((fila_total[1], fila_total[5], fila_total[8]), (16, -3, 13))
        encabezados = [c.value for c in wb['Diferencias'][1]]
        self.assertEqual(encabezados[:16], [
            'id', 'sku', 'art', 'talla', 'marca', 'costo', 'costo2', 'stk', 'pistola', 'mov',
            'final', 'pvp', 'ttcosto1', 'ttpvp1', 'ttcosto2', 'ttpvp2',
        ])

        # 8) La pantalla de detalle renderiza con la guía y el resultado por marca
        pagina = self.client.get(reverse('detalle_inventario', args=[toma.id]))
        self.assertEqual(pagina.status_code, 200)
        self.assertContains(pagina, 'id="guiaPaso"')
        self.assertContains(pagina, 'id="tbodyMarcas"')


class NoPistoleadoMantieneStockTest(BaseTomaContadaAnoche):
    """«Lo que no está en la pistola toma el antiguo»: decisiones por SKU y por
    marca en un solo plan; bolsas (operativos) y lo no pistoleado conservan stock."""

    def test_plan_mantener_por_sku_faltante_por_marca(self):
        pt = {
            'contado': self._pt('CONTADO', 9200001, 2),
            'no_pistoleado': self._pt('NO PISTOLEADO', 9200002, 3),
            'no_esta': self._pt('NO ESTA', 9200003, 1),
            'vendido_sin_pistolear': self._pt('VENDIDO SIN PISTOLEAR', 9200004, 0),  # anoche 1
            'bolsa': self._pt('45-1', 9200005, 500, marca=self.paola, descripcion='BOLSA CORPORATIVA', costo=370, sobreprecio=0, precio=990),
        }
        Movimientos_Producto.objects.create(
            ProductoTalla=pt['vendido_sin_pistolear'], cantidad=-1, concepto='VENTA_PUBLICO',
            sucursal_origen=self.sucursal, responsable='POS',
            fecha=self.venta_hoy.date(), hora=self.venta_hoy.time(),
        )
        data = self.client.post(reverse('api_crear_inventario'), data=json.dumps({
            'nombre': 'Completo', 'tipo_inventario': 'COMPLETO', 'conteo_tienda_cerrada': True,
            'fecha_corte': self.corte.strftime('%Y-%m-%dT%H:%M'), 'filtros': {'solo_con_stock': True},
        }), content_type='application/json').json()
        self.assertTrue(data['success'], data.get('error'))
        toma = TomaInventario.objects.get(id=data['inventario_id'])
        self.assertTrue(self._importar_pistola(toma.id, 'sku,stock\n9200001,2\n')['success'])

        # Agrupado para la pantalla: la bolsa va aparte (operativo), el resto por marca
        resp = self._post('api_resolver_no_contados', toma.id, {'agrupar': True})
        self.assertTrue(resp['success'], resp)
        self.assertEqual([o['sku'] for o in resp['operativos']], ['9200005'])
        # 3 + 1 + 1 (el vendido hoy tenía 1 al corte)
        self.assertEqual([(g['etiqueta'], g['lineas'], g['unidades']) for g in resp['grupos']], [('NIKE', 3, 5)])

        detalle = {d.sku: d for d in toma.detalles.all()}
        resp = self._post('api_resolver_no_contados', toma.id, {'plan': [
            {'accion': 'operativos'},
            {'accion': 'sin_diferencia', 'detalle_ids': [detalle['9200002'].id]},
            {'accion': 'faltante', 'marcas': ['NIKE']},
        ]})
        self.assertTrue(resp['success'], resp)
        self.assertEqual(resp['pendientes'], 0)
        detalle = {d.sku: d for d in toma.detalles.all()}
        self.assertTrue(detalle['9200005'].excluir_de_analisis)                    # bolsa: fuera, stock intacto
        self.assertEqual((detalle['9200002'].stock_fisico, detalle['9200002'].diferencia), (3, 0))  # toma el antiguo
        self.assertIn(informe.OBSERVACION_MANTENIDO, detalle['9200002'].observaciones)
        self.assertEqual(detalle['9200003'].diferencia, -1)                        # faltante
        # Vendido hoy sin pistolear: existía al contar → no es faltante (no deja negativo)
        self.assertEqual((detalle['9200004'].stock_fisico, detalle['9200004'].diferencia), (1, 0))

        for url in ('api_finalizar_conteo', 'api_enviar_aprobacion', 'api_aprobar_inventario'):
            resp = self._post(url, toma.id)
            self.assertTrue(resp['success'], (url, resp))
        _iniciar_tarea_ajustes(toma, self.user)
        _ejecutar_ajustes_background(toma.id, self.user.id, cerrar_conexion=False)
        toma.refresh_from_db()
        self.assertEqual(toma.estado, 'COMPLETADO')
        for clave, stock in {'contado': 2, 'no_pistoleado': 3, 'no_esta': 0,
                             'vendido_sin_pistolear': 0, 'bolsa': 500}.items():
            pt[clave].refresh_from_db()
            self.assertEqual(pt[clave].stock, stock, clave)

        analisis = informe.analizar(informe.filas_desde_toma(toma))
        self.assertEqual((analisis['total']['ant_stock'], analisis['total']['nue_stock']), (7, 6))
        self.assertEqual((analisis['resumen']['mantenidos'], analisis['resumen']['mantenidos_unidades']), (1, 3))
        self.assertEqual([f['sku'] for f in analisis['mantenidas']], ['9200002'])
        self.assertEqual([f['sku'] for f in analisis['diferencias']], ['9200003'])


class SucursalYRevisionTest(BaseTomaContadaAnoche):
    """Elegir la sucursal al crear, aviso de toma abierta, buscar/ordenar y «No ajustar» en bloque."""

    def _crear(self, **extra):
        payload = {'nombre': 'Completo', 'tipo_inventario': 'COMPLETO', 'filtros': {'solo_con_stock': True},
                   'fecha_corte': self.corte.strftime('%Y-%m-%dT%H:%M'), 'conteo_tienda_cerrada': True}
        payload.update(extra)
        return self.client.post(reverse('api_crear_inventario'), data=json.dumps(payload),
                                content_type='application/json').json()

    def test_crear_en_la_sucursal_elegida_y_aviso_de_toma_abierta(self):
        otra = crear_sucursal(self.empresa, alias='PAO4-TEST', direccion='Matta 2458')
        self._pt('ZAPATILLA', 9300001, 2, sucursal=otra)
        data = self._crear(sucursal_id=otra.id)  # la sesión está en self.sucursal
        self.assertTrue(data['success'], data)
        self.assertEqual(data['sucursal'], 'PAO4-TEST')
        self.assertEqual(TomaInventario.objects.get(id=data['inventario_id']).sucursal_id, otra.id)

        # Otra toma en la misma tienda: pide confirmación (aplicar las dos duplicaría el ajuste)
        data = self._crear(sucursal_id=otra.id)
        self.assertFalse(data['success'])
        self.assertTrue(data['requiere_confirmacion'])
        self.assertEqual(len(data['abiertas']), 1)
        self.assertTrue(self._crear(sucursal_id=otra.id, permitir_otra_abierta=True)['success'])

        # Sucursal inexistente / sin acceso
        self.assertIn('acceso', self._crear(sucursal_id=999999)['error'])

        # El listado filtra por la sucursal elegida
        resp = self.client.get(reverse('api_obtener_inventarios'), {'sucursal': otra.id}).json()
        self.assertEqual(resp['pagination']['total_items'], 2)
        self.assertEqual({i['sucursal'] for i in resp['inventarios']}, {'PAO4-TEST'})

    def test_buscar_por_descripcion_ordenar_y_no_ajustar_en_bloque(self):
        bolsa = self._pt('45-1', 9300010, 50, marca=self.paola, descripcion='BOLSA CORPORATIVA')
        sobra = self._pt('SOBRA MUCHO', 9300011, 1)
        falta = self._pt('FALTA POCO', 9300012, 5)
        self._pt('EXACTO', 9300013, 3)
        toma = TomaInventario.objects.get(id=self._crear()['inventario_id'])
        self.assertTrue(self._importar_pistola(toma.id, 'sku,stock\n9300011,8\n9300012,4\n9300013,3\n')['success'])
        url = reverse('api_productos_conteo', args=[toma.id])

        # «BOLSA» encuentra el artículo 45-1 por su descripción, y la trae para mostrarla
        resp = self.client.get(url, {'search': 'BOLSA'}).json()
        self.assertEqual([p['sku'] for p in resp['productos']], ['9300010'])
        self.assertEqual(resp['productos'][0]['descripcion'], 'BOLSA CORPORATIVA')

        # Mayor diferencia primero: el pistoleado 8 veces (+7) encabeza
        resp = self.client.get(url, {'orden': 'dif_unidades', 'estado_conteo': 'contado'}).json()
        self.assertEqual([p['sku'] for p in resp['productos']][:2], ['9300011', '9300012'])

        # «No ajustar» en bloque: excluidas, no mueven stock al aplicar
        ids = list(toma.detalles.filter(sku__in=['9300011', '9300012']).values_list('id', flat=True))
        resp = self._post('api_excluir_detalles_inventario', toma.id, {'ids': ids, 'excluir': True})
        self.assertEqual(resp['actualizados'], 2)
        resp = self.client.get(url, {'estado_conteo': 'excluido'}).json()
        self.assertEqual({p['sku'] for p in resp['productos']}, {'9300011', '9300012'})

        resp = self._post('api_resolver_no_contados', toma.id, {'plan': [{'accion': 'operativos'}, {'accion': 'sin_diferencia'}]})
        self.assertTrue(resp['success'], resp)
        for url_paso in ('api_finalizar_conteo', 'api_enviar_aprobacion', 'api_aprobar_inventario'):
            self.assertTrue(self._post(url_paso, toma.id)['success'], url_paso)
        _iniciar_tarea_ajustes(toma, self.user)
        _ejecutar_ajustes_background(toma.id, self.user.id, cerrar_conexion=False)
        for pt, stock in ((sobra, 1), (falta, 5), (bolsa, 50)):
            pt.refresh_from_db()
            self.assertEqual(pt.stock, stock, pt.sku)
        self.assertFalse(Movimientos_Producto.objects.filter(referencia_externa=toma.numero_inventario).exists())


class CorreccionManualTiendaCerradaTest(BaseTomaContadaAnoche):
    """Corregir a mano en la tabla un conteo de anoche (código mal leído) en una
    toma de tienda cerrada: vale al corte, las ventas de hoy no lo inflan."""

    def test_conteo_manual_vale_al_corte(self):
        pt = self._pt('HT3900 GORRO', 9400001, 2)  # anoche 3, hoy se vendió 1
        Movimientos_Producto.objects.create(
            ProductoTalla=pt, cantidad=-1, concepto='VENTA_PUBLICO', sucursal_origen=self.sucursal,
            responsable='POS', fecha=self.venta_hoy.date(), hora=self.venta_hoy.time(),
        )
        data = self.client.post(reverse('api_crear_inventario'), data=json.dumps({
            'nombre': 'Completo', 'tipo_inventario': 'COMPLETO', 'conteo_tienda_cerrada': True,
            'fecha_corte': self.corte.strftime('%Y-%m-%dT%H:%M'), 'filtros': {'solo_con_stock': True},
        }), content_type='application/json').json()
        toma = TomaInventario.objects.get(id=data['inventario_id'])
        d = toma.detalles.get(sku='9400001')
        self.assertEqual(d.stock_sistema, 3)
        # La pistola lo leyó mal anoche; se corrige a mano con lo que había: 3
        resp = self._post('api_registrar_conteo', toma.id, {'conteos': [{'detalle_id': d.id, 'stock_fisico': 3}]})
        self.assertTrue(resp['success'], resp)
        d.refresh_from_db()
        self.assertEqual((d.stock_movimientos_post_corte, d.diferencia), (0, 0))  # antes: −1 de base → +1 falso


class TituloPorCategoriaTest(BaseTomaContadaAnoche):
    """El título del informe dice el tipo de toma y qué cubre (pedido del usuario 09-10, toma 12 de NICK2:
    era «por categoría» de calzado y el Excel decía «Inventario General»)."""

    def setUp(self):
        super().setUp()
        self.calzado = Categoria.objects.create(nombre='Calzado')
        self.zapatillas = Categoria.objects.create(nombre='Zapatillas', padre=self.calzado)
        self.botines = Categoria.objects.create(nombre='Botines', padre=self.calzado)
        self.sucursal.direccion = 'Matta 2438'
        self.sucursal.save(update_fields=['direccion'])
        pt = self._pt('ZAP', 9900001, 2)
        Producto.objects.filter(id=pt.producto_id).update(categoria=self.zapatillas)
        pt = self._pt('BOT', 9900002, 1)
        Producto.objects.filter(id=pt.producto_id).update(categoria=self.botines)

    def _toma(self, categorias):
        data = self.client.post(reverse('api_crear_inventario'), data=json.dumps({
            'nombre': 'Calzado', 'tipo_inventario': 'POR_CATEGORIA', 'conteo_tienda_cerrada': True,
            'fecha_corte': self.corte.strftime('%Y-%m-%dT%H:%M'),
            'filtros': {'solo_con_stock': True, 'categorias': [str(c.id) for c in categorias]},
        }), content_type='application/json').json()
        self.assertTrue(data['success'], data)
        return TomaInventario.objects.get(id=data['inventario_id'])

    def test_todas_las_hijas_de_calzado_dicen_calzado(self):
        toma = self._toma([self.zapatillas, self.botines])
        self.assertTrue(self._importar_pistola(toma.id, 'sku,stock\n9900001,2\n9900002,1\n')['success'])
        anio = timezone.localtime(toma.fecha_corte).year
        titulo = f'{anio} 2438  Inventario por Categoría · Calzado'
        self.assertEqual(self.client.get(reverse('api_informe_marcas_inventario', args=[toma.id])).json()['titulo'], titulo)

        excel = self.client.get(reverse('api_informe_final_inventario', args=[toma.id]))
        self.assertEqual(excel.status_code, 200)
        self.assertIn(f'informe_{toma.numero_inventario}_{self.sucursal.alias}_Calzado.xlsx', excel['Content-Disposition'])
        wb = openpyxl.load_workbook(io.BytesIO(excel.content))
        self.assertEqual(wb['Por marca']['A1'].value, titulo)
        resumen = {fila[0]: fila[1] for fila in wb['Resumen'].iter_rows(values_only=True)}
        self.assertEqual(resumen['Tipo'], 'Por Categoría/Departamento')
        self.assertEqual(resumen['Qué cubre'], 'Calzado: Botines, Zapatillas')

    def test_parte_de_calzado_nombra_las_categorias(self):
        toma = self._toma([self.zapatillas])
        cabecera = informe.cabecera_desde_toma(toma)
        self.assertEqual((cabecera['alcance'], cabecera['alcance_detalle']), ('Calzado: Zapatillas', 'Calzado: Zapatillas'))
        self.assertTrue(informe.titulo_informe(cabecera).endswith('Inventario por Categoría · Calzado: Zapatillas'))


class InformePuroTest(SimpleTestCase):
    """El cálculo sin BD (lo usa también el script que simula una toma)."""

    def test_titulo_y_archivo_dicen_tipo_y_alcance(self):
        cabecera = {'anio': 2026, 'direccion': 'Matta 2438', 'alias': 'NICK2', 'numero': 'INV-7-20261009-001',
                    'tipo_codigo': 'POR_CATEGORIA', 'alcance': 'Calzado (10 de 13 categorías)'}
        self.assertEqual(informe.titulo_informe(cabecera), '2026 2438  Inventario por Categoría · Calzado (10 de 13 categorías)')
        self.assertEqual(informe.nombre_archivo_informe(cabecera), 'informe_INV-7-20261009-001_NICK2_Calzado_10_de_13_categorias.xlsx')
        completo = {**cabecera, 'tipo_codigo': 'COMPLETO', 'alcance': ''}
        self.assertEqual(informe.titulo_informe(completo), '2026 2438  Inventario General')
        self.assertEqual(informe.nombre_archivo_informe(completo), 'informe_INV-7-20261009-001_NICK2.xlsx')

    def test_lectura_repetida(self):
        # Un código pistoleado 8 veces con 1 en el sistema
        self.assertIn('lectura repetida', informe.alerta_diferencia(1, 8, 7))
        self.assertIn('doble lectura', informe.alerta_diferencia(3, 6, 3))
        self.assertEqual(informe.alerta_diferencia(4, 5, 1), '')

    def test_sin_contar_cuenta_cero_y_excluidas_no_suman(self):
        filas = [
            {'id': 1, 'sku': '1', 'marca': 'adidas', 'costo': 100, 'sobreprecio': 10, 'pvp': 300,
             'stk': 3, 'mov': 0, 'pistola': 2},
            {'id': 2, 'sku': '2', 'marca': 'ADIDAS ORIGINALS', 'costo': 50, 'sobreprecio': 0, 'pvp': 90,
             'stk': 1, 'mov': 0, 'pistola': None},
            {'id': 3, 'sku': '3', 'marca': 'champion', 'costo': 10, 'sobreprecio': 0, 'pvp': 20,
             'stk': 0, 'mov': -1, 'pistola': 0},
            {'id': 4, 'sku': '4', 'marca': 'PAOLA', 'costo': 0, 'sobreprecio': 0, 'pvp': 0,
             'stk': 9914, 'mov': 0, 'pistola': None, 'excluida': True},
            # PAO4: curva CHALADA 23-MITSU-2 leída dos veces (3 → 6)
            {'id': 5, 'sku': '5', 'marca': 'CHALADA', 'costo': 9000, 'sobreprecio': 0, 'pvp': 19990,
             'stk': 3, 'mov': 0, 'pistola': 6},
        ]
        a = informe.analizar(filas)
        # orden alfabético sin distinguir mayúsculas, como el informe antiguo
        self.assertEqual([m['marca'] for m in a['marcas']], ['adidas', 'ADIDAS ORIGINALS', 'CHALADA', 'champion'])
        adidas = a['marcas'][0]
        self.assertEqual((adidas['ant_stock'], adidas['nue_stock'], adidas['dif']), (3, 2, -1))
        self.assertEqual((adidas['ant_costo'], adidas['nue_costo']), (330.0, 220.0))
        self.assertEqual(a['total']['ant_stock'], 3 + 1 - 1 + 3)
        self.assertEqual(a['resumen']['sin_contar'], 1)
        self.assertEqual(a['resumen']['excluidos_unidades'], 9914)
        por_sku = {f['sku']: f for f in a['diferencias']}
        # stk 0 con venta posterior (−1) y pistola 0: el ajuste deja el negativo en 0
        self.assertEqual(por_sku['3']['final'], 1)
        self.assertIn('negativo', por_sku['3']['alerta'])
        self.assertIn('doble lectura', por_sku['5']['alerta'])
        self.assertEqual((a['resumen']['posibles_dobles'], a['resumen']['posibles_dobles_unidades']), (1, 3))
        self.assertEqual(por_sku['2']['estado'], 'No apareció')
        self.assertEqual(informe.titulo_informe({'anio': 2026, 'direccion': 'Matta 2458'}), '2026 2458  Inventario General')
