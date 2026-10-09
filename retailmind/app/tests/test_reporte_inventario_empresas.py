"""
«Inventario por empresa» (pedido del usuario, 09-10): el cuadro «2026 ENERO INV» de
gerencia, solo para el Maestro. Por local: la última toma del mes (lo contado, como
su informe final) o, sin toma, el stock del sistema; totales por empresa y holding.

Correr (sin tocar la BD del .env):
    DATABASE_URL='sqlite://:memory:' python manage.py test app.tests.test_reporte_inventario_empresas
"""
import io
import json

import openpyxl
from django.urls import reverse
from django.utils import timezone

from app.models import TomaInventario
from .factories import crear_empresa, crear_empresa_user, crear_sucursal, crear_usuario
from .test_toma_inventario_informe import BaseTomaContadaAnoche


class ReporteInventarioEmpresasTest(BaseTomaContadaAnoche):

    def setUp(self):
        super().setUp()
        # Empresa A: tienda (con toma) + bodega (sin toma). Empresa B: tienda sin toma.
        self.sucursal.alias, self.sucursal.direccion = 'PAO2', 'Matta 2422'
        self.sucursal.save()
        self.bodega = crear_sucursal(self.empresa, alias='PA00', direccion='Maipu 676',
                                     es_centro_distribucion=True, tipo_sucursal='CENTRO_DISTRIBUCION')
        self.empresa_b = crear_empresa('Importadora Test', rut='76.111.111-1')
        self.tienda_b = crear_sucursal(self.empresa_b, alias='NICK1', direccion='Matta 2479')

        # costo 10.000 + sobreprecio 1.500 = P. Interno 11.500; venta 19.990
        self._pt('UNO', 9900001, 5)
        self._pt('DOS', 9900002, 2)
        self._pt('BOLSA', 9900003, 3, costo=1000, sobreprecio=0, precio=2000, sucursal=self.bodega)
        self._pt('TRES', 9900004, 4, sucursal=self.tienda_b)

        data = self.client.post(reverse('api_crear_inventario'), data=json.dumps({
            'nombre': 'Completo', 'tipo_inventario': 'COMPLETO', 'conteo_tienda_cerrada': True,
            'fecha_corte': self.corte.strftime('%Y-%m-%dT%H:%M'), 'filtros': {'solo_con_stock': True},
        }), content_type='application/json').json()
        self.assertTrue(data['success'], data)
        self.toma = TomaInventario.objects.get(id=data['inventario_id'])
        # UNO 5 → 4 (−1), DOS 2 → 2
        self.assertTrue(self._importar_pistola(self.toma.id, 'sku,stock\n9900001,4\n9900002,2\n')['success'])

        self.maestro = crear_usuario(username='maestro_rep', rol='maestro')
        crear_empresa_user(self.maestro, self.empresa, self.sucursal)
        corte = timezone.localtime(self.corte)
        self.periodo = {'anio': corte.year, 'mes': corte.month}

    def _como_maestro(self):
        self.client.force_login(self.maestro)
        session = self.client.session
        session['idSucursalActual'] = self.sucursal.id
        session.save()

    def _cuadro(self, periodo=None):
        resp = self.client.get(reverse('api_reporte_inventario_empresas'), periodo or self.periodo).json()
        self.assertTrue(resp['success'], resp)
        return resp

    @staticmethod
    def _local(cuadro, alias):
        return next(l for b in cuadro['empresas'] for l in b['locales'] if l['alias'] == alias)

    def test_cuadro_por_empresa_con_toma_y_sin_toma(self):
        self._como_maestro()
        cuadro = self._cuadro()
        self.assertTrue(cuadro['titulo'].startswith(f"{self.periodo['anio']} "))
        self.assertTrue(cuadro['titulo'].endswith(' INV'))
        self.assertEqual([b['nombre'] for b in cuadro['empresas']], ['EMPRESA TEST', 'IMPORTADORA TEST'])

        tienda = self._local(cuadro, 'PAO2')
        self.assertEqual(tienda['local'], 'MATTA 2422')
        self.assertEqual(tienda['toma']['id'], self.toma.id)
        # Lo contado: 4 + 2 = 6 pares (el sistema tenía 7)
        self.assertEqual((tienda['pares'], tienda['pares_sistema'], tienda['diferencia']), (6, 7, -1))
        self.assertEqual((tienda['costo'], tienda['p_interno'], tienda['p_venta']), (60000, 69000, 6 * 19990))

        bodega = self._local(cuadro, 'PA00')
        self.assertEqual(bodega['local'], 'BODEGA PA00')
        self.assertIsNone(bodega['toma'])
        self.assertIsNone(bodega['diferencia'])
        self.assertEqual((bodega['pares'], bodega['costo'], bodega['p_interno'], bodega['p_venta']), (3, 3000, 3000, 6000))

        empresa_a = cuadro['empresas'][0]
        # Tiendas primero, bodegas al final
        self.assertEqual([l['alias'] for l in empresa_a['locales']], ['PAO2', 'PA00'])
        self.assertEqual((empresa_a['total']['pares'], empresa_a['total']['diferencia'],
                          empresa_a['total']['con_toma'], empresa_a['total']['locales']), (9, -1, 1, 2))
        self.assertEqual(self._local(cuadro, 'NICK1')['pares'], 4)
        self.assertEqual((cuadro['total']['pares'], cuadro['total']['p_interno'], cuadro['total']['diferencia']),
                         (13, 69000 + 3000 + 4 * 11500, -1))

    def test_toma_cancelada_no_entra_y_otro_mes_es_stock_del_sistema(self):
        self._como_maestro()
        TomaInventario.objects.filter(pk=self.toma.pk).update(estado='CANCELADO')
        tienda = self._local(self._cuadro(), 'PAO2')
        self.assertIsNone(tienda['toma'])
        self.assertEqual(tienda['pares'], 7)

        TomaInventario.objects.filter(pk=self.toma.pk).update(estado='EN_CONTEO')
        anio, mes = self.periodo['anio'], self.periodo['mes']
        anterior = {'anio': anio - (mes == 1), 'mes': 12 if mes == 1 else mes - 1}
        cuadro = self._cuadro(anterior)
        self.assertEqual(cuadro['total']['con_toma'], 0)
        self.assertEqual(self._local(cuadro, 'PAO2')['pares'], 7)  # sin movimientos: = stock de hoy

    def test_excel_con_el_mismo_cuadro(self):
        self._como_maestro()
        resp = self.client.get(reverse('api_exportar_reporte_inventario_empresas'), self.periodo)
        self.assertEqual(resp.status_code, 200)
        self.assertIn('inventario_por_empresa_', resp['Content-Disposition'])
        ws = openpyxl.load_workbook(io.BytesIO(resp.content)).active
        self.assertTrue(ws['A1'].value.endswith(' INV'))
        filas = {ws.cell(row=r, column=1).value: r for r in range(1, ws.max_row + 1)}
        self.assertEqual(ws.cell(row=filas['MATTA 2422'], column=3).value, 6)
        self.assertEqual(ws.cell(row=filas['MATTA 2422'], column=7).value, -1)
        self.assertEqual(ws.cell(row=filas['Total Holding'], column=3).value, 13)

    def test_solo_el_maestro(self):
        # self.user es administrador: ve los costos de las tomas pero no este cuadro
        self.assertNotIn('reporte-empresas', self.client.get(reverse('gestion_inventarios')).content.decode())
        self.assertEqual(self.client.get(reverse('reporte_inventario_empresas')).status_code, 403)
        self.assertEqual(self.client.get(reverse('api_reporte_inventario_empresas')).status_code, 403)
        self.assertEqual(self.client.get(reverse('api_exportar_reporte_inventario_empresas')).status_code, 403)

        self._como_maestro()
        self.assertIn('reporte-empresas', self.client.get(reverse('gestion_inventarios')).content.decode())
        self.assertEqual(self.client.get(reverse('reporte_inventario_empresas'), self.periodo).status_code, 200)

    def test_toma_parcial_suma_el_stock_de_lo_que_quedo_fuera(self):
        self._como_maestro()
        self._pt('FUERA', 9900005, 3)  # no está en la toma
        # Completa: es la foto del local al corte; lo que no tiene llegó después y no se suma
        self.assertEqual(self._local(self._cuadro(), 'PAO2')['pares'], 6)

        TomaInventario.objects.filter(pk=self.toma.pk).update(tipo_inventario='POR_CATEGORIA')
        tienda = self._local(self._cuadro(), 'PAO2')
        self.assertEqual((tienda['contado']['pares'], tienda['fuera_de_la_toma']['pares'], tienda['pares']), (6, 3, 9))
        self.assertEqual(tienda['p_interno'], 69000 + 3 * 11500)
        self.assertEqual(tienda['diferencia'], -1)  # la diferencia es solo de lo contado
