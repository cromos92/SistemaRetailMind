"""
Avisos de alcance en los Excel de Existencias por Marca y Resumen de
Existencias (auditoría reportes 2026-08, §5 P3): el archivo debe llevar los
mismos avisos que la pantalla —lista cortada por el límite de artículos,
exclusiones del modal y stock marcado `excluir_de_analitica`— y solo cuando
aplican, sin tocar las filas de datos.

Correr:
    python manage.py test app.tests.test_excel_notas_existencias --settings=test_settings_sqlite
"""
from io import BytesIO

import openpyxl
from django.test import TestCase
from django.urls import reverse

from app.models import AtributoOpcion, Productos_Atributos
from .factories import (
    crear_empresa, crear_empresa_user, crear_producto_con_talla, crear_sucursal,
    crear_usuario, otorgar_ver_pantalla,
)


def _columna_a(resp):
    """Valores de la columna A de la primera hoja del Excel devuelto."""
    ws = openpyxl.load_workbook(BytesIO(resp.content)).worksheets[0]
    return [ws.cell(row=r, column=1).value for r in range(1, ws.max_row + 1)]


class BaseNotasExcelTest(TestCase):
    def setUp(self):
        self.empresa = crear_empresa()
        self.tienda = crear_sucursal(self.empresa, alias='TIENDA')
        self.bodega = crear_sucursal(self.empresa, alias='IMP')
        self.user = crear_usuario(rol='administrador')
        crear_empresa_user(self.user, self.empresa, self.tienda)
        otorgar_ver_pantalla('administrador', 'reporte_existencias_marca',
                             'resumen_existencias', puede_exportar=True)

        attr = Productos_Atributos.objects.create(nombre='Marca', descripcion='m')
        self.marca = AtributoOpcion.objects.create(atributo=attr, valor='NIKE')

        self.client.force_login(self.user)
        session = self.client.session
        session['idSucursalActual'] = self.tienda.id
        session.save()
        self._sku = 7100000

    def _producto(self, sucursal, articulo, stock, **kwargs):
        self._sku += 1
        producto, _ = crear_producto_con_talla(
            sucursal, articulo=articulo, sku=self._sku, stock=stock,
            costo=15000, precioventa=20000, atributo1=self.marca, **kwargs)
        return producto


class ExcelExistenciasMarcaTest(BaseNotasExcelTest):
    def setUp(self):
        super().setUp()
        for articulo in ('AIR-1', 'AIR-2', 'AIR-3'):
            self._producto(self.tienda, articulo, stock=4)

    def _exportar(self, **params):
        return self.client.get(reverse('exportar_existencias_marca_excel'),
                               {'marca_id': self.marca.id, **params})

    def test_lista_truncada_lleva_el_aviso_con_las_cifras_de_la_api(self):
        api = self.client.get(reverse('obtener_reporte_existencias_marca'),
                              {'marca_id': self.marca.id, 'limite': 2}).json()
        self.assertTrue(api['truncado'])

        col_a = _columna_a(self._exportar(limite=2))
        self.assertTrue(col_a[0].startswith('LISTA INCOMPLETA: se incluyen 2 de 3 artículos'),
                        col_a[0])
        self.assertTrue(col_a[1].startswith('Alcance: totales calculados sobre 2 artículos incluidos'))
        self.assertIsNone(col_a[2])
        self.assertEqual(col_a[3], 'MARCA: NIKE')

    def test_sin_truncado_solo_la_nota_neutra_de_alcance(self):
        col_a = _columna_a(self._exportar())
        self.assertEqual(
            col_a[0],
            'Alcance: totales calculados sobre 3 artículos · sucursal TIENDA. '
            'Solo se incluyen artículos con stock.')
        self.assertFalse(any(str(v).startswith('LISTA INCOMPLETA') for v in col_a if v))
        self.assertEqual(col_a[2], 'MARCA: NIKE')
        # Las filas de datos siguen iguales (solo desplazadas por la nota)
        self.assertEqual(col_a[5:8], ['AIR-1', 'AIR-2', 'AIR-3'])


class ExcelResumenExistenciasTest(BaseNotasExcelTest):
    def _exportar(self, **params):
        return self.client.get(reverse('exportar_resumen_existencias_excel'), params)

    def test_sin_exclusiones_la_hoja_queda_como_antes(self):
        self._producto(self.tienda, 'ZAPATILLA', stock=5)
        col_a = _columna_a(self._exportar())
        self.assertTrue(col_a[0].startswith('RESUMEN DE EXISTENCIAS'))
        self.assertEqual(col_a[1], 'Sucursal')  # encabezado en la fila 2, sin avisos
        self.assertEqual(col_a[2], 'TIENDA')
        self.assertFalse(any('EXCLU' in str(v) for v in col_a if v))

    def test_exclusiones_y_stock_oculto_se_avisan_arriba(self):
        self._producto(self.tienda, 'ZAPATILLA', stock=5)
        polera = self._producto(self.tienda, 'POLERA', stock=30)
        self._producto(self.bodega, 'BOLSA REAL', stock=2200, excluir_de_analitica=True)

        col_a = _columna_a(self._exportar(excluir_articulos=str(polera.id)))
        self.assertEqual(
            col_a[1],
            'EXCLUSIONES: 1 artículo(s) excluidos del análisis — 30 unidad(es) por '
            '$600.000 a precio venta ($450.000 costo) (filtro temporal de la sesión). '
            'No suman en los totales.')
        self.assertTrue(col_a[2].startswith('EXCLUIDAS DE ANALÍTICA: 2.200 u'))
        self.assertTrue(col_a[2].endswith('Por sucursal: IMP 2.200 u (sin stock analítico).'))
        # Encabezado y datos intactos, corridos 2 filas; la nota vieja del pie ya no está
        self.assertEqual(col_a[3], 'Sucursal')
        self.assertEqual(col_a[4:6], ['IMP', 'TIENDA'])
        self.assertFalse(any(str(v).startswith('Nota:') for v in col_a if v))
