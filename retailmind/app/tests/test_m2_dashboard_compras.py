"""
Dashboard de Compras mejorado (views_modulo_compras.dashboard_compras_mejorado_api
y exportar_dashboard_compras): bordes de período, comparaciones y alertas.

Fija los hallazgos de la auditoría del 26-09-2026:
- A4-01: la base "año anterior" usa los MISMOS filtros (proveedor/temporada)
  que el período actual, así que la tarjeta cuadra con la comparativa.
- A4-02: trend_roi compara markup neto contra neto; sin base -> None.
- A4-03: el año en curso (y un personalizado con fin futuro) se compara hasta
  el mismo día del año anterior.
- A4-04: un rango de más de un año no se compara consigo mismo.
- A4-05: fechas del rango personalizado acotadas (antes 500 / decenas de MB).
- A4-06: "Proveedores Críticos" cuenta sobre todos, no sobre el top-12.
- A4-07: un período sin unidades pedidas no dispara "Cumplimiento Bajo".
- A4-08: "Últimos N días" son N días calendario incluido hoy.
- A4-09: proveedor no numérico se ignora; el 500 no filtra el traceback.
- B15-06: desglose OC reales vs ingresos sin OC ("Compra Manual").
- M2-R1: el Excel no deja como fórmula un nombre '=...' de compra, proveedor,
  temporada o producto.
"""
import io
import json
from datetime import date
from unittest import mock

from django.test import TestCase
from django.utils import timezone

from app.models import (
    Compras, Compras_Producto, Compras_Producto_Talla, Productos_Recepcionados,
)

from .factories import crear_empresa, crear_usuario, otorgar_ver_pantalla

API = '/app/dashboard_compras_mejorado_api/'
EXPORT = '/app/exportar_dashboard_compras/'
HOY = date(2026, 9, 26)
_localdate_original = timezone.localdate


def _localdate_fijo(value=None, timezone=None):
    """timezone.localdate() sin argumentos = HOY; con argumentos, el original."""
    if value is None:
        return HOY
    return _localdate_original(value, timezone)


class DashboardComprasM2Tests(TestCase):
    _correlativo = 0

    @classmethod
    def _compra(cls, proveedor, fecha, unidades, costo, precio, recibidas=0,
                nombre=None, estado='ACTIVA', familia=None, temporada=''):
        cls._correlativo += 1
        compra = Compras.objects.create(
            empresa=proveedor, nombre=nombre or f'OC {cls._correlativo}',
            correlativo=cls._correlativo, responsable='test', temporada=temporada,
            fecha=fecha, estado=estado, temporada_familia=familia,
            temporada_anio=fecha.year if familia else None,
        )
        cp = Compras_Producto.objects.create(
            compras=compra, nombre='Zapatilla', atributo1='MARCA', atributo2='',
            atributo3='', atributo4='', costo=costo, precioSugerido=precio,
        )
        cpt = Compras_Producto_Talla.objects.create(compra_producto=cp, stock=unidades, talla='42')
        if recibidas:
            Productos_Recepcionados.objects.create(compra_producto_talla=cpt, stockArribado=recibidas)
        return compra

    @classmethod
    def setUpTestData(cls):
        # 'maestro' pasa el middleware de permisos; además se otorga la pantalla
        # por si el mapeo de permisos se endurece.
        cls.user = crear_usuario(username='m2_maestro', rol='maestro')
        otorgar_ver_pantalla('maestro', 'dashboard_compras_estrategico', puede_exportar=True)
        cls.prov_a = crear_empresa(nombre='Proveedor A', rut='76.100.100-1', esProveedor=True)
        cls.prov_b = crear_empresa(nombre='Proveedor B', rut='76.200.200-2', esProveedor=True)

        # 2025 (año anterior). Markup de lista neto: 238.000 / 1,19 = 200.000 → 100 %.
        cls._compra(cls.prov_a, date(2025, 3, 10), 100, 1000, 2380, recibidas=50, familia='INVIERNO')
        cls._compra(cls.prov_a, date(2025, 11, 15), 100, 1000, 2380)       # después del 26-09
        cls._compra(cls.prov_b, date(2025, 4, 1), 1000, 1000, 2380)         # "toda la red"
        cls._compra(cls.prov_b, date(2025, 5, 1), 50, 1000, 2380, estado='ELIMINADA')

        # 2026 (año en curso): una OC real y un ingreso manual (se registra ya recibido).
        cls._compra(cls.prov_a, date(2026, 2, 1), 150, 1000, 2975, recibidas=75, familia='INVIERNO')
        cls._compra(cls.prov_a, date(2026, 5, 1), 10, 1000, 2380, recibidas=10,
                    nombre='Compra Manual - Proveedor A - 01/05/2026')

    def setUp(self):
        self.client.force_login(self.user)
        p = mock.patch('django.utils.timezone.localdate', side_effect=_localdate_fijo)
        p.start()
        self.addCleanup(p.stop)

    def _get(self, **params):
        r = self.client.get(API, params)
        self.assertEqual(r.status_code, 200, r.content[:300])
        return json.loads(r.content)

    @staticmethod
    def _oraculo_trend(data):
        """Variación de la comparativa mes a mes (misma base que la tarjeta)."""
        c = data['comparativa_anual']
        actual, anterior = sum(c['actual']), sum(c['anterior'])
        return round((actual - anterior) / anterior * 100, 1) if anterior > 0 else None

    # ---------- A4-01 ----------

    def test_tendencia_con_proveedor_compara_contra_el_mismo_proveedor(self):
        d = self._get(anio=2026, proveedor=self.prov_a.id)
        m = d['metricas']
        # Actual A 2026: 150.000 + 10.000; base A 2025 al 26-09: 100.000 (no la red: 1.100.000).
        self.assertEqual(m['inversion_total'], 160000)
        self.assertEqual(m['trend_inversion'], 60.0)
        self.assertEqual(m['trend_compras'], 100.0)
        self.assertEqual(m['trend_inversion'], self._oraculo_trend(d))

    def test_tendencia_con_temporada_usa_la_misma_temporada(self):
        d = self._get(anio=2026, temporada='Invierno')
        # Invierno 2026: 150.000; Invierno 2025 al 26-09: 100.000.
        self.assertEqual(d['metricas']['trend_inversion'], 50.0)
        self.assertEqual(d['metricas']['trend_inversion'], self._oraculo_trend(d))

    # ---------- A4-02 ----------

    def test_trend_roi_compara_markup_neto_contra_neto(self):
        d = self._get(anio=2026, proveedor=self.prov_a.id)
        m = d['metricas']
        # 2026: (446.250 + 23.800) / 1,19 = 395.000 sobre 160.000 → 146,9 %.
        self.assertEqual(m['roi_promedio'], 146.9)
        # 2025 neto = 100 % (bruto sería 138 % y el delta daría 8,9).
        self.assertEqual(m['trend_roi'], 46.9)

    def test_sin_base_las_tendencias_son_none(self):
        d = self._get(anio=2025)       # no hay compras en 2024
        m = d['metricas']
        self.assertIsNone(m['trend_compras'])
        self.assertIsNone(m['trend_inversion'])
        self.assertIsNone(m['trend_roi'])

    # ---------- A4-03 ----------

    def test_anio_en_curso_compara_hasta_el_mismo_dia(self):
        d = self._get(anio=2026)
        f = d['filtros_aplicados']
        self.assertEqual(f['fecha_desde'], '2026-01-01')
        self.assertEqual(f['fecha_hasta'], '2026-12-31')       # el lado actual no cambia
        self.assertEqual(f['fecha_desde_anterior'], '2025-01-01')
        self.assertEqual(f['fecha_hasta_anterior'], '2025-09-26')
        self.assertIn('al 26-09', f['etiqueta_comparacion'])
        # Base 2025 al 26-09: 100.000 + 1.000.000 (la OC de nov. y la eliminada no cuentan).
        self.assertEqual(d['metricas']['trend_inversion'], -85.5)
        self.assertEqual(d['metricas']['trend_inversion'], self._oraculo_trend(d))
        self.assertEqual(len(d['comparativa_anual']['anterior']), 9)   # ene..sep
        self.assertEqual(len(d['comparativa_anual']['actual']), 12)

    def test_anio_cerrado_compara_contra_el_anio_completo(self):
        d = self._get(anio=2025)
        self.assertEqual(d['filtros_aplicados']['fecha_hasta_anterior'], '2024-12-31')
        self.assertEqual(d['filtros_aplicados']['etiqueta_comparacion'], 'vs 2024')

    # ---------- A4-04 ----------

    def test_rango_mayor_a_un_anio_no_se_compara(self):
        d = self._get(periodo='personalizado', fecha_desde='2025-01-01', fecha_hasta='2026-06-12')
        f = d['filtros_aplicados']
        self.assertFalse(f['comparable'])
        self.assertIsNone(f['fecha_desde_anterior'])
        self.assertIsNone(f['fecha_hasta_anterior'])
        self.assertEqual(d['comparativa_anual']['anterior'], [])
        self.assertIsNone(d['metricas']['trend_inversion'])
        self.assertIsNone(d['metricas']['trend_compras'])
        self.assertNotIn('Crecimiento Positivo', [i['titulo'] for i in d['insights']])

    def test_rango_que_cruza_el_anio_rotula_ambos_anios(self):
        d = self._get(periodo='personalizado', fecha_desde='2025-12-01', fecha_hasta='2026-01-31')
        f = d['filtros_aplicados']
        self.assertTrue(f['comparable'])
        self.assertEqual(f['etiqueta_comparacion'], 'vs mismo período 2024-2025')

    # ---------- A4-05 ----------

    def test_fechas_extremas_se_acotan(self):
        d = self._get(periodo='personalizado', fecha_desde='0001-01-01', fecha_hasta='2026-06-12')
        self.assertEqual(d['filtros_aplicados']['fecha_desde'], '2000-01-01')
        d = self._get(periodo='personalizado', fecha_desde='2026-01-01', fecha_hasta='9999-12-31')
        self.assertEqual(d['filtros_aplicados']['fecha_hasta'], '2027-12-31')
        self.assertEqual(len(d['evolucion_mensual']), 24)

    # ---------- A4-06 ----------

    def test_alerta_proveedores_criticos_cuenta_sobre_todos(self):
        for i in range(14):
            prov = crear_empresa(nombre=f'Prov crítico {i:02d}', rut=f'77.{i:03d}.000-{i % 10}',
                                 esProveedor=True)
            self._compra(prov, date(2024, 6, 1), 10, 1000, 2380)       # 0 % recibido
        d = self._get(anio=2024)
        self.assertEqual(len(d['cumplimiento_proveedores']), 12)        # el gráfico sigue en 12
        critica = [a for a in d['alertas'] if a['titulo'] == 'Proveedores Críticos']
        self.assertEqual(len(critica), 1)
        self.assertTrue(critica[0]['mensaje'].startswith('14 de 14 '), critica[0]['mensaje'])

    # ---------- A4-07 ----------

    def test_periodo_sin_compras_no_alerta_cumplimiento(self):
        d = self._get(periodo='semana')
        self.assertEqual(d['metricas']['unidades_esperadas'], 0)
        self.assertNotIn('Cumplimiento Bajo', [a['titulo'] for a in d['alertas']])
        self.assertNotIn('Mejorar Cumplimiento', [i['titulo'] for i in d['insights']])

    def test_periodo_con_bajo_cumplimiento_si_alerta(self):
        d = self._get(anio=2025)     # 50 de 1.200 unidades recibidas (la eliminada no cuenta)
        self.assertIn('Cumplimiento Bajo', [a['titulo'] for a in d['alertas']])

    # ---------- A4-08 ----------

    def test_presets_son_n_dias_incluido_hoy(self):
        d = self._get(periodo='semana')
        self.assertEqual(d['filtros_aplicados']['fecha_desde'], '2026-09-20')
        self.assertEqual(d['filtros_aplicados']['fecha_hasta'], '2026-09-26')
        d = self._get(periodo='mes')
        self.assertEqual(d['filtros_aplicados']['fecha_desde'], '2026-08-28')
        d = self._get(periodo='trimestre')
        self.assertEqual(d['filtros_aplicados']['fecha_desde'], '2026-06-29')

    # ---------- A4-09 ----------

    def test_proveedor_no_numerico_se_ignora(self):
        d = self._get(anio=2026, proveedor='abc')
        self.assertEqual(d['filtros_aplicados']['proveedor_id'], '')
        self.assertEqual(d['metricas']['total_compras'], 2)

    def test_error_interno_no_expone_traceback(self):
        with mock.patch('app.views_modulo_compras.calcular_metricas_principales_mejorado',
                        side_effect=RuntimeError('detalle interno secreto')), \
                self.assertLogs('app', level='ERROR'):
            r = self.client.get(API, {'anio': 2026})
        self.assertEqual(r.status_code, 500)
        data = json.loads(r.content)
        self.assertNotIn('traceback', data)
        self.assertNotIn('secreto', data['error'])

    # ---------- B15-06 ----------

    def test_desglose_oc_vs_ingresos_manuales(self):
        m = self._get(anio=2026)['metricas']
        oc, manual = m['origen']['oc'], m['origen']['manual']
        self.assertEqual((oc['compras'], manual['compras']), (1, 1))
        self.assertEqual((oc['inversion'], manual['inversion']), (150000.0, 10000.0))
        self.assertEqual(oc['cumplimiento'], 50.0)          # 75 de 150
        self.assertEqual(manual['cumplimiento'], 100.0)     # entra ya recibido
        # Los totales no cambian: suman ambos orígenes.
        self.assertEqual(m['total_compras'], 2)
        self.assertEqual(m['inversion_total'], 160000)
        self.assertEqual(m['cumplimiento_general'], 53.1)   # 85 de 160

    def test_sin_oc_el_cumplimiento_oc_es_none(self):
        m = self._get(periodo='personalizado', fecha_desde='2026-04-01', fecha_hasta='2026-06-30')['metricas']
        self.assertIsNone(m['origen']['oc']['cumplimiento'])
        self.assertEqual(m['origen']['manual']['compras'], 1)

    def test_excel_rotula_origen_y_comparacion(self):
        import openpyxl
        r = self.client.get(EXPORT, {'periodo': 'personalizado',
                                     'fecha_desde': '2025-01-01', 'fecha_hasta': '2026-06-12'})
        self.assertEqual(r.status_code, 200)
        filas = {row[0]: row[1] for row in
                 openpyxl.load_workbook(io.BytesIO(r.content))['Métricas'].iter_rows(values_only=True)
                 if row[0]}
        self.assertEqual(filas['  de ellas: ingresos sin OC (Compra Manual)'], 1)
        self.assertEqual(filas['Comparado con'], 'Sin comparación (rango mayor a 1 año)')

    # ---------- A4-03 (revisión): rango personalizado en curso ----------

    def test_personalizado_en_curso_compara_igual_que_el_anual(self):
        """01-01..31-12 del año en curso da la misma tendencia elegido como
        personalizado que como anual (antes: -85,5 % vs el año anterior
        completo, que en producción daba -96,8 % contra -74,4 %)."""
        anual = self._get(anio=2026)
        pers = self._get(periodo='personalizado', fecha_desde='2026-01-01', fecha_hasta='2026-12-31')
        f = pers['filtros_aplicados']
        self.assertEqual(f['fecha_hasta'], '2026-12-31')             # el lado actual no cambia
        self.assertEqual(f['fecha_hasta_anterior'], '2025-09-26')
        self.assertEqual(f['etiqueta_comparacion'], 'vs mismo período 2025 al 26-09')
        for clave in ('trend_inversion', 'trend_compras', 'trend_roi'):
            self.assertEqual(pers['metricas'][clave], anual['metricas'][clave], clave)
        self.assertEqual(pers['metricas']['trend_inversion'], -85.5)
        self.assertEqual(pers['comparativa_anual']['anterior'], anual['comparativa_anual']['anterior'])

    def test_personalizado_en_curso_de_mas_de_un_anio_sigue_sin_comparar(self):
        # Lo transcurrido (01-10-2025..hoy) cabe en un año, pero el rango pedido
        # no: el recorte no lo vuelve comparable.
        d = self._get(periodo='personalizado', fecha_desde='2025-10-01', fecha_hasta='2026-12-31')
        self.assertFalse(d['filtros_aplicados']['comparable'])
        self.assertIsNone(d['metricas']['trend_inversion'])
        self.assertEqual(d['comparativa_anual']['anterior'], [])

    def test_personalizado_ya_transcurrido_no_se_recorta(self):
        d = self._get(periodo='personalizado', fecha_desde='2026-01-01', fecha_hasta='2026-06-12')
        f = d['filtros_aplicados']
        self.assertEqual(f['fecha_hasta_anterior'], '2025-06-12')
        self.assertEqual(f['etiqueta_comparacion'], 'vs mismo período 2025')

    # ---------- Excel: inyección de fórmulas (revisión M2-R1) ----------

    def test_excel_no_deja_formulas_con_textos_de_usuario(self):
        import openpyxl
        prov = crear_empresa(nombre='=2+2', rut='76.300.300-3', esProveedor=True)
        compra = Compras.objects.create(
            empresa=prov, nombre='=HYPERLINK("http://evil.example","clic")', correlativo=9001,
            responsable='test', temporada='=1+1', fecha=date(2026, 3, 1), estado='ACTIVA',
        )
        cp = Compras_Producto.objects.create(
            compras=compra, nombre="=cmd|' /C calc'!A0", atributo1='@marca', atributo2='',
            atributo3='', atributo4='', costo=5000, precioSugerido=11900,
        )
        Compras_Producto_Talla.objects.create(compra_producto=cp, stock=500, talla='42')

        r = self.client.get(EXPORT, {'anio': 2026})
        self.assertEqual(r.status_code, 200, r.content[:300])
        wb = openpyxl.load_workbook(io.BytesIO(r.content))

        formulas = [(ws.title, c.coordinate, c.value) for ws in wb for row in ws.iter_rows()
                    for c in row if c.data_type == 'f']
        self.assertEqual(formulas, [])

        def celda(hoja, texto):
            for row in wb[hoja].iter_rows():
                for c in row:
                    if c.value == texto:
                        return c
            self.fail(f'{texto!r} no está en la hoja {hoja}')

        for hoja, texto in (('Rendimiento', '=HYPERLINK("http://evil.example","clic")'),
                            ('Rendimiento', '=1+1'),
                            ('Rendimiento', '=2+2'),
                            ('Proveedores', '=2+2'),
                            ('Top Productos', "=cmd|' /C calc'!A0"),
                            ('Top Productos', '@marca')):
            c = celda(hoja, texto)
            self.assertEqual(c.data_type, 's', (hoja, texto))    # texto, se ve tal cual
            self.assertTrue(c.quotePrefix, (hoja, texto))
        # Los números siguen siendo números: inversión 500 × 5.000.
        fila = celda('Rendimiento', '=HYPERLINK("http://evil.example","clic")').row
        self.assertEqual(wb['Rendimiento'].cell(row=fila, column=4).value, 2500000)
