"""
Consolidado de inventario por empresa, solo para el Maestro (pedido del usuario, 09-10):
el cuadro de gerencia «2026 ENERO INV». Por empresa y local: Pares, Costo, P. Interno,
P. Venta y la Diferencia de la toma del mes; total por empresa y «Total Holding».

- Local con toma en el mes (la última por fecha de corte; sin borradores ni canceladas):
  el «inventario nuevo» de su informe final (services/informe_toma_inventario.analizar),
  así que cuadra con el informe de cada toma. Pares = lo contado (lo que no apareció
  cuenta 0; lo resuelto «mantener» toma el sistema); Costo = pares × costo de la ficha;
  P. Interno = pares × (costo + sobreprecio), el «P COSTO» del informe antiguo; P. Venta
  = pares × precio al corte; Diferencia = nuevo − antiguo en pares. Una toma parcial
  (p. ej. solo Calzado) suma además el stock del sistema de lo que quedó fuera de ella
  (NICK2 09-10: 7.794 pares contados de ~27 mil), y lo informa aparte.
- Local sin toma en el mes: stock del sistema (al cierre del mes si ya pasó; si no, el
  de hoy), sin diferencia. Mismo cálculo que Resumen de Existencias (sin los productos
  marcados «excluir de analítica»).
"""
from collections import OrderedDict, defaultdict
from datetime import datetime, timedelta

from django.utils import timezone

from app.models import Producto_Talla, TomaInventario
from app.services import informe_toma_inventario as informe_toma

MESES = ['ENERO', 'FEBRERO', 'MARZO', 'ABRIL', 'MAYO', 'JUNIO', 'JULIO', 'AGOSTO',
         'SEPTIEMBRE', 'OCTUBRE', 'NOVIEMBRE', 'DICIEMBRE']

# Una toma en borrador no se ha contado y una cancelada no vale: no entran al cuadro
ESTADOS_FUERA = ('BORRADOR', 'CANCELADO')

VALORES = ('pares', 'costo', 'p_interno', 'p_venta')


def rango_mes(anio, mes):
    """[inicio, fin) del mes en hora de Chile."""
    tz = timezone.get_current_timezone()
    inicio = timezone.make_aware(datetime(anio, mes, 1), tz)
    fin = timezone.make_aware(datetime(anio + (mes == 12), mes % 12 + 1, 1), tz)
    return inicio, fin


def titulo(anio, mes):
    return f'{anio} {MESES[mes - 1]} INV'


def _nombre_local(sucursal):
    """«MATTA 2422» para una tienda (como el cuadro de gerencia); «BODEGA PA00» para un CD."""
    if sucursal.es_centro_distribucion:
        return f'BODEGA {sucursal.alias}'.upper()
    return (sucursal.direccion or '').strip().upper() or sucursal.alias


def _dia_del_stock(fin):
    """Día al que se lee el stock del sistema: el cierre del mes, o hoy si el mes no terminó."""
    return min(timezone.localtime(fin - timedelta(seconds=1)).date(), timezone.localdate())


def _stock_sistema(tallas, dia):
    """
    {sucursal_id: {pares, costo, p_interno, p_venta}} del sistema al cierre de `dia`, sin lo
    marcado «excluir de analítica». Mismo cálculo que Resumen de Existencias: con `dia` = hoy
    no hay movimientos posteriores y es el stock actual.
    """
    # import local: views_resumen_existencias es un módulo de vistas
    from app.views_resumen_existencias import _acumulados_historicos
    return _acumulados_historicos(tallas.filter(producto__excluir_de_analitica=False), dia)


def _fuera_de_la_toma(toma, dia):
    """Stock del sistema de lo que una toma parcial no cubre (las tallas que no tiene)."""
    en_toma = toma.detalles.filter(producto_talla_id__isnull=False).values('producto_talla_id')
    tallas = Producto_Talla.objects.filter(producto__sucursal_id=toma.sucursal_id).exclude(id__in=en_toma)
    return _stock_sistema(tallas, dia).get(toma.sucursal_id)


def _local_con_toma(sucursal, toma, otras_en_el_mes, dia):
    analisis = informe_toma.analizar(informe_toma.filas_desde_toma(toma))
    total = analisis['total']
    alcance, _ = informe_toma.alcance_de_toma(toma)
    contado = {
        'pares': int(total['nue_stock']),
        'costo': int(round(total['nue_costo_puro'])),
        'p_interno': int(round(total['nue_costo'])),
        'p_venta': int(round(total['nue_venta'])),
    }
    # Una toma completa es la foto del local al corte: lo que no tiene llegó después. Una
    # parcial no miró el resto: va el stock del sistema para que el local quede entero.
    fuera = None
    if toma.tipo_inventario != 'COMPLETO':
        fuera = {k: int((_fuera_de_la_toma(toma, dia) or {}).get(k) or 0) for k in VALORES}
    return {
        **{k: contado[k] + (fuera[k] if fuera else 0) for k in VALORES},
        'contado': contado,
        'fuera_de_la_toma': fuera,
        'pares_sistema': int(total['ant_stock']),
        'diferencia': int(total['dif']),
        'toma': {
            'id': toma.id,
            'numero': toma.numero_inventario,
            'estado': toma.estado,
            'estado_display': toma.get_estado_display(),
            'corte': timezone.localtime(toma.fecha_corte).strftime('%d/%m/%Y %H:%M'),
            'tipo': informe_toma.TIPO_EN_TITULO.get(toma.tipo_inventario, toma.get_tipo_inventario_display()),
            'alcance': alcance,
            'sin_contar': analisis['resumen']['sin_contar'],
            'otras_en_el_mes': otras_en_el_mes,
        },
    }


def _local_sin_toma(stock):
    stock = stock or {}
    return {**{k: int(stock.get(k) or 0) for k in VALORES}, 'contado': None, 'fuera_de_la_toma': None,
            'pares_sistema': None, 'diferencia': None, 'toma': None}


def _sumar(locales):
    total = {k: sum(l[k] for l in locales) for k in VALORES}
    total['diferencia'] = sum(l['diferencia'] or 0 for l in locales)
    total['locales'] = len(locales)
    total['con_toma'] = sum(1 for l in locales if l['toma'])
    return total


def consolidar(sucursales, anio, mes):
    """
    `sucursales`: queryset de las sucursales del usuario. Devuelve el cuadro:
    {titulo, anio, mes, empresas: [{id, nombre, solo_bodegas, locales, total}], total}.
    Empresas con tiendas primero; dentro, las tiendas por alias y las bodegas al final.
    """
    inicio, fin = rango_mes(anio, mes)
    sucursales = sorted(sucursales.select_related('empresa'), key=lambda s: (s.es_centro_distribucion, s.alias))

    tomas, otras = {}, defaultdict(int)
    candidatas = (TomaInventario.objects
                  .filter(sucursal_id__in=[s.id for s in sucursales], fecha_corte__gte=inicio, fecha_corte__lt=fin)
                  .exclude(estado__in=ESTADOS_FUERA).order_by('-fecha_corte', '-id'))
    for t in candidatas:
        if t.sucursal_id in tomas:
            otras[t.sucursal_id] += 1
        else:
            tomas[t.sucursal_id] = t

    dia = _dia_del_stock(fin)
    sin_toma = [s.id for s in sucursales if s.id not in tomas]
    sistema = _stock_sistema(Producto_Talla.objects.filter(producto__sucursal_id__in=sin_toma), dia) if sin_toma else {}

    empresas = OrderedDict()
    for s in sucursales:
        if s.id in tomas:
            local = _local_con_toma(s, tomas[s.id], otras[s.id], dia)
        else:
            local = _local_sin_toma(sistema.get(s.id))
        local.update({'sucursal_id': s.id, 'alias': s.alias, 'local': _nombre_local(s),
                      'direccion': (s.direccion or '').strip(), 'es_bodega': s.es_centro_distribucion})
        bloque = empresas.setdefault(s.empresa_id, {
            'id': s.empresa_id, 'nombre': (s.empresa.nombre or '').strip().upper(),
            'solo_bodegas': True, 'locales': [],
        })
        bloque['locales'].append(local)
        bloque['solo_bodegas'] = bloque['solo_bodegas'] and s.es_centro_distribucion

    bloques = sorted(empresas.values(), key=lambda b: (b['solo_bodegas'], b['id']))
    for b in bloques:
        b['total'] = _sumar(b['locales'])
    return {
        'titulo': titulo(anio, mes),
        'anio': anio,
        'mes': mes,
        'dia_stock': dia.strftime('%d/%m/%Y'),
        'empresas': bloques,
        'total': _sumar([l for b in bloques for l in b['locales']]),
    }


# ---------------------------------------------------------------------------
# Excel (mismo formato del cuadro de gerencia)
# ---------------------------------------------------------------------------

def nombre_archivo(datos):
    return f"inventario_por_empresa_{datos['anio']}_{datos['mes']:02d}.xlsx"


def construir_workbook(datos):
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

    PESOS = '"$" #,##0;[Red]-"$" #,##0;"$" -'
    UNIDADES = '#,##0;[Red]-#,##0;0'
    DIFERENCIA = '+#,##0;[Red]-#,##0;0'
    fino = Side(style='thin', color='808080')
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    negrita = Font(bold=True)
    total_font = Font(bold=True, color='C00000')
    centro = Alignment(horizontal='center', vertical='center')
    cabecera_fill = PatternFill(start_color='DDE3EF', end_color='DDE3EF', fill_type='solid')
    columnas = ['LOCAL', 'Toma', 'Pares', 'Costo', 'P. Interno', 'P. Venta', 'Diferencia']
    ultima = 'G'

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = 'Por empresa'
    ws.merge_cells(f'A1:{ultima}1')
    ws['A1'] = datos['titulo']
    ws['A1'].font = Font(bold=True, size=16, underline='single')
    ws['A1'].alignment = centro

    def _fila(n, nombre, toma_txt, valores, diferencia, font=None):
        celdas = [nombre, toma_txt, valores['pares'], valores['costo'], valores['p_interno'], valores['p_venta'], diferencia]
        for col, valor in enumerate(celdas, 1):
            c = ws.cell(row=n, column=col, value=valor)
            c.border = borde
            if col == 3:
                c.number_format = UNIDADES
            elif col in (4, 5, 6):
                c.number_format = PESOS
            elif col == 7:
                c.number_format = DIFERENCIA
                c.alignment = centro
            if font:
                c.font = font

    n = 3
    for b in datos['empresas']:
        ws.merge_cells(f'A{n}:{ultima}{n}')
        ws.cell(row=n, column=1, value=b['nombre']).font = Font(bold=True, size=12)
        ws.cell(row=n, column=1).alignment = centro
        n += 1
        for col, texto in enumerate(columnas, 1):
            c = ws.cell(row=n, column=col, value=texto)
            c.font = negrita
            c.alignment = centro
            c.border = borde
            c.fill = cabecera_fill
        n += 1
        for l in b['locales']:
            t = l['toma']
            if t:
                toma_txt = f"{t['numero']} · {t['estado_display']}" + (f" · {t['alcance']}" if t['alcance'] else '')
                if l['fuera_de_la_toma']:
                    toma_txt += (f" · contados {l['contado']['pares']:,} + {l['fuera_de_la_toma']['pares']:,} "
                                 f"fuera de la toma (sistema)").replace(',', '.')
            else:
                toma_txt = 'Sin toma: stock del sistema'
            _fila(n, l['local'], toma_txt, l, l['diferencia'])
            n += 1
        _fila(n, 'TOTAL', f"{b['total']['con_toma']} de {b['total']['locales']} con toma", b['total'],
              b['total']['diferencia'], font=total_font)
        n += 2

    _fila(n, 'Total Holding', f"{datos['total']['con_toma']} de {datos['total']['locales']} con toma", datos['total'],
          datos['total']['diferencia'], font=Font(bold=True, size=12, color='C00000'))
    n += 2
    notas = [
        'Con toma: inventario nuevo del informe final de la toma (lo contado). Diferencia = contado − sistema, en pares.',
        'Toma parcial (p. ej. solo Calzado): se suma el stock del sistema de lo que quedó fuera de la toma.',
        f"Sin toma en el mes: stock del sistema al {datos['dia_stock']}, sin diferencia "
        '(sin los productos marcados «excluir de analítica»).',
        'Costo = pares × costo · P. Interno = pares × (costo + sobreprecio), el P COSTO del informe antiguo · '
        'P. Venta = pares × precio de venta.',
    ]
    for texto in notas:
        ws.cell(row=n, column=1, value=texto).font = Font(italic=True, size=9, color='595959')
        n += 1

    for letra, ancho in {'A': 30, 'B': 44, 'C': 11, 'D': 17, 'E': 17, 'F': 17, 'G': 12}.items():
        ws.column_dimensions[letra].width = ancho
    return wb
