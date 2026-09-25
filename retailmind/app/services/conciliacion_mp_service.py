"""Conciliación Mercado Pago: cobros ↔ documentos ↔ liberaciones ↔ banco.

Cuatro preguntas, cuatro funciones:

1. ``cobros_vs_documentos``: cada cobro MP aprobado, ¿terminó en una venta con
   boleta/factura? Solo datos locales (no llama a MP): sirve para el arqueo.
2. ``procesar_reporte_liberaciones``: el reporte de Liberaciones de MP dice qué
   pagos componen cada retiro al banco. Crea ``RetiroMercadoPago`` y amarra
   cada ``TransaccionMercadoPago`` a su retiro.
3. ``conciliar_cartola``: marca "visto en cartola" los retiros cuyo abono
   aparece en el extracto del banco (mismo monto, fecha ±N días).
4. ``diferencias_contra_mp``: cruce contra la API (payments/search) para ver
   cobros que MP tiene y el sistema no, y el cierre por caja/día.

Todo es READ-ONLY salvo (2) y (3) con ``aplicar=True``.
"""
import csv
import io
import logging
import os
import re
from datetime import date, datetime, timedelta, timezone as dt_timezone
from decimal import Decimal, InvalidOperation, ROUND_HALF_UP

from django.core.cache import cache
from django.db import IntegrityError, transaction
from django.db.models import Count, Min, Q, Sum
from django.db.models.functions import Coalesce
from django.utils import timezone

from app.models import (
    ArqueoCaja,
    Dte,
    MercadoPagoConfig,
    RetiroMercadoPago,
    Ticket,
    TicketDetallePago,
    TransaccionMercadoPago,
)
from app.services import mercadopago_service as mp

logger = logging.getLogger('app')

METODOS_MP = ('MP_QR', 'MP_POINT', 'MP_POINT_DEBITO', 'MP_POINT_CREDITO')
TIPOS_DTE_REINTENTABLES = ('BOLETA_ELECTRONICA', 'BOLETA_PAPEL', 'FACTURA_ELECTRONICA')

# Categorías de cobros_vs_documentos (orden = prioridad de revisión).
CAT_SIN_VENTA = 'SIN_VENTA'
CAT_TICKET_ANULADO = 'TICKET_ANULADO'
CAT_SIN_DOCUMENTO = 'SIN_DOCUMENTO'
CAT_COBRO_DIRECTO = 'COBRO_DIRECTO'
CAT_MANUAL = 'MANUAL'
CAT_CON_DOCUMENTO = 'CON_DOCUMENTO'
CATEGORIAS = {
    CAT_SIN_VENTA: 'Cobro aprobado que ninguna venta usa',
    CAT_TICKET_ANULADO: 'Cobrado con la venta anulada (devolver)',
    CAT_SIN_DOCUMENTO: 'Venta pagada sin boleta/factura',
    CAT_COBRO_DIRECTO: 'Cobro directo en terminal (sin ticket)',
    CAT_MANUAL: 'Registrado a mano (verificar contra MP)',
    CAT_CON_DOCUMENTO: 'Con documento',
}


def _fecha(valor, por_defecto):
    if isinstance(valor, date):
        return valor
    try:
        return datetime.strptime(str(valor), '%Y-%m-%d').date()
    except (TypeError, ValueError):
        return por_defecto


def rango_fechas(desde, hasta, dias_defecto=7, max_dias=93):
    """(desde, hasta) como date, con defaults y un tope de días."""
    hoy = timezone.localdate()
    h = _fecha(hasta, hoy)
    d = _fecha(desde, h - timedelta(days=dias_defecto - 1))
    if d > h:
        d, h = h, d
    if (h - d).days >= max_dias:
        d = h - timedelta(days=max_dias - 1)
    return d, h


# ───────────────────────────── 1. Cobros ↔ documentos ─────────────────────────

def _documento_de_ticket(ticket, dtes_por_ref):
    """(texto, dte_id) del documento tributario del ticket, o ('', None)."""
    if ticket is None:
        return '', None
    dte = dtes_por_ref.get((ticket.sucursal_id, ticket.correlativo))
    if dte is not None:
        return f'{dte.tipo_documento} #{dte.numero_documento}', dte.id
    if ticket.folio_dte:
        return f'{(ticket.tipo_dte or "DTE").replace("_", " ")} #{ticket.folio_dte}', None
    return '', None


def _dtes_por_ref(tickets):
    """{(sucursal_id, correlativo): Dte} por la referencia exacta 'TICKET-<corr>' (una consulta)."""
    dtes_por_ref = {}
    tickets = [tk for tk in tickets if tk is not None]
    if not tickets:
        return dtes_por_ref
    refs = {f'TICKET-{tk.correlativo}' for tk in tickets}
    sucs = {tk.sucursal_id for tk in tickets}
    for dte in (Dte.objects.filter(referencias__in=refs, sucursal_id__in=sucs)
                .exclude(descartado=True).only('id', 'tipo_documento', 'numero_documento',
                                               'referencias', 'sucursal_id')):
        try:
            corr = int(str(dte.referencias).split('-', 1)[1])
        except (IndexError, ValueError):
            continue
        dtes_por_ref[(dte.sucursal_id, corr)] = dte
    return dtes_por_ref


def _ticket_de_cobros(trxs):
    """{trx.id: Ticket|None}: FK directa, vía el pago, o por (sucursal, correlativo)."""
    salida, faltan = {}, {}
    for t in trxs:
        if t.ticket_id:
            salida[t.id] = t.ticket
        elif t.detalle_pago_id and t.detalle_pago.ticket_id:
            salida[t.id] = t.detalle_pago.ticket
        else:
            salida[t.id] = None
            corr = str(t.correlativo_ticket or '')
            if corr.isdigit():
                faltan[(t.sucursal_id, int(corr))] = None
    if faltan:
        q = Q()
        for suc, corr in faltan:
            q |= Q(sucursal_id=suc, correlativo=corr)
        for tk in Ticket.objects.filter(q):
            faltan[(tk.sucursal_id, tk.correlativo)] = tk
        for t in trxs:
            corr = str(t.correlativo_ticket or '')
            if salida[t.id] is None and corr.isdigit():
                salida[t.id] = faltan.get((t.sucursal_id, int(corr)))
    return salida


def cobros_vs_documentos(desde, hasta, sucursal_id=None):
    """Cada cobro MP del período clasificado según si terminó en documento.

    Devuelve {'filas': [...], 'resumen': {cat: {'cantidad', 'monto'}},
    'desde', 'hasta'}. Incluye los pagos MP registrados a MANO en el POS
    (origen MANUAL), que no tienen transacción: van en la categoría MANUAL.
    """
    d, h = rango_fechas(desde, hasta)
    trxs = list(
        TransaccionMercadoPago.objects
        .filter(tipo='VENTA', estado__in=('APROBADA', 'DEVUELTA', 'CONTRACARGO'),
                creado_en__date__gte=d, creado_en__date__lte=h)
        .exclude(correlativo_ticket__startswith='PRUEBA-')
        .select_related('sucursal', 'config', 'ticket', 'detalle_pago__ticket', 'retiro')
        .order_by('-creado_en')
    )
    if sucursal_id:
        trxs = [t for t in trxs if t.sucursal_id == int(sucursal_id)]

    # Ticket de cada cobro: FK directa, vía el pago, o por (sucursal, correlativo).
    faltan = {}
    for t in trxs:
        if t.ticket_id is None and not (t.detalle_pago_id and t.detalle_pago.ticket_id):
            corr = str(t.correlativo_ticket or '')
            if corr.isdigit():
                faltan[(t.sucursal_id, int(corr))] = None
    if faltan:
        q = Q()
        for suc, corr in faltan:
            q |= Q(sucursal_id=suc, correlativo=corr)
        for tk in Ticket.objects.filter(q):
            faltan[(tk.sucursal_id, tk.correlativo)] = tk

    def _ticket(t):
        if t.ticket_id:
            return t.ticket
        if t.detalle_pago_id and t.detalle_pago.ticket_id:
            return t.detalle_pago.ticket
        corr = str(t.correlativo_ticket or '')
        return faltan.get((t.sucursal_id, int(corr))) if corr.isdigit() else None

    tickets = {t.id: _ticket(t) for t in trxs}

    # Pagos MP registrados a mano en el período (sin transacción MP).
    manuales = list(
        TicketDetallePago.objects
        .filter(metodo_pago__in=METODOS_MP, origen_pago='MANUAL',
                ticket__estado='PAGADO', transacciones_mercadopago__isnull=True,
                creado_en__date__gte=d, creado_en__date__lte=h)
        .select_related('ticket', 'ticket__sucursal')
        .order_by('-creado_en')
    )
    if sucursal_id:
        manuales = [p for p in manuales if p.ticket.sucursal_id == int(sucursal_id)]

    dtes_por_ref = _dtes_por_ref([tk for tk in tickets.values() if tk] + [p.ticket for p in manuales])

    filas = []
    for t in trxs:
        tk = tickets[t.id]
        documento, dte_id = _documento_de_ticket(tk, dtes_por_ref)
        corr = str(t.correlativo_ticket or '')
        if t.estado != 'APROBADA':
            categoria = CAT_CON_DOCUMENTO if documento else CAT_SIN_VENTA
        elif corr.startswith('DIRECTO-') and not t.consumida:
            categoria = CAT_COBRO_DIRECTO
        elif tk is not None and tk.estado in ('ANULADO', 'DEVUELTO'):
            categoria = CAT_TICKET_ANULADO
        elif not t.consumida:
            # Aprobado y sin respaldar ningún pago: aunque el ticket esté PAGADO
            # (se cobró con otro medio), esta plata no está en la venta. Es el
            # caso NICK2 05-09: MP cobró y la venta se registró como tarjeta.
            categoria = CAT_SIN_VENTA
        elif documento:
            categoria = CAT_CON_DOCUMENTO
        else:
            categoria = CAT_SIN_DOCUMENTO
        filas.append({
            'origen': 'COBRO_MP',
            'id': t.id,
            'fecha': timezone.localtime(t.creado_en).strftime('%Y-%m-%d %H:%M'),
            'sucursal': t.sucursal.alias if t.sucursal_id else '',
            'sucursal_id': t.sucursal_id,
            'config_id': t.config_id,
            'caja': t.config.nombre if t.config_id else '',
            'correlativo': corr,
            'monto': t.monto,
            'monto_neto': t.monto_neto,
            'estado_mp': t.estado,
            'medio': mp.etiqueta_medio_mp(t.metodo_pago_mp),
            'payment_id_mp': t.payment_id_mp or '',
            'canal': t.canal,
            'ticket_id': tk.id if tk else None,
            'ticket_correlativo': tk.correlativo if tk else None,
            'ticket_estado': tk.estado if tk else '',
            'documento': documento,
            'dte_id': dte_id,
            'dte_fallido': bool(tk and tk.dte_generacion_fallida),
            'dte_error': (tk.dte_error_detalle or '')[:200] if tk else '',
            'puede_reintentar_dte': bool(
                categoria == CAT_SIN_DOCUMENTO and tk and tk.estado == 'PAGADO'
                and not tk.folio_dte and tk.tipo_dte in TIPOS_DTE_REINTENTABLES
            ),
            'liberacion': (timezone.localtime(t.money_release_date).strftime('%Y-%m-%d')
                           if t.money_release_date else ''),
            'retiro': t.retiro.withdrawal_id if t.retiro_id else '',
            'categoria': categoria,
        })
    for p in manuales:
        documento, dte_id = _documento_de_ticket(p.ticket, dtes_por_ref)
        filas.append({
            'origen': 'PAGO_MANUAL',
            'id': p.id,
            'fecha': timezone.localtime(p.creado_en).strftime('%Y-%m-%d %H:%M'),
            'sucursal': p.ticket.sucursal.alias if p.ticket.sucursal_id else '',
            'sucursal_id': p.ticket.sucursal_id,
            'config_id': None,
            'caja': '',
            'correlativo': str(p.ticket.correlativo),
            'monto': p.monto,
            'monto_neto': None,
            'estado_mp': 'MANUAL',
            'medio': p.tipo_tarjeta or p.metodo_pago,
            'payment_id_mp': p.voucher or '',
            'canal': 'MANUAL',
            'ticket_id': p.ticket_id,
            'ticket_correlativo': p.ticket.correlativo,
            'ticket_estado': p.ticket.estado,
            'documento': documento,
            'dte_id': dte_id,
            'dte_fallido': bool(p.ticket.dte_generacion_fallida),
            'dte_error': '',
            'puede_reintentar_dte': False,
            'liberacion': '',
            'retiro': '',
            'categoria': CAT_MANUAL,
        })

    orden = list(CATEGORIAS)
    filas.sort(key=lambda f: (orden.index(f['categoria']), f['fecha']), reverse=False)
    resumen = {c: {'etiqueta': CATEGORIAS[c], 'cantidad': 0, 'monto': 0} for c in CATEGORIAS}
    for f in filas:
        if f['estado_mp'] in ('DEVUELTA', 'CONTRACARGO'):
            continue
        resumen[f['categoria']]['cantidad'] += 1
        resumen[f['categoria']]['monto'] += int(f['monto'] or 0)
    return {'desde': str(d), 'hasta': str(h), 'filas': filas, 'resumen': resumen}


# ───────────────────────── 2. Reporte de Liberaciones → retiros ───────────────

_ALIAS_COLUMNAS = {
    'DATE': ('DATE', 'FECHA', 'FECHA DE LIBERACION', 'RELEASE_DATE'),
    'SOURCE_ID': ('SOURCE_ID', 'ID DE OPERACION', 'ID DE LA OPERACION', 'OPERATION_ID'),
    'EXTERNAL_REFERENCE': ('EXTERNAL_REFERENCE', 'REFERENCIA EXTERNA', 'CODIGO DE REFERENCIA'),
    'RECORD_TYPE': ('RECORD_TYPE', 'TIPO DE REGISTRO'),
    'DESCRIPTION': ('DESCRIPTION', 'DESCRIPCION', 'TIPO DE OPERACION'),
    'NET_CREDIT_AMOUNT': ('NET_CREDIT_AMOUNT', 'NET_CREDIT', 'CREDITO NETO', 'MONTO NETO ACREDITADO'),
    'NET_DEBIT_AMOUNT': ('NET_DEBIT_AMOUNT', 'NET_DEBIT', 'DEBITO NETO', 'MONTO NETO DEBITADO'),
    'GROSS_AMOUNT': ('GROSS_AMOUNT', 'MONTO BRUTO'),
    'EXTERNAL_POS_ID': ('EXTERNAL_POS_ID', 'ID EXTERNO DE LA CAJA', 'ID EXTERNO DEL PUNTO DE VENTA'),
    'POS_ID': ('POS_ID', 'ID DE LA CAJA', 'ID DEL PUNTO DE VENTA'),
    'POS_NAME': ('POS_NAME', 'NOMBRE DE LA CAJA', 'NOMBRE DEL PUNTO DE VENTA'),
}
# Comparación EXACTA (en minúsculas): 'tax_withholding_payout' es una
# retención sobre el retiro, NO el retiro; con búsqueda por substring contaba
# como retiro. Glosario oficial: payout = "withdrawal of available money".
DESCRIPCIONES_RETIRO = ('payout', 'withdrawal', 'retiro')
DESCRIPCIONES_PAGO = ('payment', 'pago')
# 'reserve_for_payout': segundos antes del payout MP aparta la plata (débito) y
# la suelta (crédito), las dos filas con el N° del retiro. Suman cero y no son
# pagos: leído el crédito como liberación sin venta, el retiro entero quedaba en
# «Sin caja (online…)» (Paola, retiro de $6.147.490 del 23-09).
DESCRIPCIONES_INTERNAS = ('reserve_for_payout', 'reserva para retiro', 'reserva_para_retiro')


def _solo_digitos(valor):
    """N° de operación tal como lo digitó el cajero ('179 000-000') → solo dígitos."""
    return re.sub(r'\D', '', str(valor or ''))


def _distancia_edicion(a, b):
    """Cambios de a un carácter (quitar, poner, cambiar o invertir dos vecinos)
    para pasar de `a` a `b` (Damerau-Levenshtein restringida)."""
    previa, actual = None, list(range(len(b) + 1))
    for i in range(1, len(a) + 1):
        anterior, previa, actual = previa, actual, [i] + [0] * len(b)
        for j in range(1, len(b) + 1):
            costo = 0 if a[i - 1] == b[j - 1] else 1
            actual[j] = min(previa[j] + 1, actual[j - 1] + 1, previa[j - 1] + costo)
            if i > 1 and j > 1 and a[i - 1] == b[j - 2] and a[i - 2] == b[j - 1]:
                actual[j] = min(actual[j], anterior[j - 2] + 1)
    return actual[len(b)]


def n_parecido(digitado, real):
    """Si el N° de operación `digitado` parece el `real` con UN error de tipeo,
    dice cuál («falta un 8»); si no, ''. Solo para N° largos (8+ dígitos): los
    N° de MP de una misma cuenta son casi correlativos y con N° cortos cualquier
    cosa se parece. Caso real 21-09: 18014409060 digitado por 180148409060."""
    a, b = _solo_digitos(digitado), _solo_digitos(real)
    if not a or not b or a == b or min(len(a), len(b)) < 8 or abs(len(a) - len(b)) > 1:
        return ''
    if _distancia_edicion(a, b) != 1:
        return ''
    corto = min(len(a), len(b))
    i = next((k for k in range(corto) if a[k] != b[k]), corto)
    if len(a) < len(b):
        return f'falta un {b[i]}'
    if len(a) > len(b):
        return f'sobra un {a[i]}'
    if i + 1 < len(a) and a[i] == b[i + 1] and a[i + 1] == b[i]:
        return f'«{b[i]}{b[i + 1]}» escrito al revés'
    return f'dice {a[i]} donde va {b[i]}'


def _sin_tildes(texto):
    return (str(texto or '').upper()
            .replace('Á', 'A').replace('É', 'E').replace('Í', 'I')
            .replace('Ó', 'O').replace('Ú', 'U').replace('Ñ', 'N').strip())


def _monto(valor):
    """'1.234.567', '1234567.00', '-5.000', '$ 12.000' → int (CLP)."""
    texto = str(valor or '').strip().replace('$', '').replace(' ', '')
    if not texto:
        return 0
    negativo = texto.startswith('-') or (texto.startswith('(') and texto.endswith(')'))
    texto = texto.strip('-()')
    if re.fullmatch(r'\d{1,3}(\.\d{3})+(,\d+)?', texto):      # 1.234.567,00
        texto = texto.replace('.', '').replace(',', '.')
    elif re.fullmatch(r'\d{1,3}(,\d{3})+(\.\d+)?', texto):    # 1,234,567.00
        texto = texto.replace(',', '')
    else:
        texto = texto.replace(',', '.')
    try:
        # Half-up, nunca round()/int(): round(1234.5) da 1234 (redondeo bancario).
        n = int(Decimal(texto).quantize(Decimal('1'), rounding=ROUND_HALF_UP))
    except (InvalidOperation, ValueError):
        return 0
    return -n if negativo else n


def _fecha_libre(valor):
    """ISO con/sin hora, dd/mm/yyyy o dd-mm-yyyy → date (o None)."""
    texto = str(valor or '').strip()
    if not texto:
        return None
    m = re.match(r'(\d{4})-(\d{2})-(\d{2})', texto)
    if m:
        return date(int(m.group(1)), int(m.group(2)), int(m.group(3)))
    m = re.match(r'(\d{1,2})[/-](\d{1,2})[/-](\d{2,4})', texto)
    if m:
        anio = int(m.group(3))
        anio = anio + 2000 if anio < 100 else anio
        try:
            return date(anio, int(m.group(2)), int(m.group(1)))
        except ValueError:
            return None
    return None


def leer_csv(contenido):
    """Texto CSV (',' o ';') → lista de dicts con claves normalizadas (mayúsculas, sin tildes)."""
    if isinstance(contenido, bytes):
        for enc in ('utf-8-sig', 'latin-1'):
            try:
                contenido = contenido.decode(enc)
                break
            except UnicodeDecodeError:
                continue
    muestra = contenido[:4096]
    delimitador = ';' if muestra.count(';') > muestra.count(',') else ','
    lector = csv.DictReader(io.StringIO(contenido), delimiter=delimitador)
    filas = []
    for fila in lector:
        filas.append({_sin_tildes(k): (v or '').strip() for k, v in fila.items() if k})
    return filas


def _col(fila, clave):
    for alias in _ALIAS_COLUMNAS[clave]:
        if alias in fila and fila[alias] != '':
            return fila[alias]
    return ''


def _instante(valor):
    """DATE del reporte ('2026-09-23T10:48:00.000-04:00') → datetime aware (o None).

    Se ordena por instante real y no por texto: el reporte puede mezclar
    offsets (-04:00 / -03:00 según horario de verano) y el orden alfabético
    pondría un retiro antes de los pagos que lo componen.
    """
    texto = str(valor or '').strip()
    if not texto:
        return None
    try:
        dt = datetime.fromisoformat(texto.replace('Z', '+00:00'))
    except ValueError:
        f = _fecha_libre(texto)
        if f is None:
            return None
        dt = datetime(f.year, f.month, f.day)
    if timezone.is_naive(dt):
        dt = timezone.make_aware(dt, timezone.get_current_timezone())
    return dt


def _cuenta_efectiva(config):
    """Cuenta MP de la caja con la misma regla que usa el cobro (`mp._cuenta_de`)."""
    try:
        return mp._cuenta_de(config)
    except AttributeError:   # caja sin sucursal ni cuenta
        return None


def _configs_de_la_cuenta(config):
    """Cajas que comparten la cuenta MP de `config` (el reporte es por cuenta).

    Misma regla que el cobro: la cuenta elegida en la caja o, si quedó en
    «automática» (FK vacío), la cuenta activa de la empresa. Antes solo miraba
    el FK y con «automática» quedaba UNA caja: las ventas de las otras cajas de
    la misma cuenta no se cruzaban ni se completaban.
    """
    cuenta = _cuenta_efectiva(config)
    if cuenta is None:
        return [config.id]
    ids = [config.id]
    for c in MercadoPagoConfig.objects.exclude(id=config.id).select_related('cuenta', 'sucursal'):
        otra = _cuenta_efectiva(c)
        if otra is not None and otra.id == cuenta.id:
            ids.append(c.id)
    return ids


def _remanente_previo(config, antes_de, monto, excluir_ids=(), retiros_del_reporte=(), piso=None):
    """Ventas liberadas ANTES del reporte y todavía sin retiro que explican el
    saldo inicial: las más recientes hacia atrás hasta cubrir `monto`.

    Sin esto, lo que quedó en MP después de un retiro parcial era un saldo
    anónimo en el reporte siguiente y esas ventas nunca se asociaban al retiro
    que finalmente se las llevó.

    Cuentan también las ventas que ya están amarradas a un retiro que viene en
    este mismo reporte: al reprocesar (el reporte pedido parte el día del
    último retiro) ese saldo ES lo que ese retiro se llevó. Si solo se miraban
    las sin retiro, el saldo se rellenaba con ventas viejas ajenas que quedaban
    amarradas a ese retiro, y el error crecía en cada pasada.

    `piso`: la liberación más antigua entre las ventas de esos retiros. Una
    venta sin retiro liberada ANTES ya salió de MP (FIFO: un retiro se lleva
    primero lo más antiguo), aunque en su momento no se haya podido cruzar:
    no sirve para rellenar el saldo. Lo que falte queda como «Saldo anterior».
    Devuelve [(trx, neto)] en orden de liberación (más antigua primero).
    """
    if monto <= 0 or antes_de is None:
        return []
    qs = (TransaccionMercadoPago.objects
          .filter(Q(retiro__isnull=True) | Q(retiro__withdrawal_id__in=list(retiros_del_reporte)),
                  config_id__in=_configs_de_la_cuenta(config), tipo='VENTA', estado='APROBADA',
                  money_release_date__isnull=False,
                  money_release_date__lt=antes_de,
                  money_release_date__gte=antes_de - timedelta(days=90))
          .exclude(id__in=list(excluir_ids))
          .select_related('config__sucursal')
          .order_by('-money_release_date', '-id'))
    if piso is not None:
        qs = qs.filter(Q(retiro__withdrawal_id__in=list(retiros_del_reporte)) | Q(money_release_date__gte=piso))
    elegidas, suma = [], 0
    for t in qs[:20000].iterator(chunk_size=500):
        neto = int(t.monto_neto if t.monto_neto is not None else t.monto)
        if neto <= 0:
            continue
        usa = min(neto, monto - suma)
        elegidas.append((t, usa))
        suma += usa
        if suma >= monto:
            break
    return list(reversed(elegidas))


def reportes_aplicados():
    """file_name de los reportes ya aplicados (lista acumulativa en los retiros)."""
    aplicados = set()
    for raw in (RetiroMercadoPago.objects.filter(raw_reporte__isnull=False)
                .order_by('-fecha').values_list('raw_reporte', flat=True)[:1000]):
        if not isinstance(raw, dict):
            continue
        aplicados.update(a for a in (raw.get('archivos') or []) if a)
        if raw.get('archivo'):
            aplicados.add(raw['archivo'])  # formato anterior
    return aplicados


CAJA_SALDO_ANTERIOR = 'Saldo anterior'
CAJA_NO_EXPLICADA = 'No explicado por el reporte'


def _no_explicado(por_caja):
    """Parte de un retiro que el reporte no alcanzó a explicar: el saldo con que
    partió («Saldo anterior») o un faltante («No explicado por el reporte»)."""
    return sum(int(x.get('monto') or 0) for x in (por_caja or [])
               if isinstance(x, dict) and x.get('caja') in (CAJA_SALDO_ANTERIOR, CAJA_NO_EXPLICADA))


def procesar_reporte_liberaciones(filas, config, aplicar=False, archivo=''):
    """Crea/actualiza retiros y asocia cada venta liberada a su retiro.

    `filas`: salida de ``leer_csv`` sobre el reporte de Liberaciones de MP
    (columnas DATE, SOURCE_ID, EXTERNAL_REFERENCE, RECORD_TYPE, DESCRIPTION,
    NET_CREDIT_AMOUNT, NET_DEBIT_AMOUNT; también acepta encabezados en
    castellano).

    Regla (FIFO de saldo): el saldo disponible es una cola en orden de
    liberación. Parte del saldo inicial del reporte, reconstruido con las ventas
    previas que siguen sin retiro. Cada retiro (DESCRIPTION 'payout' o
    'withdrawal') se lleva plata desde el frente; una venta queda asociada al
    retiro que se lleva su último peso. Funciona igual con el retiro automático
    (se lleva todo) y con un retiro manual parcial.

    Movimientos que restan y no son retiros (devolución, contracargo, comisión,
    retención): primero se descuentan de la MISMA operación si está en la cola
    (esa venta no llega al banco); el resto sale del frente, y las ventas que
    se agotan así se asocian al retiro siguiente. Si la cola está vacía, la
    diferencia queda como deuda que pagan los créditos siguientes.

    Idempotente y a prueba de reprocesos: al aplicar, las ventas que un retiro de
    este reporte NO se llevó pierden ese vínculo (quedan disponibles en MP).

    - CONCILIADO: el retiro está cubierto por lo liberado.
    - CON_DIFERENCIA: el retiro supera lo que el reporte muestra disponible.
    `aplicar=False` no escribe nada (vista previa).
    """
    registros = []
    saldo_inicial = 0
    for fila in filas:
        tipo = _sin_tildes(_col(fila, 'RECORD_TYPE')).lower()
        credito = abs(_monto(_col(fila, 'NET_CREDIT_AMOUNT')))
        debito = abs(_monto(_col(fila, 'NET_DEBIT_AMOUNT')))
        if tipo in ('initial_available_balance', 'saldo inicial disponible', 'saldo inicial'):
            saldo_inicial += (credito - debito) or _monto(_col(fila, 'GROSS_AMOUNT'))
            continue
        if tipo and tipo not in ('release', 'liberacion', 'liberado', 'liberaciones'):
            continue  # total, available_balance (pre_/pos_payout), block, …
        descripcion = _sin_tildes(_col(fila, 'DESCRIPTION')).lower()
        if descripcion in DESCRIPCIONES_INTERNAS:
            continue
        instante = _instante(_col(fila, 'DATE'))
        if instante is None:
            continue
        registros.append({
            'instante': instante,
            'source_id': str(_col(fila, 'SOURCE_ID')).split('.')[0],
            'external_reference': _col(fila, 'EXTERNAL_REFERENCE'),
            'descripcion': descripcion,
            'credito': credito,
            'pos': (_col(fila, 'EXTERNAL_POS_ID'), _col(fila, 'POS_ID'), _col(fila, 'POS_NAME')),
            'debito': debito,
        })
    registros.sort(key=lambda r: r['instante'])

    ids = [r['source_id'] for r in registros if r['source_id']]
    refs = [r['external_reference'] for r in registros if r['external_reference']]
    locales = {}
    if ids or refs:
        for t in TransaccionMercadoPago.objects.filter(
                Q(payment_id_mp__in=ids) | Q(payment_id__in=ids) | Q(external_reference__in=refs)
        ).select_related('config__sucursal'):
            for clave in (t.payment_id_mp, t.payment_id, t.external_reference):
                if clave:
                    locales[str(clave)] = t

    # Pagos «MP manual» del POS: el N° de operación quedó en el voucher (a veces
    # digitado con espacios o guiones: se compara solo por dígitos).
    manuales = {}
    ids_num = {i for i in ids if i.isdigit()}
    if ids_num and registros:
        margen = timedelta(days=5)
        for pago in (TicketDetallePago.objects
                     .filter(metodo_pago__in=METODOS_MP, creado_en__gte=registros[0]['instante'] - margen - timedelta(days=35),
                             creado_en__lte=registros[-1]['instante'] + margen)
                     .exclude(voucher__isnull=True).exclude(voucher='')
                     .select_related('ticket__sucursal')):
            voucher = _solo_digitos(pago.voucher)
            if voucher in ids_num:
                manuales[voucher] = pago

    # Caja por el punto de venta que informa el reporte (external_pos_id / pos_id).
    cajas_por_pos = {}
    for c in MercadoPagoConfig.objects.filter(id__in=_configs_de_la_cuenta(config)).select_related('sucursal'):
        etiqueta = f'{c.sucursal.alias} · {c.nombre}' if c.sucursal_id else c.nombre
        for clave in (c.external_pos_id, c.pos_id):
            if clave:
                cajas_por_pos[str(clave)] = etiqueta

    cajas_por_nombre = {}
    for c in MercadoPagoConfig.objects.filter(id__in=_configs_de_la_cuenta(config)).select_related('sucursal'):
        etiqueta = f'{c.sucursal.alias} · {c.nombre}' if c.sucursal_id else c.nombre
        for clave in (c.nombre, c.sucursal.alias if c.sucursal_id else ''):
            if clave:
                cajas_por_nombre[_sin_tildes(clave).lower()] = etiqueta

    def _caja_por_nombre(pos_nombre):
        """El POS_NAME del reporte contra el nombre de la caja o el alias de la tienda."""
        nombre = _sin_tildes(pos_nombre or '').lower()
        if not nombre:
            return ''
        if nombre in cajas_por_nombre:
            return cajas_por_nombre[nombre]
        return next((et for clave, et in cajas_por_nombre.items() if clave in nombre), '')

    def _caja_de(item):
        trx = item.get('trx')
        if trx is not None and trx.config_id:
            cfg = trx.config
            return f'{cfg.sucursal.alias} · {cfg.nombre}' if cfg.sucursal_id else cfg.nombre
        manual = item.get('manual')
        if manual is not None:
            return f'{manual.ticket.sucursal.alias} · MP manual'
        ext_pos, pos_id, pos_nombre = item.get('pos') or ('', '', '')
        return (cajas_por_pos.get(ext_pos) or cajas_por_pos.get(pos_id) or _caja_por_nombre(pos_nombre)
                or (f'Caja MP «{pos_nombre}»' if pos_nombre else '')
                or (CAJA_SALDO_ANTERIOR if item.get('id') == 'saldo inicial' else 'Sin caja (online, link de pago u otro)'))

    resultado = {'retiros': [], 'pagos_sin_local': 0, 'pagos_amarrados': 0,
                 'pagos_manuales': 0, 'muestra_sin_local': [], 'dias_sin_local': []}
    dias_sin_local = set()
    dias_manuales = set()
    reembolsos = {}     # source_id -> total devuelto en este reporte
    # Cola FIFO del saldo disponible: {'trx': TransaccionMercadoPago|None, 'resto': $, 'id': str}
    cola = []
    deuda = 0
    pendientes_asociar = []   # ventas agotadas por débitos que no son retiros
    sin_banco = []            # ventas devueltas: su plata no llegó al banco

    # Retiros que trae este reporte (mismo id que se les guarda más abajo), los
    # que ya estaban registrados y las ventas que cada uno ya tiene amarradas.
    ids_retiros = [r['source_id'] or f'RET-{r["instante"]:%Y%m%d%H%M}-{r["debito"]}'
                   for r in registros if r['descripcion'] in DESCRIPCIONES_RETIRO and r['debito']]
    registrados = ({x.withdrawal_id: x for x in RetiroMercadoPago.objects.filter(withdrawal_id__in=ids_retiros)}
                   if ids_retiros else {})
    ya_registrados = set(registrados)
    # «Cerrado»: su desglose guardado explica todo el monto. Esos no se recalculan
    # al reprocesar; los que quedaron con algo sin explicar sí (FIFO de nuevo).
    cerrados = {wid for wid, x in registrados.items()
                if isinstance(x.raw_reporte, dict) and x.raw_reporte.get('por_caja')
                and _no_explicado(x.raw_reporte['por_caja']) == 0
                and not _desglose_envenenado(wid, x.raw_reporte)}
    # Retiros guardados antes de existir el desglose: cerrados si las ventas que
    # ya tienen amarradas cubren su monto (si no, al reprocesar se desamarraban).
    legado = [wid for wid, x in registrados.items()
              if not (isinstance(x.raw_reporte, dict) and x.raw_reporte.get('por_caja'))]
    if legado:
        sumas = dict(TransaccionMercadoPago.objects.filter(retiro__withdrawal_id__in=legado)
                     .values('retiro__withdrawal_id')
                     .annotate(s=Sum(Coalesce('monto_neto', 'monto')))
                     .values_list('retiro__withdrawal_id', 's'))
        cerrados |= {wid for wid in legado if (sumas.get(wid) or 0) >= registrados[wid].monto > 0}
    propias_de = {}
    if cerrados:
        for trx_id, wid in (TransaccionMercadoPago.objects.filter(retiro__withdrawal_id__in=list(cerrados))
                            .values_list('id', 'retiro__withdrawal_id')):
            propias_de.setdefault(wid, set()).add(trx_id)
    # El piso lo fijan los retiros cerrados anteriores al primer retiro abierto
    # del reporte: uno cerrado posterior (p. ej. el nocturno) no puede impedir
    # que un retiro abierto anterior rearme su saldo con ventas viejas.
    orden_retiros = [(r['instante'], r['source_id'] or f'RET-{r["instante"]:%Y%m%d%H%M}-{r["debito"]}')
                     for r in registros if r['descripcion'] in DESCRIPCIONES_RETIRO and r['debito']]
    primer_abierto = min((inst for inst, wid in orden_retiros if wid not in cerrados), default=None)
    cerrados_previos = {wid for inst, wid in orden_retiros
                        if wid in cerrados and (primer_abierto is None or inst < primer_abierto)}
    ids_piso = set().union(*(propias_de.get(w, set()) for w in cerrados_previos)) if cerrados_previos else set()
    piso = (TransaccionMercadoPago.objects
            .filter(id__in=ids_piso, money_release_date__isnull=False)
            .aggregate(m=Min('money_release_date'))['m'] if ids_piso else None)
    # Ventas que el reporte trae como liberación (no las que solo aparecen en
    # una devolución: esas pueden ser de un retiro y deben poder rearmar el saldo).
    ids_credito = set()
    for r in registros:
        if r['credito'] and r['descripcion'] in DESCRIPCIONES_PAGO:
            t = locales.get(r['source_id']) or locales.get(r['external_reference'])
            if t is not None:
                ids_credito.add(t.id)

    if saldo_inicial < 0:
        deuda = -saldo_inicial
    elif saldo_inicial > 0:
        previas = _remanente_previo(
            config, registros[0]['instante'] if registros else None, saldo_inicial,
            excluir_ids=ids_credito, retiros_del_reporte=ids_retiros, piso=piso,
        )
        anonimo = saldo_inicial - sum(n for _t, n in previas)
        if anonimo > 0:
            cola.append({'trx': None, 'resto': anonimo, 'id': 'saldo inicial'})
        for trx, neto in previas:
            cola.append({'trx': trx, 'resto': neto, 'id': trx.payment_id_mp or trx.external_reference})
    resultado['remanente_previo'] = sum(it['resto'] for it in cola if it['trx'] is not None)

    tomado = []   # (item, monto) que se lleva el retiro en curso

    def _consumir(monto):
        """Saca `monto` del frente de la cola. Devuelve (faltante, ítems agotados)."""
        restante, agotados = monto, []
        while restante > 0 and cola:
            item = cola[0]
            usa = min(item['resto'], restante)
            item['resto'] -= usa
            restante -= usa
            tomado.append((item, usa))
            if item['resto'] <= 0:
                agotados.append(cola.pop(0))
        return restante, agotados

    retiros_ids = []   # RetiroMercadoPago escritos en esta pasada (solo con aplicar)
    liberadas = {}     # trx.id -> (trx, instante de la liberación) de las ventas del POS
    devueltas_fuera = []   # devueltas enteras cuya liberación quedó antes del reporte
    ids_retiros_vistos = set()   # retiros ya procesados en este reporte (para detectar un reverso)
    devoluciones_reporte = []    # (trx, monto, instante) devoluciones que muestra el reporte
    for r in registros:
        desc = r['descripcion']
        if desc in DESCRIPCIONES_RETIRO and r['debito']:
            monto = r['debito']
            withdrawal_id = r['source_id'] or f'RET-{r["instante"]:%Y%m%d%H%M}-{monto}'
            ids_retiros_vistos.add(withdrawal_id)
            tomado.clear()
            # Un retiro cerrado se lleva primero SUS ventas: al reprocesar no se le
            # cambian por otras ni se le pasa una suya al retiro siguiente.
            propias, sacado = [], 0
            ids_propias = propias_de.get(withdrawal_id) or set()
            for it in list(cola):
                if sacado >= monto:
                    break
                if it['trx'] is not None and it['trx'].id in ids_propias:
                    usa = min(it['resto'], monto - sacado)
                    it['resto'] -= usa
                    sacado += usa
                    tomado.append((it, usa))
                    if it['resto'] <= 0:
                        cola.remove(it)
                        propias.append(it)
            faltante, agotados = _consumir(monto - sacado)
            agotados = propias + agotados
            por_caja = {}
            for item, usado in tomado:
                por_caja[_caja_de(item)] = por_caja.get(_caja_de(item), 0) + usado
            if faltante:
                por_caja[CAJA_NO_EXPLICADA] = faltante
            trxs = [it['trx'] for it in pendientes_asociar + agotados if it['trx'] is not None]
            pendientes_asociar = []
            quedan = sum(it['resto'] for it in cola)
            estado = 'CONCILIADO' if faltante == 0 else 'CON_DIFERENCIA'
            if faltante:
                detalle = (
                    f'Hay una diferencia de ${faltante:,}: el retiro de ${monto:,} supera lo liberado '
                    f'que muestra el reporte. Pida un reporte que empiece antes (saldo previo) '
                    'o revise movimientos que no pasan por el POS.'
                ).replace(',', '.')
            else:
                detalle = (f'Retiro parcial: quedaron ${quedan:,} disponibles en Mercado Pago.'
                           .replace(',', '.') if quedan else '')
            por_caja_lista = sorted(({'caja': k, 'monto': v} for k, v in por_caja.items()),
                                    key=lambda x: -x['monto'])
            # Retiro ya registrado: si su desglose guardado lo explica todo, o mejor
            # que esta pasada (este reporte parte más tarde), se conserva tal cual,
            # con sus vínculos. Solo se recalcula lo que quedó sin explicar.
            previo = registrados.get(withdrawal_id)
            raw_prev = previo.raw_reporte if previo is not None and isinstance(previo.raw_reporte, dict) else {}
            guardado = raw_prev.get('por_caja')
            conservar = (bool(guardado) and not _desglose_envenenado(withdrawal_id, raw_prev)
                         and (_no_explicado(guardado) == 0
                              or _no_explicado(guardado) < _no_explicado(por_caja_lista)))
            if conservar:
                trxs = []
                por_caja_lista = guardado
                estado = previo.estado
                detalle = previo.detalle_diferencia or detalle
            resultado['retiros'].append({
                'withdrawal_id': withdrawal_id,
                'fecha': str(timezone.localtime(r['instante']).date()),
                'hora': timezone.localtime(r['instante']).strftime('%H:%M'),
                'monto': monto,
                'pagos': len(agotados),
                'pagos_pos': raw_prev.get('pagos_pos', len(ids_propias)) if conservar else len(trxs),
                'suma_pagos': monto - faltante,
                'por_caja': por_caja_lista,
                'quedan_disponibles': quedan,
                'estado': estado,
                'detalle': detalle,
                'nuevo': withdrawal_id not in ya_registrados,
            })
            if aplicar:
                with transaction.atomic():
                    retiro = (RetiroMercadoPago.objects.select_for_update()
                              .filter(withdrawal_id=withdrawal_id).first())
                    raw = dict(retiro.raw_reporte or {}) if retiro and isinstance(retiro.raw_reporte, dict) else {}
                    archivos = set(raw.get('archivos') or [])
                    if raw.get('archivo'):
                        archivos.add(raw['archivo'])
                    if archivo:
                        archivos.add(archivo)
                    if retiro is not None and conservar:
                        raw['archivos'] = sorted(archivos)
                        raw.pop('archivo', None)
                        if archivo and _no_explicado(por_caja_lista):
                            raw['reintento_en'] = timezone.now().isoformat()
                        retiro.raw_reporte = raw
                        retiro.save(update_fields=['raw_reporte', 'actualizado_en'])
                        # Sus vínculos quedan como estaban; solo se amarran los cobros
                        # que ya había contado (por su N°) y que se registraron después,
                        # p. ej. un pago «MP manual» importado.
                        prev_ids = ({str(x) for x in raw.get('pagos') or []}
                                    | {str(x) for x in (raw.get('netos') or {})})
                        nuevas = [it['trx'].id for it, _u in tomado
                                  if it['trx'] is not None and it['trx'].retiro_id is None
                                  and str(it['id']) in prev_ids]
                        if nuevas:
                            TransaccionMercadoPago.objects.filter(
                                id__in=nuevas, retiro__isnull=True).update(retiro=retiro)
                        continue
                    raw_nuevo = {'archivos': sorted(archivos),
                                 'pagos': [it['id'] for it in agotados][:500],
                                 'instante': r['instante'].isoformat(),
                                 'por_caja': por_caja_lista,
                                 'pagos_pos': len(trxs),
                                 # lo que quedó disponible en MP: 0 = se llevó todo; > 0 = retiro por un monto
                                 'quedan': quedan,
                                 # neto liberado de cada pago «MP manual» que este retiro terminó
                                 # de llevarse (el pago entero, no solo el trozo de este retiro)
                                 'netos': {str(it['id']): it.get('total', usado) for it, usado in tomado[:2000]
                                           if it['trx'] is None and it.get('manual') is not None
                                           and it['resto'] <= 0},
                                 # cuánto se llevó de cada venta (una venta partida entre dos retiros)
                                 'partes': {str(it['id']): usado for it, usado in tomado[:2000] if it['trx'] is not None},
                                 # lo que no se pudo explicar con una venta del sistema
                                 'sin_venta': faltante + sum(usado for it, usado in tomado
                                                             if it['trx'] is None and it.get('manual') is None),
                                 # y por operación, para descontarlo si después aparece su cobro
                                 'sin_local': {str(it['id']): usado for it, usado in tomado[:2000]
                                               if it['trx'] is None and it.get('manual') is None
                                               and it['id'] and it['id'] != 'saldo inicial'}}
                    # Solo una pasada completa (con archivo) gasta el reintento de 6 h: una
                    # incompleta (faltó tiempo) se rehace en la vuelta siguiente.
                    if archivo and (_no_explicado(por_caja_lista) or _desglose_envenenado(withdrawal_id, raw_nuevo)):
                        raw_nuevo['reintento_en'] = timezone.now().isoformat()
                    datos = {'config': config, 'fecha': timezone.localtime(r['instante']).date(),
                             'monto': monto, 'estado': estado, 'detalle_diferencia': detalle,
                             'raw_reporte': raw_nuevo}
                    if retiro is None:
                        try:
                            with transaction.atomic():
                                retiro = RetiroMercadoPago.objects.create(withdrawal_id=withdrawal_id, **datos)
                        except IntegrityError:   # otra pasada lo creó en este instante
                            retiro = RetiroMercadoPago.objects.select_for_update().get(withdrawal_id=withdrawal_id)
                            for campo, valor in datos.items():
                                setattr(retiro, campo, valor)
                            retiro.save()
                    else:
                        for campo, valor in datos.items():
                            setattr(retiro, campo, valor)
                        retiro.save()
                    retiros_ids.append(retiro.id)
                    if trxs:
                        TransaccionMercadoPago.objects.filter(
                            id__in=[t.id for t in trxs]).update(retiro=retiro)
            resultado['pagos_amarrados'] += len(trxs)
            continue

        neto = r['credito'] - r['debito']
        if not neto:
            continue
        if neto < 0:
            debito = -neto
            if r['source_id']:
                reembolsos[r['source_id']] = reembolsos.get(r['source_id'], 0) + debito
            # 1) La misma operación (devolución de esa venta): no llega al banco.
            en_cola = False
            for it in list(cola):
                if r['source_id'] and it['id'] == r['source_id'] and it['trx'] is not None:
                    en_cola = True
                    usa = min(it['resto'], debito)
                    it['resto'] -= usa
                    devoluciones_reporte.append((it['trx'], usa, r['instante']))
                    debito -= usa
                    if it['resto'] <= 0:
                        cola.remove(it)
                        sin_banco.append(it['trx'])
                    break
            # Devuelta entera una venta liberada antes de este reporte y que no se
            # llevó ningún retiro: su plata ya no está en MP. Se le borra la fecha de
            # liberación para que ningún reporte posterior la dé por depositada.
            t_dev = locales.get(r['source_id']) if (r['source_id'] and not en_cola) else None
            if t_dev is not None:
                devoluciones_reporte.append((t_dev, -neto, r['instante']))
            if (t_dev is not None and t_dev.retiro_id is None
                    and -neto >= int(t_dev.monto_neto if t_dev.monto_neto is not None else t_dev.monto)):
                devueltas_fuera.append(t_dev)
            # 2) El resto sale del frente; esas ventas van con el retiro siguiente.
            if debito:
                faltante, agotados = _consumir(debito)
                pendientes_asociar.extend(agotados)
                deuda += faltante
            continue
        # Crédito con el N° de un retiro (o que dice reverso/rechazo): el banco
        # devolvió el retiro y la plata volvió a MP. Ese retiro no llegó al
        # banco: se desamarran sus ventas y vuelven a la cola para el siguiente.
        if r['source_id'] and (r['source_id'] in ids_retiros_vistos or r['source_id'] in registrados
                               or any(k in desc for k in ('revers', 'rechaz', 'reject'))):
            revertido = registrados.get(r['source_id']) or RetiroMercadoPago.objects.filter(
                withdrawal_id=r['source_id']).first()
            resto_credito = neto
            if revertido is not None:
                for t in TransaccionMercadoPago.objects.filter(retiro=revertido).select_related('config__sucursal'):
                    neto_t = int(t.monto_neto if t.monto_neto is not None else t.monto)
                    usa = min(neto_t, resto_credito)
                    if usa <= 0:
                        break
                    cola.append({'trx': t, 'resto': usa, 'total': usa, 'id': t.payment_id_mp or t.external_reference,
                                 'pos': ('', '', ''), 'manual': None})
                    resto_credito -= usa
            if resto_credito > 0:
                cola.append({'trx': None, 'resto': resto_credito, 'total': resto_credito, 'id': r['source_id'],
                             'pos': ('', '', ''), 'manual': None})
            for fila_ret in resultado['retiros']:
                if fila_ret['withdrawal_id'] == r['source_id']:
                    fila_ret['estado'] = 'REVERTIDO'
                    fila_ret['detalle'] = 'Retiro revertido: la plata volvió a Mercado Pago.'
            resultado['revertidos'] = resultado.get('revertidos', 0) + 1
            if aplicar and revertido is not None:
                TransaccionMercadoPago.objects.filter(retiro=revertido).update(retiro=None)
                raw_rev = dict(revertido.raw_reporte or {}) if isinstance(revertido.raw_reporte, dict) else {}
                raw_rev['revertido'] = {'instante': r['instante'].isoformat(), 'monto': neto}
                revertido.raw_reporte = raw_rev
                revertido.estado = 'REVERTIDO'
                revertido.detalle_diferencia = (f'Retiro revertido el {timezone.localtime(r["instante"]):%d/%m %H:%M}: '
                                                'la plata volvió a Mercado Pago (el banco lo rechazó).')
                revertido.save(update_fields=['raw_reporte', 'estado', 'detalle_diferencia', 'actualizado_en'])
            continue
        # Crédito: primero paga la deuda (saldo negativo), después entra a la cola.
        es_pago = desc in DESCRIPCIONES_PAGO
        trx = (locales.get(r['source_id']) or locales.get(r['external_reference'])) if es_pago else None
        if trx is not None:
            liberadas.setdefault(trx.id, (trx, r['instante']))
        manual = manuales.get(r['source_id']) if (es_pago and trx is None) else None
        if manual is not None:
            resultado['pagos_manuales'] += 1
            # Días a leer para registrar este pago como cobro (el cobro en la
            # máquina fue minutos antes de registrarlo: si fue tras la medianoche, el anterior).
            local = timezone.localtime(manual.creado_en)
            dias_manuales.add(local.date())
            if (local.hour, local.minute) < (0, 30):
                dias_manuales.add(local.date() - timedelta(days=1))
        elif es_pago and trx is None:
            resultado['pagos_sin_local'] += 1
            dias_sin_local.add(timezone.localtime(r['instante']).date())
            if len(resultado['muestra_sin_local']) < 5:
                resultado['muestra_sin_local'].append({
                    'fecha': timezone.localtime(r['instante']).strftime('%Y-%m-%d %H:%M'),
                    'source_id': r['source_id'], 'external_reference': r['external_reference'],
                    'monto': neto, 'pos': ' / '.join(x for x in r['pos'] if x),
                })
        total = neto   # neto del pago entero (antes de pagar una deuda de saldo)
        if deuda:
            usa = min(deuda, neto)
            deuda -= usa
            neto -= usa
            if not neto:
                continue
        cola.append({'trx': trx, 'resto': neto, 'total': total, 'id': r['source_id'], 'pos': r['pos'],
                     'manual': manual})

    resultado['liberado_sin_retirar'] = sum(it['resto'] for it in cola)
    resultado['dias_sin_local'] = sorted(str(d) for d in dias_sin_local)
    resultado['dias_manuales'] = sorted(str(d) for d in dias_manuales)
    resultado['deuda'] = deuda
    resultado['rango'] = ([registros[0]['instante'].isoformat(), registros[-1]['instante'].isoformat()]
                          if registros else [])
    resultado['filas'] = len(filas)
    resultado['registros'] = len(registros)

    # La fila «release» ES la liberación: se guarda en la venta si no la tenía.
    # Sin esa fecha `_remanente_previo` no veía lo que dejó un retiro parcial y
    # esas ventas quedaban para siempre como «liberado sin retirar».
    for t in sin_banco + devueltas_fuera:
        liberadas.pop(t.id, None)
    if aplicar and (sin_banco or devueltas_fuera):
        ids_dev = [t.id for t in sin_banco + devueltas_fuera]
        TransaccionMercadoPago.objects.filter(id__in=ids_dev, retiro__isnull=True).update(money_release_date=None)
        # Un pago «MP manual» importado no recibe avisos de MP por referencia: si el
        # reporte muestra devuelto su neto entero, se marca DEVUELTA acá. Una
        # devolución parcial no: la venta sigue vigente por el resto.
        enteras = [t.id for t in sin_banco + devueltas_fuera
                   if reembolsos.get(t.payment_id_mp or t.payment_id or '', 0)
                   >= int(t.monto_neto if t.monto_neto is not None else t.monto)]
        TransaccionMercadoPago.objects.filter(
            id__in=enteras, estado='APROBADA',
        ).update(estado='DEVUELTA', estado_detalle='Devuelta en Mercado Pago (reporte de Liberaciones)')
    if aplicar and devoluciones_reporte:
        _registrar_devoluciones_del_reporte(devoluciones_reporte)
    if aplicar and liberadas:
        sin_fecha = []
        for trx, instante in liberadas.values():
            if trx.money_release_date is None:
                trx.money_release_date = instante
                sin_fecha.append(trx)
        if sin_fecha:
            TransaccionMercadoPago.objects.bulk_update(sin_fecha, ['money_release_date'], batch_size=200)

    # Reproceso: lo que ningún retiro de ESTE reporte se llevó no puede quedar
    # asociado a esos retiros (quedaría "depositado" estando en MP).
    if aplicar and retiros_ids:
        sueltas = [it['trx'].id for it in cola + pendientes_asociar if it['trx'] is not None]
        sueltas += [t.id for t in sin_banco]
        if sueltas:
            resultado['desasociadas'] = (TransaccionMercadoPago.objects
                                         .filter(id__in=sueltas, retiro_id__in=retiros_ids)
                                         .update(retiro=None))
    return resultado


def _vouchers_mp_manual(configs, desde, hasta):
    """{N° de operación: TicketDetallePago} de los pagos «MP manual» de ventas
    cerradas (ticket PAGADO) de las tiendas de la cuenta, sin transacción
    todavía y con un N° con forma válida."""
    sucursales = set(MercadoPagoConfig.objects.filter(id__in=configs).values_list('sucursal_id', flat=True))
    salida = {}
    for pago in (TicketDetallePago.objects
                 .filter(Q(origen_pago='MANUAL') | Q(origen_pago__isnull=True),
                         metodo_pago__in=METODOS_MP, ticket__sucursal_id__in=sucursales,
                         ticket__estado='PAGADO', transacciones_mercadopago__isnull=True,
                         creado_en__date__gte=desde, creado_en__date__lte=hasta)
                 .select_related('ticket').order_by('creado_en')):
        voucher = _solo_digitos(pago.voucher)
        if len(voucher) >= 9:
            salida[voucher] = pago
    # Un N° digitado en más de un pago (sin transacción, en cualquier fecha) es
    # ambiguo: no se elige uno, queda para «Asociar».
    if salida:
        conteo = {}
        for voucher in (TicketDetallePago.objects
                        .filter(metodo_pago__in=METODOS_MP, transacciones_mercadopago__isnull=True,
                                creado_en__date__gte=desde - timedelta(days=60))
                        .exclude(ticket__estado='ANULADO').exclude(voucher__isnull=True)
                        .values_list('voucher', flat=True)):
            v = _solo_digitos(voucher)
            if v in salida:
                conteo[v] = conteo.get(v, 0) + 1
        for voucher, n in conteo.items():
            if n > 1:
                salida.pop(voucher, None)
    return salida


def _caja_del_pago(payment, cajas):
    """Caja (id de MercadoPagoConfig) de la tienda donde se cobró el pago: la única
    si hay una; si no, la que tiene el punto de venta que informa MP; si no, la
    habilitada y principal. Nunca «la de menor id» a ciegas: el cierre de cada
    caja aprende de estos cobros qué pos_id/store_id le pertenecen."""
    if len(cajas) == 1:
        return cajas[0].id
    claves = {str(payment.get(k) or '') for k in ('pos_id', 'store_id')} - {''}
    for c in cajas:
        if claves & ({c.pos_id, c.external_pos_id} - {''}):
            return c.id
    return cajas[0].id if cajas else None


def _retiros_por_operacion(configs, dias=180):
    """{N° de operación: RetiroMercadoPago} de los retiros recientes de la cuenta
    que ya contaron ese pago (raw_reporte 'pagos'/'netos'), para amarrar al
    registrarlo un pago «MP manual» que un retiro ya se llevó."""
    salida = {}
    desde = timezone.localdate() - timedelta(days=dias)
    qs = RetiroMercadoPago.objects.filter(fecha__gte=desde).exclude(estado='REVERTIDO').order_by('fecha', 'id')
    if configs is not None:
        qs = qs.filter(config_id__in=list(configs))
    for retiro in qs:
        raw = retiro.raw_reporte if isinstance(retiro.raw_reporte, dict) else {}
        for clave in list(raw.get('pagos') or []) + list((raw.get('netos') or {}).keys()):
            salida[str(clave)] = retiro
    return salida


def _importar_pago_manual(pago, payment, config_id, retiro=None):
    """Transacción (consumida) para un pago «MP manual» con los datos de MP.

    No se importa (queda para «Asociar») si: el pago no está aprobado, el monto
    no calza, ya tuvo devoluciones, o MP lo trae con la referencia de un cobro
    que el sistema ya tiene (cobro del POS o «Tu caja» que todavía no tiene su
    N°: importarlo lo duplicaría).
    """
    pid = str(payment.get('id') or '')
    monto = _monto(payment.get('transaction_amount'))
    if payment.get('status') != 'approved' or monto != int(pago.monto or 0):
        return None
    if _monto(payment.get('transaction_amount_refunded') or 0) > 0 or payment.get('refunds'):
        return None
    ya_existe = Q(payment_id_mp=pid) | Q(payment_id=pid) | Q(external_reference=f'MANUAL-{pid}')
    referencia = str(payment.get('external_reference') or '').strip()
    if referencia:
        ya_existe |= Q(external_reference=referencia)
    if TransaccionMercadoPago.objects.filter(ya_existe).exists():
        return None
    neto = (payment.get('transaction_details') or {}).get('net_received_amount')
    neto = _monto(neto) if neto is not None else None
    tipo_pago = str(payment.get('payment_type_id') or '')
    try:
        with transaction.atomic():
            # Lock del pago (lo toman también «Asociar» e importar_y_asociar) y
            # revisión de nuevo: otra pasada o un Maestro pudo haberlo resuelto.
            bloqueado = (TicketDetallePago.objects.select_for_update()
                         .filter(pk=pago.pk, ticket__estado='PAGADO').first())
            # Revalidar con el pago bloqueado (un Maestro pudo corregirlo mientras
            # corría la pasada): mismo N°, mismo monto y sigue siendo MP manual.
            if (bloqueado is None or bloqueado.metodo_pago not in METODOS_MP
                    or bloqueado.origen_pago not in ('MANUAL', None)
                    or _solo_digitos(bloqueado.voucher) != pid or int(bloqueado.monto or 0) != monto
                    or TransaccionMercadoPago.objects.filter(ya_existe | Q(detalle_pago_id=pago.pk)).exists()):
                return None
            pago = bloqueado
            trx = TransaccionMercadoPago.objects.create(
                config_id=config_id, sucursal_id=pago.ticket.sucursal_id, ticket=pago.ticket,
                detalle_pago=pago, correlativo_ticket=str(pago.ticket.correlativo), tipo='VENTA',
                canal='POINT' if tipo_pago in ('debit_card', 'credit_card', 'prepaid_card') else 'QR',
                external_reference=f'MANUAL-{pid}', payment_id=pid, payment_id_mp=pid,
                monto=monto, monto_neto=neto, fee_mp=(monto - neto) if neto is not None else None,
                installments=int(payment.get('installments') or 1), estado='APROBADA',
                estado_detalle=f'Pago MP manual del POS, datos traídos de MP ({payment.get("status_detail") or ""})'[:120],
                metodo_pago_mp=tipo_pago[:40],
                ultimos_4_digitos=str((payment.get('card') or {}).get('last_four_digits') or '')[:4],
                codigo_autorizacion=str(payment.get('authorization_code') or '')[:30],
                money_release_date=_instante(payment.get('money_release_date')),
                consumida=True, raw_response=payment, retiro=retiro,
            )
            # Fecha del cobro real (creado_en es auto_now_add: quedaría con la de hoy).
            TransaccionMercadoPago.objects.filter(pk=trx.pk).update(
                creado_en=_instante(payment.get('date_created')) or pago.creado_en)
    except IntegrityError:
        return None   # otra pasada lo registró al mismo tiempo
    logger.info("Conciliación MP: pago MP manual N° %s ($%s) del ticket #%s registrado como cobro",
                pid, monto, pago.ticket.correlativo)
    return trx


def _registrar_devoluciones_del_reporte(devoluciones):
    """Deja registrada (tipo DEVOLUCION) cada devolución que el reporte muestra
    sobre una venta local, si ninguna otra fila la registró ya (el aviso de MP
    o una NC). Sin esto, sin webhook, la venta seguía «en el banco» entera."""
    for trx, monto, instante in devoluciones:
        monto = int(monto)
        if monto <= 0:
            continue
        clave = f'{trx.external_reference}-REF-REP-{instante:%Y%m%d%H%M%S}'[:80]
        ya = TransaccionMercadoPago.objects.filter(
            tipo='DEVOLUCION', transaccion_origen=trx, monto=monto,
            creado_en__gte=instante - timedelta(days=3), creado_en__lte=instante + timedelta(days=3))
        if ya.exists() or TransaccionMercadoPago.objects.filter(external_reference=clave).exists():
            continue
        try:
            with transaction.atomic():
                dev = TransaccionMercadoPago.objects.create(
                    config_id=trx.config_id, sucursal_id=trx.sucursal_id, ticket_id=trx.ticket_id,
                    correlativo_ticket=trx.correlativo_ticket, tipo='DEVOLUCION', canal=trx.canal,
                    transaccion_origen=trx, external_reference=clave, monto=monto, estado='DEVUELTA',
                    estado_detalle='Devolución vista en el reporte de Liberaciones de MP',
                    metodo_pago_mp=trx.metodo_pago_mp, consumida=True,
                )
                TransaccionMercadoPago.objects.filter(pk=dev.pk).update(creado_en=instante)
        except IntegrityError:
            continue


def completar_numeros_mp(config, dias, presupuesto_seg=35, importar=False, saltar=()):
    """Rellena `payment_id_mp` (y neto/comisión si faltan) de los cobros de la
    cuenta leyendo `payments/search` de esos días, cruzado por external_reference.

    El reporte de Liberaciones identifica cada pago por ese número; el aviso de
    pagos (webhook) casi nunca lo dejó guardado, y sin él ninguna venta del POS
    se cruzaba con su retiro. También guarda la fecha de liberación (solo pagos
    aprobados), aunque el cobro ya tuviera el N°: sin ella el saldo con que
    parte un reporte no se puede explicar con ventas y un retiro quedaba como
    «Saldo anterior». Solo escribe campos vacíos. Devuelve
    {'dias': n, 'completados': n (N° nuevos), 'actualizados': n (filas tocadas),
    'importados': n (pagos «MP manual» registrados como cobro), 'sin_tiempo': bool,
    'fallidos': n}. Con `importar=True` (solo al APLICAR; nunca en una vista
    previa), los pagos «MP manual» cuyo N° aparece ese día en MP se registran
    como cobro (`_importar_pago_manual`).
    `fallidos` son los días que no se pudieron leer por un error de red o de
    MP saturado (429/5xx): quien llama debe tratarlo como «no terminó».
    `saltar`: días que otra vuelta ya leyó (la página los devuelve: la caché es
    por worker). `leidos` en la respuesta: los días que esta vuelta dio por leídos.
    """
    import time as _time
    fin = _time.monotonic() + presupuesto_seg
    completados, leidos, fallidos, seguidos, actualizados, importados = 0, 0, 0, 0, 0, 0
    lista_leidos = []
    configs = _configs_de_la_cuenta(config)
    hoy = timezone.localdate()
    dias = sorted(set(dias) - set(saltar))
    manuales = (_vouchers_mp_manual(configs, dias[0] - timedelta(days=1), dias[-1] + timedelta(days=1))
                if dias and importar else {})
    cajas_de_tienda, retiro_de = {}, {}
    if manuales:
        # Todas las cajas de esas tiendas (el pago trae su punto de venta).
        tiendas = MercadoPagoConfig.objects.filter(id__in=configs).values_list('sucursal_id', flat=True)
        for c in (MercadoPagoConfig.objects.filter(sucursal_id__in=list(tiendas))
                  .order_by('-habilitado', '-es_principal', 'id')):
            cajas_de_tienda.setdefault(c.sucursal_id, []).append(c)
        retiro_de = _retiros_por_operacion(configs)
    for dia in dias:
        # Un día pasado ya leído hace poco no se vuelve a pedir: lo que no cruzó
        # entonces (pagos online, devueltos) tampoco va a cruzar ahora.
        # La marca es de una pasada que aplica (con importación): una vista previa
        # no la usa ni la deja, para que el «Aplicar» siguiente sí importe.
        clave_dia = f'conc_mp:dia_leido:{min(configs)}:{dia}'
        if importar and dia < hoy and cache.get(clave_dia):
            lista_leidos.append(str(dia))
            continue
        if _time.monotonic() > fin - RESERVA_DIA_SEG:
            return {'dias': leidos, 'completados': completados, 'sin_tiempo': True, 'fallidos': fallidos,
                    'actualizados': actualizados, 'importados': importados, 'leidos': lista_leidos}
        try:
            pagos = mp.buscar_pagos_dia(config, dia)
        except mp.MercadoPagoError as e:
            logger.warning("Conciliación MP: no se pudo leer payments/search %s: %s", dia, e.mensaje)
            if getattr(e, 'red', False):
                fallidos += 1
                seguidos += 1
                if seguidos >= 2:   # breaker abierto o MP caído: no insistir día por día
                    break
            continue
        seguidos = 0
        leidos += 1
        lista_leidos.append(str(dia))
        if importar and dia < hoy:
            cache.set(clave_dia, 1, 60 * 30)
        por_ref = {str(p.get('external_reference') or ''): p for p in pagos
                   if p.get('external_reference') and str(p.get('id') or '').isdigit()}
        for trx in TransaccionMercadoPago.objects.filter(
                Q(payment_id_mp='') | Q(money_release_date__isnull=True),
                config_id__in=configs, external_reference__in=list(por_ref)):
            pago = por_ref[trx.external_reference]
            if trx.payment_id_mp and trx.payment_id_mp != str(pago['id']):
                continue   # es otro pago (reintento): no se mezclan sus datos
            campos = []
            if not trx.payment_id_mp:
                trx.payment_id_mp = str(pago['id'])[:40]
                campos.append('payment_id_mp')
            neto = (pago.get('transaction_details') or {}).get('net_received_amount')
            if trx.monto_neto is None and neto is not None:
                trx.monto_neto = _monto(neto)
                trx.fee_mp = trx.monto - trx.monto_neto
                campos += ['monto_neto', 'fee_mp']
            liberacion = _instante(pago.get('money_release_date'))
            if (trx.money_release_date is None and liberacion is not None
                    and pago.get('status') == 'approved'):
                trx.money_release_date = liberacion
                campos.append('money_release_date')
            if not campos:
                continue
            trx.save(update_fields=campos + ['actualizado_en'])
            actualizados += 1
            if 'payment_id_mp' in campos:
                completados += 1
        # Después de completar los N° por referencia: así un cobro del POS que
        # recién recibe su N° no se vuelve a crear como pago manual.
        for p in pagos:
            pid = str(p.get('id') or '')
            pago = manuales.pop(pid, None)
            if pago is None:
                continue
            caja = _caja_del_pago(p, cajas_de_tienda.get(pago.ticket.sucursal_id) or []) or config.id
            if _importar_pago_manual(pago, p, caja, retiro=retiro_de.get(pid)) is not None:
                importados += 1
    return {'dias': leidos, 'completados': completados, 'sin_tiempo': False, 'fallidos': fallidos,
            'actualizados': actualizados, 'importados': importados, 'leidos': lista_leidos}


def dias_cobros_sin_numero(config, resultado, dias_antes=60):
    """Días de cobro a consultar en payments/search para cruzar un reporte.

    Son los días de las ventas de la cuenta que siguen sin N° de operación (o
    sin fecha de liberación) y sin retiro, desde `dias_antes` días antes del reporte (la liberación puede
    llegar semanas después del cobro) hasta su fin. Antes se miraban los pagos
    del reporte que no cruzaban: esos incluyen los online y de link de pago,
    que nunca cruzan, y el tope de 15 días dejaba fuera justo los días viejos,
    que son los que un retiro se lleva primero.
    """
    rango = resultado.get('rango') or []
    inicio = _instante(rango[0]) if rango else None
    fin = _instante(rango[-1]) if rango else None
    if inicio is None or fin is None:
        return []
    creados = (TransaccionMercadoPago.objects
               .filter(Q(estado='APROBADA', money_release_date__isnull=True)
                       | Q(estado__in=('APROBADA', 'DEVUELTA'), payment_id_mp='')
                       # aviso de pago perdido: MP cobró y acá quedó pendiente
                       | (Q(estado__in=('CREADA', 'PENDIENTE', 'EXPIRADA', 'ERROR'), payment_id_mp='')
                          & ~Q(order_id='')),
                       config_id__in=_configs_de_la_cuenta(config), tipo='VENTA', retiro__isnull=True,
                       creado_en__gte=inicio - timedelta(days=dias_antes), creado_en__lte=fin)
               .values_list('creado_en', flat=True))
    dias = set()
    for creado in creados:
        local = timezone.localtime(creado)
        dias.add(local.date())
        if (local.hour, local.minute) >= (23, 30):   # el pago puede quedar con fecha del día siguiente
            dias.add(local.date() + timedelta(days=1))
    return sorted(dias)


def dias_para_completar(config, resultado, importar=False):
    """Días a leer en payments/search antes de aplicar un reporte: los que
    tienen cobros propios por completar (los pagos online del reporte nunca
    cruzan y releerlos gastaba el tiempo de cada pasada) y, solo al aplicar,
    los de los pagos «MP manual» que trae el reporte (para registrarlos)."""
    dias = set(dias_cobros_sin_numero(config, resultado))
    if importar:
        dias.update(f for f in (_fecha_libre(x) for x in resultado.get('dias_manuales') or []) if f)
    return sorted(dias)


def dias_a_completar(resultado, margen=2, maximo=15):
    """Días de pago a consultar para los pagos del reporte que no se cruzaron
    (fecha de liberación y hasta `margen` días antes)."""
    dias = set()
    for texto in resultado.get('dias_sin_local') or []:
        base = _fecha_libre(texto)
        if base:
            dias.update(base - timedelta(days=k) for k in range(margen + 1))
    return sorted(dias)[-maximo:]


# ─────────────── API de reportes de Liberaciones (release_report) ───────────────
# Docs oficiales (mercadopago.cl/developers, verificadas 23-09-2026):
#   POST /v1/account/release_report          crea (202 = aceptado; 203 = NO se creó)
#   GET  /v1/account/release_report/task/{id} estado: pending|processing|processed|failed
#   GET  /v1/account/release_report/search    reportes generados (más reciente primero)
#   GET  /v1/account/release_report/{file}    descarga el CSV
#   GET/POST/PUT /v1/account/release_report/config   configuración (execute_after_withdrawal)
# La generación tarda "unos minutos": NUNCA se espera dentro del request.

_RE_ARCHIVO = re.compile(r'^[A-Za-z0-9._\-]{1,200}$')


def _utc(dt):
    return dt.astimezone(dt_timezone.utc).strftime('%Y-%m-%dT%H:%M:%SZ')


def rango_utc_reporte(desde, hasta):
    """(begin, end) en UTC para días de Chile, con fin nunca en el futuro y
    al menos 24 horas de largo (MP exige mínimo un día de información)."""
    tz = timezone.get_current_timezone()
    inicio = timezone.make_aware(datetime.combine(desde, datetime.min.time()), tz)
    fin = timezone.make_aware(datetime.combine(hasta, datetime.max.time().replace(microsecond=0)), tz)
    # En UTC: restar dos horas locales del mismo huso ignora el cambio de hora.
    inicio, fin = inicio.astimezone(dt_timezone.utc), fin.astimezone(dt_timezone.utc)
    ahora = timezone.now().astimezone(dt_timezone.utc) - timedelta(minutes=5)
    if fin > ahora:
        fin = ahora
    if fin - inicio < timedelta(days=1):
        inicio = fin - timedelta(days=1)
    maximo = timedelta(days=60) - timedelta(minutes=1)
    if fin - inicio > maximo:
        inicio = fin - maximo
    return _utc(inicio), _utc(fin)


def pedir_reporte_liberaciones(config, desde, hasta):
    """Pide a MP que genere el reporte. Devuelve al instante (no espera)."""
    begin, end = rango_utc_reporte(desde, hasta)
    resp = mp._request(config, 'POST', '/v1/account/release_report',
                       json_body={'begin_date': begin, 'end_date': end}, cuenta_breaker=False)
    try:
        cuerpo = resp.json()
    except ValueError:
        cuerpo = {}
    logger.info("MP release_report crear config=%s %s..%s status=%s cuerpo=%s",
                config.id, begin, end, resp.status_code, str(cuerpo)[:500])
    if resp.status_code == 203:
        raise mp.MercadoPagoError(
            'Mercado Pago recibió la solicitud pero NO creó el reporte con esas fechas. '
            f'Pruebe con un rango que termine ayer. Detalle: {str(cuerpo)[:200]}', detalle=cuerpo)
    if resp.status_code >= 400:
        mp._json_o_error(resp, 'release_report crear')
    tarea = cuerpo.get('id') if isinstance(cuerpo, dict) else None
    return {'task_id': tarea, 'begin_date': begin, 'end_date': end,
            'estado': (cuerpo.get('status') if isinstance(cuerpo, dict) else '') or 'pending'}


def estado_tarea_liberaciones(config, task_id):
    """{'estado', 'file_name', 'listo'} de una solicitud de reporte."""
    resp = mp._request(config, 'GET', f'/v1/account/release_report/task/{int(task_id)}',
                       cuenta_breaker=False)
    data = mp._json_o_error(resp, 'release_report task') or {}
    estado = str(data.get('status') or '')
    archivo = data.get('file_name') or ''
    # MP usa estados que no documenta ('data-ready' = juntó los datos y falta el
    # archivo; 'enabled' = listo): lo que manda es que ya exista el archivo.
    return {'estado': estado, 'estado_texto': texto_estado_reporte(estado, archivo), 'file_name': archivo,
            'listo': bool(archivo),
            'fallido': not archivo and (estado in ESTADOS_REPORTE_FALLIDO or estado not in ESTADOS_REPORTE_EN_PROCESO)}


ESTADOS_REPORTE_FALLIDO = ('failed', 'error', 'rejected', 'cancelled', 'canceled')


def texto_estado_reporte(estado, archivo=''):
    """Estado de un pedido de reporte en palabras."""
    if archivo:
        return 'listo'
    return {'pending': 'en cola en Mercado Pago', 'processing': 'Mercado Pago lo está procesando',
            'data-ready': 'datos listos, Mercado Pago está armando el archivo'}.get(
        estado, 'no se pudo generar' if estado in ESTADOS_REPORTE_FALLIDO else (estado or 'en proceso'))


def _normalizar_reporte(item):
    return {
        'id': item.get('id'),
        'file_name': item.get('file_name') or '',
        'begin_date': str(item.get('begin_date') or ''),
        'end_date': str(item.get('end_date') or ''),
        'creado': str(item.get('date_created') or item.get('generation_date')
                      or item.get('last_modified') or ''),
        'origen': str(item.get('created_from') or ''),
        'estado': str(item.get('status') or ''),
    }


_RE_HORA_ARCHIVO = re.compile(r'-(\d{4}-\d{2}-\d{2})-(\d{6})\.\w+$')


def _instante_creado_mp(texto, file_name='', como_utc=None):
    """`date_created` de un reporte como instante real.

    MP lo entrega etiquetado «-04:00» pero el instante es UTC: el archivo
    `…-2026-09-23-163054.csv` (16:30:54 hora de Chile) trae
    `date_created 2026-09-23T19:30:54.000-04:00`. Leído tal cual, cada reporte
    parecía 4 h más nuevo (la página decía «revisado hasta 19:10» a las 15:44 y
    un pedido trabado hacía 6 h parecía reciente y se seguía esperando). Con
    nombre de archivo se comprueba contra su hora; sin él (pedidos en cola) se
    usa `como_utc`, lo aprendido de los reportes terminados de la misma cuenta.
    Devuelve (instante, como_utc): True si se leyó como UTC, False si tal cual,
    None si no se pudo comprobar.
    """
    dt = _instante(texto)
    if dt is None:
        return None, None
    m = _RE_HORA_ARCHIVO.search(str(file_name or ''))
    if m:
        try:
            local = datetime.strptime(f'{m.group(1)} {m.group(2)}', '%Y-%m-%d %H%M%S')
        except ValueError:
            local = None
        if local is not None:
            en_utc = dt.replace(tzinfo=dt_timezone.utc)
            d_utc = abs(timezone.localtime(en_utc).replace(tzinfo=None) - local)
            d_tal = abs(timezone.localtime(dt).replace(tzinfo=None) - local)
            # Las dos lecturas están horas aparte: gana la más cercana a la hora
            # del archivo (tolera 1 h si MP lo nombra en otro huso en invierno).
            if min(d_utc, d_tal) <= timedelta(minutes=90):
                return (en_utc, True) if d_utc <= d_tal else (dt, False)
    if como_utc:
        return dt.replace(tzinfo=dt_timezone.utc), True
    if como_utc is None and dt > timezone.now() + timedelta(minutes=5):
        # Sin reporte terminado con que comprobar: una creación «en el futuro»
        # solo la explica la etiqueta equivocada.
        return dt.replace(tzinfo=dt_timezone.utc), True
    return dt, (None if como_utc is None else False)


def _corregir_creados(reportes):
    """Deja `creado` de cada reporte como instante real en ISO UTC (ver
    `_instante_creado_mp`). Los pedidos en cola, sin archivo, siguen la regla
    que demostró el reporte terminado más reciente de la lista."""
    como_utc, sin_comprobar = None, []
    for r in reportes:
        if not r.get('creado'):
            continue
        inst, decidido = _instante_creado_mp(r['creado'], r.get('file_name'))
        if decidido is None:
            sin_comprobar.append(r)
            continue
        r['creado'] = inst.astimezone(dt_timezone.utc).isoformat()
        r['creado_local'] = timezone.localtime(inst).strftime('%Y-%m-%d %H:%M')
        if como_utc is None:
            como_utc = decidido
    for r in sin_comprobar:
        inst, _d = _instante_creado_mp(r['creado'], '', como_utc=como_utc)
        if inst is not None:
            r['creado'] = inst.astimezone(dt_timezone.utc).isoformat()
            r['creado_local'] = timezone.localtime(inst).strftime('%Y-%m-%d %H:%M')
    return reportes


def listar_reportes_liberaciones(config, limite=20, incluir_pendientes=False):
    """Reportes ya generados en MP, más recientes primero (solo los descargables).

    Usa /search (documentado: ordenado por creación descendente y con
    file_name); si falla, cae a /list, cuyo esquema actual no garantiza file_name.
    """
    reportes = []
    try:
        resp = mp._request(config, 'GET', '/v1/account/release_report/search',
                           params={'limit': limite, 'offset': 0}, cuenta_breaker=False)
        data = mp._json_o_error(resp, 'release_report search') or {}
        reportes = [_normalizar_reporte(i) for i in (data.get('results') or [])]
    except mp.MercadoPagoError as e:
        logger.warning("MP release_report search falló (%s); se usa /list", e.mensaje)
        resp = mp._request(config, 'GET', '/v1/account/release_report/list', cuenta_breaker=False)
        data = mp._json_o_error(resp, 'release_report list') or []
        if isinstance(data, dict):
            data = data.get('results') or []
        reportes = [_normalizar_reporte(i) for i in data]
    _corregir_creados(reportes)
    reportes.sort(key=lambda r: r['creado'], reverse=True)   # ISO UTC: el orden de texto es el cronológico
    if incluir_pendientes:
        return reportes[:limite]
    return [r for r in reportes if r['file_name']][:limite]


def descargar_reporte_liberaciones(config, file_name):
    """Contenido (bytes) de un reporte ya generado."""
    if not _RE_ARCHIVO.match(str(file_name or '')):
        raise mp.MercadoPagoError('Nombre de reporte inválido.')
    resp = mp._request(config, 'GET', f'/v1/account/release_report/{file_name}',
                       cuenta_breaker=False)
    if resp.status_code >= 400:
        mp._json_o_error(resp, 'release_report descargar')
    return resp.content


COLUMNAS_REPORTE = (
    'DATE', 'SOURCE_ID', 'EXTERNAL_REFERENCE', 'RECORD_TYPE', 'DESCRIPTION',
    'NET_CREDIT_AMOUNT', 'NET_DEBIT_AMOUNT', 'GROSS_AMOUNT', 'MP_FEE_AMOUNT',
    'PAYMENT_METHOD', 'PAYMENT_METHOD_TYPE', 'POS_NAME', 'EXTERNAL_POS_ID',
    'PAYOUT_BANK_ACCOUNT_NUMBER',
)


def leer_config_reporte(config):
    """Configuración actual del reporte en MP, o None si nunca se creó."""
    resp = mp._request(config, 'GET', '/v1/account/release_report/config', cuenta_breaker=False)
    if resp.status_code == 404:
        return None
    return mp._json_o_error(resp, 'release_report config') or None


def activar_reporte_por_retiro(config):
    """Configura MP para generar el reporte de Liberaciones después de cada retiro.

    Conserva la configuración existente y solo agrega: las columnas que usa la
    conciliación, `execute_after_withdrawal=true`, `include_withdrawal_at_end=true`
    y encabezados en inglés (`report_translation='en'`, los nombres que lee el
    sistema). Devuelve la configuración resultante.
    """
    actual = leer_config_reporte(config)
    base = dict(actual or {})
    columnas = [c.get('key') for c in (base.get('columns') or []) if isinstance(c, dict)]
    for clave in COLUMNAS_REPORTE:
        if clave not in columnas:
            columnas.append(clave)
    offset = timezone.localtime().utcoffset()
    horas = int(offset.total_seconds() // 3600) if offset else -4
    cuerpo = {
        'file_name_prefix': base.get('file_name_prefix') or 'release-report-retailmind',
        'columns': [{'key': c} for c in columnas if c],
        'frequency': base.get('frequency') or {'hour': 0, 'type': 'monthly', 'value': 1},
        'include_withdrawal_at_end': True,
        'execute_after_withdrawal': True,
        'report_translation': 'en',
        'display_timezone': f'GMT{horas:+03d}',
    }
    for clave in ('separator', 'notification_email_list', 'sftp_info', 'check_available_balance',
                  'compensate_detail'):
        if base.get(clave) not in (None, ''):
            cuerpo[clave] = base[clave]
    metodo = 'PUT' if actual else 'POST'
    resp = mp._request(config, metodo, '/v1/account/release_report/config',
                       json_body=cuerpo, cuenta_breaker=False)
    if metodo == 'POST' and resp.status_code == 409:
        resp = mp._request(config, 'PUT', '/v1/account/release_report/config',
                           json_body=cuerpo, cuenta_breaker=False)
    data = mp._json_o_error(resp, f'release_report config {metodo}') or cuerpo
    logger.info("MP release_report config actualizada config=%s execute_after_withdrawal=%s",
                config.id, data.get('execute_after_withdrawal'))
    return data


def cuentas_mp():
    """[(config, cuenta, etiqueta)] una por CUENTA de Mercado Pago (empresa/RUT).

    Los retiros y los reportes de Liberaciones son de la cuenta: basta una caja
    cualquiera de cada cuenta para leerlos.
    """
    vistas = {}
    for cfg in (MercadoPagoConfig.objects.select_related('sucursal__empresa', 'cuenta__empresa')
                .order_by('-habilitado', 'id')):
        cuenta = mp._cuenta_de(cfg)
        clave = f'c{cuenta.id}' if cuenta else f'cfg{cfg.id}'
        if clave in vistas:
            vistas[clave]['cajas'].add(cfg.sucursal.alias if cfg.sucursal_id else cfg.nombre)
            continue
        empresa = cuenta.empresa if cuenta else (cfg.sucursal.empresa if cfg.sucursal_id else None)
        vistas[clave] = {
            'config': cfg,
            'nombre': (getattr(empresa, 'nombre', '') or 'Cuenta Mercado Pago'),
            'rut': getattr(empresa, 'rut', '') or '',
            'cajas': {cfg.sucursal.alias if cfg.sucursal_id else cfg.nombre},
        }
    return list(vistas.values())


# Un reporte que termina hace menos de esto se considera al día (no se pide otro).
FRESCURA_REPORTE_MIN = int(os.environ.get('MP_REPORTE_FRESCURA_MIN', '15'))
# Más que esto sin que MP genere el reporte pedido, se deja de esperar (24-09:
# la cuenta de Nicole tenía dos pedidos en 'pending' desde las 09:35 y a las
# 15:44 la página seguía «Buscando retiros recientes…»).
ESPERA_PEDIDO_MIN = int(os.environ.get('MP_ESPERA_PEDIDO_MIN', '90'))
# Un día de payments/search puede traer varias páginas (~1,5 s cada una): no se
# empieza uno si queda menos que esto del presupuesto, para no pasarse del tope.
RESERVA_DIA_SEG = 4
# Un retiro que quedó con parte sin explicar se vuelve a aplicar con su reporte
# como máximo cada tantas horas (entre medio se completan N° y liberaciones).
REINTENTO_RETIRO_ABIERTO_H = 6
# La plata de un retiro se liberó días antes (Point libera a D+N): el reporte
# parte con este margen antes del retiro para traer esas liberaciones con su N°.
MARGEN_LIBERACION_DIAS = 14
DIAS_MAX_REPORTE = 59          # MP acepta hasta 60 días por reporte
_CLAVE_PEDIDO = 'conc_mp:pedido:{}'
_CLAVE_SIN_RETIROS = 'conc_mp:sin_retiros:{}'


def _dia_desde_retiros(config, ahora=None):
    """Día del último retiro (o del abierto más antiguo) del que parte el reporte.

    El día del último retiro registrado de la cuenta: el saldo que dejó ese
    retiro queda dentro y la cola FIFO parte bien (ese retiro se reprocesa y
    queda igual). Sin retiros registrados, lo máximo que acepta MP (59 días):
    el primer retiro se arma con toda la historia disponible.
    """
    ahora = ahora or timezone.now()
    tope = timezone.localtime(ahora).date() - timedelta(days=DIAS_MAX_REPORTE)
    configs = _configs_de_la_cuenta(config)
    ultimo = (RetiroMercadoPago.objects.filter(config_id__in=configs)
              .order_by('-fecha', '-id').values_list('fecha', flat=True).first())
    if not ultimo:
        return tope
    inicio = max(ultimo, tope)
    # Un retiro que quedó con parte sin explicar (p. ej. «Saldo anterior» porque
    # a sus ventas les faltaba la fecha de liberación) tiene que volver a entrar
    # en el reporte para poder recalcularse; si no, quedaba fuera de la ventana
    # apenas se registraba el retiro siguiente.
    for r in (RetiroMercadoPago.objects.filter(config_id__in=configs, fecha__gte=tope, fecha__lt=inicio)
              .annotate(n_trx=Count('transacciones')).order_by('fecha')):
        if retiro_abierto(r):
            return r.fecha
    return inicio


def inicio_reporte_cuenta(config, ahora=None):
    """Primer día del reporte que se pide para una cuenta: el del último retiro
    registrado (o del abierto más antiguo, para que se recalcule) menos
    MARGEN_LIBERACION_DIAS, porque la plata que se lleva un retiro se liberó
    días antes y sin esas filas el retiro quedaba explicado solo por «Saldo
    anterior» (Paola 23-09: $2,5M sin explicar con un reporte que partía el
    mismo día). Sin retiros registrados, lo máximo que acepta MP (59 días)."""
    ahora = ahora or timezone.now()
    tope = timezone.localtime(ahora).date() - timedelta(days=DIAS_MAX_REPORTE)
    return max(tope, _dia_desde_retiros(config, ahora) - timedelta(days=MARGEN_LIBERACION_DIAS))


def reportes_por_reintentar(config, ahora=None):
    """Reportes ya aplicados que vale la pena volver a aplicar: los de un retiro
    reciente que quedó con parte sin explicar («Saldo anterior» / «No
    explicado»). Suele resolverse solo al reprocesar, porque entre medio se
    completaron N° y fechas de liberación de sus ventas o se registraron pagos
    «MP manual». Como máximo cada REINTENTO_RETIRO_ABIERTO_H horas (queda
    anotado en `raw_reporte['reintento_en']`). Hasta ahora un retiro así quedaba
    en «Saldo anterior» para siempre salvo que alguien aplicara el reporte a mano."""
    ahora = ahora or timezone.now()
    limite = ahora - timedelta(hours=REINTENTO_RETIRO_ABIERTO_H)
    archivos = set()
    for r in (RetiroMercadoPago.objects
              .filter(config_id__in=_configs_de_la_cuenta(config),
                      fecha__gte=timezone.localtime(ahora).date() - timedelta(days=DIAS_MAX_REPORTE))
              .annotate(n_trx=Count('transacciones'))):
        raw = r.raw_reporte if isinstance(r.raw_reporte, dict) else {}
        ultimo = _instante(raw.get('reintento_en'))
        if retiro_abierto(r) and (ultimo is None or ultimo < limite):
            archivos.update(a for a in (raw.get('archivos') or []) if a)
            if raw.get('archivo'):
                archivos.add(raw['archivo'])   # formato anterior
    return archivos


def _cubierto_por_aplicado(reporte, reportes, aplicados):
    """Reporte sin aplicar que otro YA aplicado contiene entero (empieza antes
    o igual y se generó después): no aporta nada y se salta. Sin esto, los
    reportes que MP generó ayer antes del definitivo se volvían a bajar y
    procesar en cada «Detectar retiros» (la marca «sin retiros» es por worker) y
    cada vuelta se gastaba el tiempo en ellos."""
    desde, fin = _instante(reporte.get('begin_date')), _fin_de_reporte(reporte)
    if desde is None or fin is None:
        return False
    for a in reportes:
        if a is reporte or a.get('file_name') not in aplicados:
            continue
        a_desde, a_fin = _instante(a.get('begin_date')), _fin_de_reporte(a)
        if a_desde is not None and a_fin is not None and a_desde <= desde and a_fin >= fin:
            return True
    return False


def _desglose_envenenado(withdrawal_id, raw):
    """Desglose guardado con el error del 24-09: el crédito 'reserve_for_payout'
    (que lleva el N° del propio retiro) leído como pago sin venta dejó el retiro
    entero en «Sin caja» y, como eso cuenta como explicado, «cerrado» para
    siempre. Se reconoce porque `sin_local` trae su propio N°: se recalcula."""
    raw = raw if isinstance(raw, dict) else {}
    return str(withdrawal_id) in {str(k) for k in (raw.get('sin_local') or {})}


def retiro_abierto(retiro):
    """Retiro con parte sin explicar con ventas (se recalcula al reprocesar):
    desglose con «Saldo anterior»/«No explicado», o guardado sin desglose y sin
    ventas amarradas. Un retiro revertido no cuenta."""
    raw = retiro.raw_reporte if isinstance(retiro.raw_reporte, dict) else {}
    if retiro.estado == 'REVERTIDO' or raw.get('revertido'):
        return False
    por_caja = raw.get('por_caja')
    if por_caja:
        return _no_explicado(por_caja) > 0 or _desglose_envenenado(retiro.withdrawal_id, raw)
    n = getattr(retiro, 'n_trx', None)
    if n is None:
        n = retiro.transacciones.count()
    return n == 0


def _fin_de_reporte(reporte):
    """Hasta cuándo trae datos un reporte: su fin, pero nunca después de cuando
    se generó (MP pone de fin las 23:59 del día aunque se genere a mediodía; con
    eso el reporte parecía al día y no se pedía otro tras un retiro de la tarde).
    Sin fecha de creación y con fin en el futuro, no se sabe: None."""
    fin = _instante(reporte.get('end_date'))
    creado = _instante(reporte.get('creado'))
    if fin is None:
        return creado
    if creado is not None and creado < fin:
        return creado
    if creado is None and fin > timezone.now():
        return None
    return fin


def _pedido_valido(pedido):
    """Pedido de reporte que devuelve la página en las vueltas automáticas: solo
    los campos esperados y con forma válida (el N° de tarea lo valida MP con el
    token de la cuenta; el archivo, `_RE_ARCHIVO` al bajarlo)."""
    if not isinstance(pedido, dict):
        return None
    try:
        tarea = int(pedido.get('task_id') or 0)
    except (TypeError, ValueError):
        tarea = 0
    limpio = {k: str(pedido.get(k) or '')[:40]
              for k in ('begin_date', 'end_date', 'hasta', 'desde', 'pedido_a', 'pedido_en')}
    limpio['task_id'] = tarea or None
    limpio['adoptado'] = bool(pedido.get('adoptado'))
    archivo = str(pedido.get('file_name') or '')
    limpio['file_name'] = archivo if _RE_ARCHIVO.match(archivo) else ''
    return limpio if (limpio['task_id'] or limpio['file_name']) else None


ESTADOS_REPORTE_EN_PROCESO = ('pending', 'processing', 'data-ready', 'in_process', 'queued', '')


def _pedido_en_cola_mp(pendientes, minutos=None, cubre_desde=None, cubre_hasta=None):
    """El pedido más reciente que MP todavía está generando, con la forma de un
    pedido propio, para esperarlo en vez de pedir otro encima.

    Solo cuenta si sigue en proceso (un estado desconocido sin archivo se da por
    terminado), se pidió hace menos de `minutos` (más que eso, se pide de nuevo)
    y cubre lo que se necesita: empieza en `cubre_desde` o antes y, si se pide,
    llega hasta `cubre_hasta`. Su fin útil es la hora en que se pidió: MP pone
    de fin las 23:59 del día, pero no puede traer nada posterior a su creación.
    """
    limite = timezone.now() - timedelta(minutes=ESPERA_PEDIDO_MIN if minutos is None else minutos)
    for r in pendientes:
        creado = _instante(r.get('creado'))
        if (r.get('file_name') or not r.get('id') or str(r.get('estado') or '') not in ESTADOS_REPORTE_EN_PROCESO
                or creado is None or creado < limite):
            continue
        desde = _instante(r.get('begin_date'))
        hasta = _instante(r.get('end_date'))
        if cubre_desde is not None and (desde is None or timezone.localtime(desde).date() > cubre_desde):
            continue
        if cubre_hasta is not None and hasta is not None and timezone.localtime(hasta).date() < cubre_hasta:
            continue
        return {'task_id': int(r['id']), 'begin_date': r.get('begin_date') or '',
                'end_date': creado.isoformat(),
                'desde': timezone.localtime(desde).strftime('%d/%m/%Y') if desde else '',
                'hasta': timezone.localtime(creado).strftime('%d/%m %H:%M'),
                'pedido_a': timezone.localtime(creado).strftime('%H:%M'), 'pedido_en': creado.isoformat(),
                'estado_texto': texto_estado_reporte(r.get('estado')), 'adoptado': True}
    return None


def _pedidos_atascados(pendientes, ahora=None):
    """Instantes de los pedidos que MP tiene en cola desde hace más de
    ESPERA_PEDIDO_MIN: la cola de esa cuenta está trabada y pedir otro no ayuda."""
    limite = (ahora or timezone.now()) - timedelta(minutes=ESPERA_PEDIDO_MIN)
    return sorted(c for c in (_instante(r.get('creado')) for r in pendientes
                              if not r.get('file_name') and r.get('id')
                              and str(r.get('estado') or '') in ESTADOS_REPORTE_EN_PROCESO)
                  if c is not None and c < limite)


def _pedido_en_curso(cfg, reportes, pedido_pagina=None, pedido_mp=None):
    """Qué pasó con el reporte que se le pidió antes a MP para esta cuenta.

    El pedido se busca en la caché del proceso y, si no está (otro worker de
    gunicorn), en el que devuelve la página. Devuelve
    (pendiente, extra, resuelto, fallido, vencido):
    - pendiente: el pedido, mientras MP lo genera o hasta aplicar su archivo.
    - extra: el reporte que MP ya terminó y /search todavía no lista; se procesa igual.
    - resuelto: el pedido llegó; aunque haya quedado viejo porque MP tardó, no
      se pide otro en esta pasada (eso encadenaba pedidos sin fin).
    - fallido: MP no pudo generarlo.
    - vencido: el pedido (con 'minutos') si MP lleva más de ESPERA_PEDIDO_MIN
      sin generarlo: se deja de esperar; solo «forzar» pide otro.
    """
    clave = _CLAVE_PEDIDO.format(cfg.id)
    pedido = cache.get(clave) or _pedido_valido(pedido_pagina) or pedido_mp
    if not pedido:
        return None, None, False, False, None
    fin_pedido = _instante(pedido.get('end_date'))
    if fin_pedido is not None and any(
            f is not None and f >= fin_pedido - timedelta(minutes=1)
            for f in (_fin_de_reporte(r) for r in reportes)):
        cache.delete(clave)
        return None, None, True, False, None
    archivo = pedido.get('file_name') or ''

    def _vencido():
        pedido_en = _instante(pedido.get('pedido_en'))
        ahora = timezone.now()
        if pedido_en is not None and pedido_en < ahora - timedelta(minutes=ESPERA_PEDIDO_MIN):
            return dict(pedido, minutos=int((ahora - pedido_en).total_seconds() // 60))
        return None

    if not archivo and pedido.get('task_id'):
        try:
            estado = estado_tarea_liberaciones(cfg, pedido['task_id'])
        except mp.MercadoPagoError as e:
            if getattr(e, 'red', False):
                vencido = _vencido()   # MP saturado no alarga la espera
                if vencido:
                    cache.delete(clave)
                    return None, None, False, False, vencido
                return pedido, None, False, False, None   # sin red: se sigue esperando
            cache.delete(clave)                           # MP no reconoce la tarea (4xx): se da por perdida
            return None, None, False, True, None
        if estado['fallido']:
            cache.delete(clave)
            return None, None, False, True, None
        if estado['listo']:
            archivo = estado['file_name']
        pedido = dict(pedido, estado_texto=estado.get('estado_texto') or '')
    if not archivo:
        vencido = _vencido()
        if vencido:
            cache.delete(clave)
            return None, None, False, False, vencido
        return pedido, None, False, False, None
    pedido = dict(pedido, file_name=archivo)
    cache.set(clave, pedido, 60 * 20)   # hasta aplicarlo
    extra = {'file_name': archivo, 'begin_date': pedido.get('begin_date') or '',
             'end_date': pedido.get('end_date') or '', 'creado': '', 'origen': 'manual',
             'estado': 'processed'}
    return pedido, extra, True, False, None


def _pedir_si_hace_falta(cfg, reportes, fila, ahora=None, forzar=False, por_retiro_activo=False):
    """Si el reporte más nuevo de la cuenta quedó viejo, pide uno hasta ahora.

    Sin esto «Detectar retiros» solo miraba reportes ya generados: con el
    reporte automático apagado en MP, un retiro recién hecho no aparecía nunca.
    Deja en `fila['pedido']` el pedido (para que la página lo espere).
    """
    ahora = ahora or timezone.now()
    ultimo_fin = max((f for f in (_fin_de_reporte(r) for r in reportes) if f), default=None)
    # Con el reporte automático tras cada retiro activo en MP, un reporte de
    # hoy basta: MP genera otro solo cuando haya un retiro nuevo. Sin él, 15 min.
    frescura = timedelta(hours=24) if por_retiro_activo else timedelta(minutes=FRESCURA_REPORTE_MIN)
    if not forzar and ultimo_fin is not None and ultimo_fin >= ahora - frescura:
        fila['al_dia_hasta'] = timezone.localtime(ultimo_fin).strftime('%d/%m %H:%M')
        return
    try:
        pedido = pedir_reporte_liberaciones(cfg, inicio_reporte_cuenta(cfg, ahora),
                                            timezone.localtime(ahora).date())
    except mp.MercadoPagoError as e:
        fila['error_pedido'] = e.mensaje
        return
    fin = _instante(pedido.get('end_date'))
    pedido['hasta'] = timezone.localtime(fin).strftime('%d/%m %H:%M') if fin else ''
    pedido['pedido_a'] = timezone.localtime(ahora).strftime('%H:%M')
    pedido['pedido_en'] = ahora.isoformat()
    pedido['estado_texto'] = 'en cola en Mercado Pago'
    desde = _instante(pedido.get('begin_date'))
    pedido['desde'] = timezone.localtime(desde).strftime('%d/%m/%Y') if desde else ''
    cache.set(_CLAVE_PEDIDO.format(cfg.id), pedido, 60 * 20)
    fila['pedido'] = pedido


def detectar_retiros(presupuesto_seg=45, pedir=True, pedir_ids=(), pedidos=None, forzar=False):
    """Recorre TODAS las cuentas, aplica los reportes de Liberaciones nuevos y
    devuelve los retiros encontrados. Es lo que hace el botón «Detectar
    retiros» y el comando con --todas (para cron).

    - Solo aplica reportes que traen retiros (uno sin retiros no tiene nada que
      asociar) y que no estaban aplicados.
    - Antes de aplicar completa el N° de operación de los cobros del POS para
      que las ventas se crucen. Si no alcanzó el tiempo (o MP falló), el reporte
      se aplica igual (el retiro aparece) pero NO se marca aplicado: la próxima
      pasada lo rehace con más ventas cruzadas. Reprocesar es idempotente.
    - Si el reporte más nuevo de la cuenta quedó viejo y `pedir=True` (o la
      cuenta viene en `pedir_ids`), le pide a MP uno hasta ahora y lo devuelve
      en `fila['pedido']`: MP tarda unos minutos y la página vuelve a detectar
      (con `pedir=False` y los pedidos que espera en `pedidos`) cuando termina.
      Si no se alcanzó a pedir (faltó tiempo), `fila['pedir_pendiente']`.
      `forzar=True` ignora un pedido en curso (trabado en MP) y pide otro.
    - Cada retiro trae `nuevo`: False si ya estaba registrado (el reporte pedido
      parte el día del último retiro, así que lo vuelve a traer).
    - Tiene techo de tiempo: lo que no alcance queda para la próxima pasada.
    """
    import time as _time
    fin = _time.monotonic() + presupuesto_seg
    aplicados = reportes_aplicados()
    salida = []
    for info in cuentas_mp():
        cfg = info['config']
        fila = {'cuenta': info['nombre'], 'rut': info['rut'], 'config_id': cfg.id,
                'cajas': ', '.join(sorted(info['cajas'])), 'reportes_revisados': 0,
                'reportes_aplicados': 0, 'retiros': [], 'por_retiro_activo': None,
                'error': '', 'falto_tiempo': False, 'incompleto': False,
                'pedido': None, 'error_pedido': '', 'revisado_hasta': '', 'pedir_pendiente': False,
                'al_dia_hasta': '', 'pedido_vencido': None, 'atascados_mp': None,
                'espera_min': ESPERA_PEDIDO_MIN}
        salida.append(fila)
        pedir_cuenta = pedir or cfg.id in pedir_ids
        pedido_pagina = (pedidos or {}).get(cfg.id)
        if _time.monotonic() > fin:
            fila['falto_tiempo'] = True
            fila['pedir_pendiente'] = pedir_cuenta
            if pedido_pagina:
                fila['pedido'] = _pedido_valido(pedido_pagina)
            continue
        try:
            cfg_reporte = leer_config_reporte(cfg)
            fila['por_retiro_activo'] = bool(cfg_reporte and cfg_reporte.get('execute_after_withdrawal'))
        except mp.MercadoPagoError:
            pass
        try:
            todos = listar_reportes_liberaciones(cfg, limite=30, incluir_pendientes=True)
        except mp.MercadoPagoError as e:
            fila['error'] = e.mensaje
            continue
        reportes = [r for r in todos if r['file_name']]
        atascados = _pedidos_atascados([r for r in todos if not r['file_name']])
        if atascados:
            fila['atascados_mp'] = {'n': len(atascados),
                                    'desde': timezone.localtime(atascados[0]).strftime('%d/%m %H:%M')}
        if forzar and pedir_cuenta:
            cache.delete(_CLAVE_PEDIDO.format(cfg.id))
            pedido_pagina, pedido_mp = None, None
        else:
            # Basta con que parta el día del retiro (sin el margen de liberación):
            # trae el retiro, y el reintento de retiros abiertos completa el resto.
            pedido_mp = _pedido_en_cola_mp([r for r in todos if not r['file_name']],
                                           cubre_desde=_dia_desde_retiros(cfg))
        pendiente, extra, resuelto, fallido, vencido = _pedido_en_curso(cfg, reportes, pedido_pagina, pedido_mp)
        if vencido:
            fila['pedido_vencido'] = {'pedido_a': vencido.get('pedido_a') or '', 'desde': vencido.get('desde') or '',
                                      'minutos': int(vencido.get('minutos') or 0)}
        if extra and extra['file_name'] not in {r['file_name'] for r in reportes}:
            reportes = [extra] + reportes
        # Del más antiguo al más nuevo: los retiros se reconstruyen en orden.
        reintentar = reportes_por_reintentar(cfg)
        nuevos = [r for r in reportes
                  if (r['file_name'] in reintentar
                      or (r['file_name'] not in aplicados and not _cubierto_por_aplicado(r, reportes, aplicados)))
                  and not cache.get(_CLAVE_SIN_RETIROS.format(r['file_name']))]
        for rep in reversed(nuevos):
            if _time.monotonic() > fin:
                fila['falto_tiempo'] = True
                break
            try:
                filas = leer_csv(descargar_reporte_liberaciones(cfg, rep['file_name']))
            except mp.MercadoPagoError as e:
                fila['error'] = e.mensaje
                continue
            fila['reportes_revisados'] += 1
            try:
                previa = procesar_reporte_liberaciones(filas, cfg, aplicar=False, archivo=rep['file_name'])
                if filas and not previa.get('registros') and not previa.get('remanente_previo'):
                    # Ni fechas ni tipos reconocibles: no se cachea como «sin retiros», se avisa.
                    fila['error'] = (f'Reporte {rep["file_name"]} ilegible (encabezados: '
                                     f'{", ".join(list(filas[0].keys())[:12])}). Revise el idioma del reporte en MP.')
                    continue
                if not previa['retiros']:
                    # Sin retiros no hay nada que asociar: no se vuelve a bajar en un día.
                    cache.set(_CLAVE_SIN_RETIROS.format(rep['file_name']), 1, 60 * 60 * 24)
                    continue
                completo = True
                dias = dias_para_completar(cfg, previa, importar=True)
                if dias:
                    restante = max(5, int(fin - _time.monotonic()) - 5)
                    comp = completar_numeros_mp(cfg, dias, presupuesto_seg=restante, importar=True)
                    completo = not comp['sin_tiempo'] and not comp.get('fallidos')
                # Todo o nada: si el proceso muere a mitad (tope de gunicorn), no queda
                # un retiro creado con el archivo «aplicado» y los demás sin crear.
                with transaction.atomic():
                    res = procesar_reporte_liberaciones(filas, cfg, aplicar=True,
                                                        archivo=rep['file_name'] if completo else '')
            except Exception as e:  # noqa: BLE001 — un reporte roto no debe tumbar las demás cuentas
                logger.exception("Conciliación MP: error procesando %s", rep['file_name'])
                fila['error'] = f'No se pudo procesar {rep["file_name"]}: {e}'
                continue
            fila['reportes_aplicados'] += 1
            fila['retiros'].extend(res['retiros'])
            if completo:
                aplicados.add(rep['file_name'])
            else:
                fila['falto_tiempo'] = fila['incompleto'] = True
                break
        ultimo_fin = max((f for f in (_fin_de_reporte(r) for r in reportes) if f), default=None)
        if ultimo_fin is not None:
            fila['revisado_hasta'] = timezone.localtime(ultimo_fin).strftime('%d/%m %H:%M')
        if extra and (extra['file_name'] in aplicados
                      or cache.get(_CLAVE_SIN_RETIROS.format(extra['file_name']))):
            cache.delete(_CLAVE_PEDIDO.format(cfg.id))   # ya se aplicó entero
            pendiente = None
        # Un pedido adoptado (hecho por otro flujo antes del retiro) que llega
        # puede no traer el retiro: en un clic manual se vuelve a mirar la frescura.
        adoptado = bool((pedido_mp or {}).get('adoptado') or (pedido_pagina or {}).get('adoptado'))
        if pendiente:
            fila['pedido'] = pendiente
        elif pedir_cuenta and (not resuelto or adoptado or forzar) and (not vencido or forzar) \
                and (not atascados or forzar):
            if fila['incompleto'] or _time.monotonic() > fin:
                fila['pedir_pendiente'] = True   # la página lo pide en la vuelta siguiente
            else:
                _pedir_si_hace_falta(cfg, reportes, fila, forzar=forzar,
                                     por_retiro_activo=bool(fila['por_retiro_activo']))
        if fallido and not fila['pedido']:
            fila['error_pedido'] = ('Mercado Pago no pudo generar el reporte pedido. '
                                    'Vuelva a apretar «Detectar retiros».')
    return salida


def detalle_retiro(retiro):
    """Qué se llevó un retiro: las ventas del POS amarradas, los pagos «MP
    manual» y lo que el reporte no pudo explicar con ventas (saldo anterior,
    online). Una venta partida entre dos retiros aparece en el que se llevó su
    último peso."""
    raw = retiro.raw_reporte if isinstance(retiro.raw_reporte, dict) else {}
    netos = raw.get('netos') or {}
    ids = [str(i) for i in (raw.get('pagos') or []) if str(i).isdigit()]
    # Amarradas + las que este retiro ya contó por su N° y se registraron después
    # (p. ej. un pago «MP manual» importado): todavía sin retiro.
    trxs = list(TransaccionMercadoPago.objects
                .filter(Q(retiro=retiro) | Q(retiro__isnull=True, tipo='VENTA', payment_id_mp__in=ids))
                .select_related('sucursal', 'config', 'ticket', 'detalle_pago__ticket')
                .order_by('money_release_date', 'creado_en')) if ids else list(
        retiro.transacciones.select_related('sucursal', 'config', 'ticket', 'detalle_pago__ticket')
        .order_by('money_release_date', 'creado_en'))
    tickets = _ticket_de_cobros(trxs)
    # «MP manual» con la misma regla que al procesar: solo si ninguna transacción
    # tiene ese N° (si no, la venta del POS salía dos veces).
    if ids:
        con_trx = {str(x) for par in (TransaccionMercadoPago.objects
                                      .filter(Q(payment_id_mp__in=ids) | Q(payment_id__in=ids))
                                      .values_list('payment_id_mp', 'payment_id')) for x in par if x}
        ids = [i for i in ids if i not in con_trx]
    lineas = {}
    if ids:
        for pago in (TicketDetallePago.objects
                     .filter(Q(origen_pago='MANUAL') | Q(origen_pago__isnull=True),
                             metodo_pago__in=METODOS_MP, transacciones_mercadopago__isnull=True,
                             creado_en__date__gte=retiro.fecha - timedelta(days=60),
                             creado_en__date__lte=retiro.fecha + timedelta(days=2))
                     .exclude(ticket__estado='ANULADO').exclude(voucher__isnull=True)
                     .select_related('ticket__sucursal').order_by('creado_en')):
            if _solo_digitos(pago.voucher) not in ids:
                continue
            lineas.setdefault(_solo_digitos(pago.voucher), []).append(pago)
    dtes = _dtes_por_ref(list(tickets.values()) + [p.ticket for ps in lineas.values() for p in ps])
    filas, suma_neto, bruto_sin_neto = [], 0, 0
    for t in trxs:
        neto = int(t.monto_neto if t.monto_neto is not None else t.monto)
        suma_neto += neto
        tk = tickets.get(t.id)
        documento, dte_id = _documento_de_ticket(tk, dtes)
        filas.append({
            'origen': 'MP manual' if str(t.external_reference).startswith('MANUAL-') else 'POS',
            'fecha': timezone.localtime(t.creado_en).strftime('%d/%m/%Y %H:%M'),
            'liberacion': (timezone.localtime(t.money_release_date).strftime('%d/%m/%Y')
                           if t.money_release_date else ''),
            'sucursal': t.sucursal.alias if t.sucursal_id else '',
            'caja': t.config.nombre if t.config_id else '',
            'ticket': tk.correlativo if tk else (t.correlativo_ticket or ''),
            'documento': documento, 'dte_id': dte_id,
            'bruto': int(t.monto), 'comision': int(t.monto) - neto, 'neto': neto,
            'payment_id_mp': t.payment_id_mp, 'estado': t.estado,
        })
    for voucher, pagos in lineas.items():
        # Mismo N° en varias líneas: si son copias (mismo monto) es un solo pago;
        # si no, es un pago repartido y el bruto es la suma.
        montos = [int(p.monto or 0) for p in pagos]
        bruto = montos[0] if len(set(montos)) == 1 else sum(montos)
        neto = netos.get(voucher)
        neto = int(neto) if neto is not None else None
        if neto is not None and len(montos) > 1 and bruto < neto <= sum(montos):
            bruto = sum(montos)          # un pago repartido en líneas del mismo monto
        if neto is None:
            bruto_sin_neto += bruto
            comision = None
        else:
            suma_neto += neto
            comision = bruto - neto if neto <= bruto else None   # N° mal digitado: sin comisión
        pago = pagos[-1]
        documento, dte_id = _documento_de_ticket(pago.ticket, dtes)
        filas.append({
            'origen': 'MP manual', 'fecha': timezone.localtime(pago.creado_en).strftime('%d/%m/%Y %H:%M'),
            'liberacion': '', 'sucursal': pago.ticket.sucursal.alias if pago.ticket.sucursal_id else '',
            'caja': 'MP manual', 'ticket': ', '.join(sorted({str(p.ticket.correlativo) for p in pagos})),
            'documento': documento, 'dte_id': dte_id,
            'bruto': bruto, 'comision': comision, 'neto': neto, 'payment_id_mp': voucher, 'estado': '',
        })
    por_caja = raw.get('por_caja') or []
    no_ventas = [c for c in por_caja if isinstance(c, dict) and (
        c.get('caja') in (CAJA_SALDO_ANTERIOR, CAJA_NO_EXPLICADA)
        or str(c.get('caja', '')).startswith(('Sin caja', 'Caja MP «')))]
    # Lo que se contó «sin venta» y ya tiene su cobro (listado arriba) no se repite.
    sin_venta = _sin_venta_de(raw, {str(t.payment_id_mp) for t in trxs if t.payment_id_mp})
    return {
        'withdrawal_id': retiro.withdrawal_id, 'fecha': retiro.fecha.strftime('%d/%m/%Y'),
        'monto': retiro.monto, 'estado': retiro.estado, 'detalle': retiro.detalle_diferencia or '',
        'por_caja': por_caja, 'filas': filas, 'suma_neto': suma_neto,
        'bruto': sum(f['bruto'] for f in filas), 'comision': sum(f['comision'] or 0 for f in filas),
        'bruto_sin_neto': bruto_sin_neto, 'no_ventas': no_ventas, 'sin_venta': int(sin_venta or 0),
    }


def cuadre_por_sucursal(desde, hasta, sucursal_id=None, configs=None):
    """Por tienda, los cobros MP del período: vendido (bruto), devuelto,
    comisión, neto y dónde está ese neto: ya en el banco (un retiro se lo
    llevó), por liberar (MP aún no lo suelta) o disponible en MP."""
    d, h = rango_fechas(desde, hasta)
    ahora = timezone.now()
    qs = (TransaccionMercadoPago.objects
          .filter(tipo='VENTA', estado__in=('APROBADA', 'DEVUELTA', 'CONTRACARGO'),
                  creado_en__date__gte=d, creado_en__date__lte=h)
          .exclude(correlativo_ticket__startswith='PRUEBA-')
          .select_related('sucursal', 'retiro'))
    if sucursal_id:
        qs = qs.filter(sucursal_id=int(sucursal_id))
    if configs is not None:
        qs = qs.filter(config_id__in=list(configs))
    filas = {}
    campos = ('cobros', 'vendido', 'devuelto', 'comision', 'neto', 'en_banco', 'abonado', 'en_transito',
              'por_liberar', 'en_mp', 'sin_comision')
    for t in qs:
        f = filas.setdefault(t.sucursal_id, dict({'sucursal': t.sucursal.alias if t.sucursal_id else ''},
                                                 **{c: 0 for c in campos}))
        f['cobros'] += 1
        f['vendido'] += int(t.monto)
        if t.estado != 'APROBADA':
            f['devuelto'] += int(t.monto)   # la plata volvió al cliente
            continue
        neto = int(t.monto_neto if t.monto_neto is not None else t.monto)
        if t.monto_neto is None:
            f['sin_comision'] += 1
        f['comision'] += int(t.monto) - neto
        f['neto'] += neto
        if t.retiro_id:
            f['en_banco'] += neto
            f['abonado' if t.retiro.visto_en_cartola else 'en_transito'] += neto
        elif t.money_release_date and t.money_release_date > ahora:
            f['por_liberar'] += neto
        else:
            f['en_mp'] += neto
    # Devoluciones parciales: la venta sigue APROBADA y la devolución es otra fila.
    devoluciones = (TransaccionMercadoPago.objects
                    .filter(tipo='DEVOLUCION', estado='DEVUELTA', transaccion_origen__tipo='VENTA',
                            transaccion_origen__estado='APROBADA',
                            transaccion_origen__creado_en__date__gte=d, transaccion_origen__creado_en__date__lte=h)
                    .exclude(transaccion_origen__correlativo_ticket__startswith='PRUEBA-')
                    .select_related('transaccion_origen__retiro'))
    if sucursal_id:
        devoluciones = devoluciones.filter(transaccion_origen__sucursal_id=int(sucursal_id))
    if configs is not None:
        devoluciones = devoluciones.filter(transaccion_origen__config_id__in=list(configs))
    for dv in devoluciones:
        origen = dv.transaccion_origen
        f = filas.get(origen.sucursal_id)
        if f is None:
            continue
        f['devuelto'] += int(dv.monto)
        f['neto'] -= int(dv.monto)
        # De dónde sale: si se devolvió antes del retiro que se llevó la venta, ese
        # retiro llevó menos (sale de «en el banco»); si fue después, sale del saldo
        # de MP; si la venta aún no se libera, de lo por liberar.
        if origen.retiro_id:
            raw_r = origen.retiro.raw_reporte if isinstance(origen.retiro.raw_reporte, dict) else {}
            corte = _instante(raw_r.get('instante')) or timezone.make_aware(
                datetime.combine(origen.retiro.fecha, datetime.max.time()))
            clave = 'en_banco' if dv.creado_en <= corte else 'en_mp'
            if clave == 'en_banco':
                f['abonado' if origen.retiro.visto_en_cartola else 'en_transito'] -= int(dv.monto)
        elif origen.money_release_date and origen.money_release_date > ahora:
            clave = 'por_liberar'
        else:
            clave = 'en_mp'
        f[clave] -= int(dv.monto)
    filas = sorted(filas.values(), key=lambda f: -f['vendido'])
    total = {c: sum(f[c] for f in filas) for c in campos}
    total['sucursal'] = 'Total'
    return {'filas': filas, 'total': total, 'desde': str(d), 'hasta': str(h)}


# ─────────────── Asignación de retiros: cobros del mes ↔ retiros ───────────────

ESTADOS_ASIGNACION = {
    'EN_BANCO': 'En el banco',
    'EN_TRANSITO': 'Enviado al banco (sin abono)',
    'POR_LIBERAR': 'Por liberar',
    'EN_MP': 'En Mercado Pago',
    'DEVUELTA': 'Devuelto',
    'SIN_REGISTRAR': 'MP manual sin registrar',
}


def _primer_dia_mes(valor):
    """'2026-09' → date(2026, 9, 1); None si no calza o el año es absurdo."""
    m = re.match(r'^(\d{4})-(\d{2})$', str(valor or '').strip())
    if not m or not 1 <= int(m.group(2)) <= 12 or not 2000 <= int(m.group(1)) <= 2100:
        return None
    return date(int(m.group(1)), int(m.group(2)), 1)


def _sin_venta_de(raw, con_cobro=()):
    """Cuánto de un retiro no se pudo explicar con una venta del sistema, según
    su `raw_reporte`, descontando las operaciones que después sí tuvieron su
    cobro (`con_cobro`: N° de operación con transacción)."""
    por_caja = raw.get('por_caja') or []
    if 'sin_venta' not in raw:
        return sum(int(c.get('monto') or 0) for c in por_caja if isinstance(c, dict) and (
            c.get('caja') in (CAJA_SALDO_ANTERIOR, CAJA_NO_EXPLICADA)
            or str(c.get('caja', '')).startswith(('Sin caja', 'Caja MP «'))))
    sin_local = raw.get('sin_local') or {}
    return max(0, int(raw.get('sin_venta') or 0)
               - sum(int(v) for k, v in sin_local.items() if str(k) in set(con_cobro)))


def dias_habiles(desde, hasta):
    """Días de lunes a viernes entre dos fechas (sin contar `desde`)."""
    n, d = 0, desde
    while d < hasta:
        d += timedelta(days=1)
        if d.weekday() < 5:
            n += 1
    return n


def etapa_bancaria(retiro, hoy=None):
    """En qué va la plata de un retiro: ABONADO (visto en cartola), REVERTIDO
    (el banco lo devolvió a MP) o EN_TRANSITO (MP ya la envió; el banco aún no la
    acredita o nadie lo marcó). Alerta a los 2 días hábiles sin abono."""
    hoy = hoy or timezone.localdate()
    raw = retiro.raw_reporte if isinstance(retiro.raw_reporte, dict) else {}
    dias = max(0, (hoy - retiro.fecha).days)
    habiles = dias_habiles(retiro.fecha, hoy)
    base = {'dias_transito': dias, 'dias_habiles': habiles, 'fecha_abono': raw.get('fecha_abono') or '',
            'abono_origen': raw.get('abono_origen') or ''}
    if retiro.estado == 'REVERTIDO' or raw.get('revertido'):
        return dict(base, etapa='REVERTIDO', alerta_transito=False)
    if retiro.visto_en_cartola:
        return dict(base, etapa='ABONADO', alerta_transito=False)
    return dict(base, etapa='EN_TRANSITO', alerta_transito=habiles > 1)


def marcar_abono(retiro, visto, origen='MANUAL', fecha=None):
    """Marca (o desmarca) el abono en el banco de un retiro y guarda cuándo y cómo."""
    raw = dict(retiro.raw_reporte or {}) if isinstance(retiro.raw_reporte, dict) else {}
    if visto:
        raw['fecha_abono'] = str(fecha or timezone.localdate())
        raw['abono_origen'] = origen
    else:
        raw.pop('fecha_abono', None)
        raw.pop('abono_origen', None)
    retiro.visto_en_cartola = bool(visto)
    retiro.raw_reporte = raw
    retiro.save(update_fields=['visto_en_cartola', 'raw_reporte', 'actualizado_en'])


def _fin_de_mes(primero):
    return (primero + timedelta(days=32)).replace(day=1) - timedelta(days=1)


def _pagos_manuales_sin_registrar(desde, hasta, sucursales=None):
    """Pagos «MP manual» de ventas cerradas que todavía no tienen su cobro: su
    plata no se puede asignar a un retiro hasta registrarlos («Detectar
    retiros» los registra por su N°; si no calzan, «Asociar»)."""
    qs = (TicketDetallePago.objects
          .filter(Q(origen_pago='MANUAL') | Q(origen_pago__isnull=True),
                  metodo_pago__in=METODOS_MP, ticket__estado='PAGADO',
                  transacciones_mercadopago__isnull=True,
                  creado_en__date__gte=desde, creado_en__date__lte=hasta)
          .select_related('ticket__sucursal'))
    if sucursales is not None:
        qs = qs.filter(ticket__sucursal_id__in=list(sucursales))
    return qs


def _por_tienda(por_caja, amarradas):
    """[{tienda, monto, ventas}] de un retiro. El monto sale del desglose del
    reporte (cubre pagos «MP manual» y ventas partidas); el conteo, de las
    ventas amarradas. Sin desglose (retiro antiguo), solo las amarradas."""
    if not por_caja:
        return list(amarradas)
    conteo = {x['tienda']: x['ventas'] for x in amarradas}
    salida = {}
    for c in por_caja:
        etiqueta = str(c.get('caja', '')) if isinstance(c, dict) else ''
        if ' · ' not in etiqueta:
            continue
        tienda = etiqueta.split(' · ')[0]
        salida[tienda] = salida.get(tienda, 0) + int(c.get('monto') or 0)
    return [{'tienda': t, 'monto': m, 'ventas': conteo.get(t)} for t, m in sorted(salida.items(), key=lambda x: -x[1])]


def asignaciones_mp(mes=None, configs=None, sucursal_id=None, meses=6, max_cobros=5000):
    """Cuánto de lo cobrado con Mercado Pago ya se llevó un retiro al banco.

    - `serie`: los `meses` meses hasta el elegido, con neto cobrado, ya en el
      banco, por liberar, en MP, devuelto, pagos MP manual sin registrar y % completado.
    - `retiros`: los retiros con fecha en el mes y cuánto de cada uno está
      explicado con ventas.
    - `cobros`: cada cobro del mes y dónde está su plata (con su retiro).
    `configs`: cajas de UNA cuenta MP (None = todas). `sucursal_id` filtra la tienda.
    """
    hoy = timezone.localdate()
    primero = _primer_dia_mes(mes) or hoy.replace(day=1)
    ultimo = _fin_de_mes(primero)
    ahora = timezone.now()
    # Pagos «MP manual» que un retiro ya contó por su N° (aunque no se hayan
    # podido registrar como cobro): no están «sin registrar».
    retiro_de = _retiros_por_operacion(configs)
    sucursales = None
    if configs is not None:
        sucursales = set(MercadoPagoConfig.objects.filter(id__in=list(configs)).values_list('sucursal_id', flat=True))
    if sucursal_id:
        sucursales = {int(sucursal_id)} if sucursales is None else (sucursales & {int(sucursal_id)})

    # ── Avance por mes ──
    serie = []
    m = primero
    for _ in range(max(1, int(meses))):
        fin = _fin_de_mes(m)
        tot = cuadre_por_sucursal(str(m), str(fin), sucursal_id=sucursal_id, configs=configs)['total']
        pend = [p for p in _pagos_manuales_sin_registrar(m, fin, sucursales)
                if _solo_digitos(p.voucher) not in retiro_de]
        tot = {k: v for k, v in tot.items() if k != 'sucursal'}
        tot.update({
            'mes': m.strftime('%Y-%m'),
            'manual_sin_registrar': sum(int(p.monto or 0) for p in pend), 'manual_sin_registrar_n': len(pend),
            # % de lo cobrado (neto) que un retiro ya se llevó al banco
            'pct': round(100 * tot['en_banco'] / tot['neto']) if tot['neto'] > 0 else None,
        })
        serie.append(tot)
        m = (m - timedelta(days=1)).replace(day=1)
    serie.reverse()

    # ── Retiros del mes ──
    retiros_qs = RetiroMercadoPago.objects.filter(fecha__gte=primero, fecha__lte=ultimo)
    if configs is not None:
        retiros_qs = retiros_qs.filter(config_id__in=list(configs))
    elif sucursales is not None:
        cuentas = set()
        for cfg in MercadoPagoConfig.objects.filter(sucursal_id__in=list(sucursales)).select_related('sucursal', 'cuenta'):
            cuentas.update(_configs_de_la_cuenta(cfg))
        retiros_qs = retiros_qs.filter(config_id__in=list(cuentas))
    alias = {}
    if sucursal_id:
        from app.models import Sucursal
        alias_tienda = Sucursal.objects.filter(pk=int(sucursal_id)).values_list('alias', flat=True).first() or ''
    else:
        alias_tienda = ''
    retiros = []
    lista = list(retiros_qs.select_related('config__sucursal__empresa', 'config__cuenta__empresa')
                 .annotate(n_trx=Count('transacciones'),
                           neto_trx=Sum(Coalesce('transacciones__monto_neto', 'transacciones__monto')))
                 .order_by('-fecha', '-id'))
    # Ventas amarradas por tienda (lo que de verdad se llevó cada retiro) y
    # operaciones «sin venta» que después sí tuvieron su cobro.
    por_tienda = {}
    for fila in (TransaccionMercadoPago.objects.filter(retiro_id__in=[r.id for r in lista])
                 .values('retiro_id', 'sucursal__alias')
                 .annotate(s=Sum(Coalesce('monto_neto', 'monto')), n=Count('id'))):
        por_tienda.setdefault(fila['retiro_id'], []).append(
            {'tienda': fila['sucursal__alias'] or '', 'monto': int(fila['s'] or 0), 'ventas': fila['n']})
    claves = set()
    for r in lista:
        raw = r.raw_reporte if isinstance(r.raw_reporte, dict) else {}
        claves.update(str(k) for k in (raw.get('sin_local') or {}))
    con_cobro = set(TransaccionMercadoPago.objects.filter(tipo='VENTA', payment_id_mp__in=list(claves))
                    .values_list('payment_id_mp', flat=True)) if claves else set()
    for r in lista:
        raw = r.raw_reporte if isinstance(r.raw_reporte, dict) else {}
        por_caja = raw.get('por_caja') or []
        if por_caja:
            sin_venta = _sin_venta_de(raw, con_cobro)
        else:
            # Guardado antes del desglose: lo explican sus ventas amarradas.
            sin_venta = max(0, r.monto - int(r.neto_trx or 0))
        explicado = max(0, r.monto - sin_venta)
        if r.config_id not in alias:
            cuenta = _cuenta_efectiva(r.config)
            empresa = cuenta.empresa if cuenta else (r.config.sucursal.empresa if r.config.sucursal_id else None)
            alias[r.config_id] = getattr(empresa, 'nombre', '') or ''
        instante = _instante(raw.get('instante'))
        retiros.append({
            'withdrawal_id': r.withdrawal_id, 'fecha': r.fecha.strftime('%d/%m/%Y'),
            'hora': timezone.localtime(instante).strftime('%H:%M') if instante else '',
            'cuenta': alias[r.config_id], 'monto': r.monto, 'estado': r.estado,
            'visto_en_cartola': r.visto_en_cartola,
            'ventas': r.n_trx, 'neto_ventas': int(r.neto_trx or 0),
            'explicado': explicado, 'sin_venta': sin_venta,
            'pct_explicado': round(100 * explicado / r.monto) if r.monto else None,
            # qué se llevó de cada tienda: el desglose del reporte (incluye pagos MP
            # manual y ventas partidas); las ventas amarradas dan el conteo
            'por_tienda': _por_tienda(por_caja, por_tienda.get(r.id) or []),
            # con una tienda elegida: cuánto de este retiro salió de ella
            'de_la_tienda': (sum(x['monto'] for x in _por_tienda(por_caja, por_tienda.get(r.id) or [])
                                 if x['tienda'] == alias_tienda) if alias_tienda else None),
            'por_caja': por_caja,
            **etapa_bancaria(r, hoy),
        })
    vigentes = [r for r in lista if etapa_bancaria(r, hoy)['etapa'] != 'REVERTIDO']
    retirado_mes = sum(r.monto for r in vigentes)
    sin_abonar = [r for r in vigentes if not r.visto_en_cartola]
    arrastre_qs = RetiroMercadoPago.objects.filter(fecha__lt=primero, fecha__gte=primero - timedelta(days=120),
                                                   visto_en_cartola=False).exclude(estado='REVERTIDO')
    if configs is not None:
        arrastre_qs = arrastre_qs.filter(config_id__in=list(configs))
    arrastre = [r for r in arrastre_qs if not (isinstance(r.raw_reporte, dict) and r.raw_reporte.get('revertido'))]

    # ── Cobros del mes ──
    qs = (TransaccionMercadoPago.objects
          .filter(tipo='VENTA', estado__in=('APROBADA', 'DEVUELTA', 'CONTRACARGO'),
                  creado_en__date__gte=primero, creado_en__date__lte=ultimo)
          .exclude(correlativo_ticket__startswith='PRUEBA-')
          .select_related('sucursal', 'config', 'retiro', 'ticket', 'detalle_pago__ticket')
          .order_by('creado_en'))
    if sucursal_id:
        qs = qs.filter(sucursal_id=int(sucursal_id))
    if configs is not None:
        qs = qs.filter(config_id__in=list(configs))
    total_cobros = qs.count()
    trxs = list(qs[:max_cobros])
    # Devoluciones parciales (la venta sigue APROBADA): se restan del neto de la fila.
    devuelto_de = dict(TransaccionMercadoPago.objects
                       .filter(tipo='DEVOLUCION', estado='DEVUELTA', transaccion_origen_id__in=[t.id for t in trxs])
                       .values('transaccion_origen_id').annotate(s=Sum('monto'))
                       .values_list('transaccion_origen_id', 's')) if trxs else {}
    manuales = list(_pagos_manuales_sin_registrar(primero, ultimo, sucursales).order_by('creado_en')[:max_cobros])
    tickets = _ticket_de_cobros(trxs)
    dtes = _dtes_por_ref(list(tickets.values()) + [p.ticket for p in manuales])
    cobros, conteo = [], {k: 0 for k in ESTADOS_ASIGNACION}
    for t in trxs:
        neto = int(t.monto_neto if t.monto_neto is not None else t.monto)
        devuelto = int(devuelto_de.get(t.id) or 0)
        neto = max(0, neto - devuelto)
        if t.estado != 'APROBADA':
            estado = 'DEVUELTA'
        elif t.retiro_id:
            estado = 'EN_BANCO' if t.retiro.visto_en_cartola else 'EN_TRANSITO'
        elif t.money_release_date and t.money_release_date > ahora:
            estado = 'POR_LIBERAR'
        else:
            estado = 'EN_MP'
        conteo[estado] += 1
        tk = tickets.get(t.id)
        documento, _dte_id = _documento_de_ticket(tk, dtes)
        ref = str(t.external_reference or '')
        cobros.append({
            'fecha': timezone.localtime(t.creado_en).strftime('%d/%m/%Y %H:%M'),
            'sucursal': t.sucursal.alias if t.sucursal_id else '',
            'caja': t.config.nombre if t.config_id else '',
            'ticket': tk.correlativo if tk else (t.correlativo_ticket or ''),
            'documento': documento,
            'origen': 'MP manual' if ref.startswith('MANUAL-') else ('Tu caja' if 'DIRECTO-' in ref else 'POS'),
            'bruto': int(t.monto), 'neto': neto, 'devuelto': devuelto, 'payment_id_mp': t.payment_id_mp,
            'liberacion': (timezone.localtime(t.money_release_date).strftime('%d/%m/%Y')
                           if t.money_release_date else ''),
            'estado': estado,
            'retiro': t.retiro.withdrawal_id if t.retiro_id else '',
            'retiro_fecha': t.retiro.fecha.strftime('%d/%m/%Y') if t.retiro_id else '',
        })
    for pago in manuales:
        documento, _dte_id = _documento_de_ticket(pago.ticket, dtes)
        voucher = _solo_digitos(pago.voucher)
        contado = retiro_de.get(voucher)
        if contado is not None:
            estado = 'EN_BANCO' if contado.visto_en_cartola else 'EN_TRANSITO'
        else:
            estado = 'SIN_REGISTRAR'
        conteo[estado] += 1
        raw_c = contado.raw_reporte if contado is not None and isinstance(contado.raw_reporte, dict) else {}
        neto_manual = (raw_c.get('netos') or {}).get(voucher) if contado is not None else None
        cobros.append({
            'fecha': timezone.localtime(pago.creado_en).strftime('%d/%m/%Y %H:%M'),
            'sucursal': pago.ticket.sucursal.alias if pago.ticket.sucursal_id else '',
            'caja': 'MP manual', 'ticket': pago.ticket.correlativo, 'documento': documento,
            'origen': 'MP manual', 'bruto': int(pago.monto or 0),
            'neto': int(neto_manual) if neto_manual is not None else None, 'devuelto': 0,
            'payment_id_mp': str(pago.voucher or '').strip(), 'liberacion': '',
            'estado': estado, 'retiro': contado.withdrawal_id if contado is not None else '',
            'retiro_fecha': contado.fecha.strftime('%d/%m/%Y') if contado is not None else '',
        })
    return {
        'mes': primero.strftime('%Y-%m'), 'desde': str(primero), 'hasta': str(ultimo),
        'serie': serie, 'resumen': serie[-1] if serie else {}, 'retiros': retiros,
        'cobros': cobros, 'conteo': conteo, 'total_cobros': total_cobros,
        'recortado': total_cobros > len(trxs), 'estados': ESTADOS_ASIGNACION,
        # plata que MP ya envió al banco y todavía no está confirmada en la cartola
        'retirado_mes': retirado_mes,
        'sin_abonar': sum(r.monto for r in sin_abonar), 'sin_abonar_n': len(sin_abonar),
        'sin_abonar_alerta': sum(1 for r in sin_abonar if etapa_bancaria(r, hoy)['alerta_transito']),
        'revertidos': len(lista) - len(vigentes),
        # de meses anteriores, todavía sin abono confirmado
        'arrastre_sin_abonar': sum(r.monto for r in arrastre), 'arrastre_sin_abonar_n': len(arrastre),
    }


# ───────────────────────────── 3. Cartola del banco ───────────────────────────

def _filas_cartola(contenido, nombre=''):
    """Filas (dict) de la cartola: CSV con preámbulo (Banco de Chile, Santander,
    BCI escriben «Cartola», «Cuenta: …» antes de los encabezados) o XLSX.
    Empieza en la primera fila que tenga FECHA y una columna de abono/monto."""
    nombre = str(nombre or '').lower()
    filas_crudas = []
    if nombre.endswith(('.xlsx', '.xlsm')) or (isinstance(contenido, (bytes, bytearray)) and contenido[:2] == b'PK'):
        from openpyxl import load_workbook
        libro = load_workbook(io.BytesIO(contenido), read_only=True, data_only=True)
        hoja = libro[libro.sheetnames[0]]
        for fila in hoja.iter_rows(values_only=True):
            filas_crudas.append(['' if v is None else (v.strftime('%Y-%m-%d') if hasattr(v, 'strftime') else str(v))
                                 for v in fila])
    else:
        texto = contenido.decode('utf-8-sig', errors='replace') if isinstance(contenido, (bytes, bytearray)) else str(contenido)
        muestra = texto[:4000]
        sep = ';' if muestra.count(';') >= muestra.count(',') else ','
        for linea in texto.splitlines():
            filas_crudas.append(next(csv.reader([linea], delimiter=sep)) if linea.strip() else [])
    encabezado = None
    salida = []
    for fila in filas_crudas:
        celdas = [_sin_tildes(str(c or '')).strip().upper() for c in fila]
        if encabezado is None:
            if any('FECHA' in c or c == 'DATE' for c in celdas) and any(
                    any(p in c for p in ('ABONO', 'DEPOSITO', 'CREDITO', 'HABER', 'MONTO', 'IMPORTE', 'AMOUNT')) for c in celdas):
                encabezado = celdas
            continue
        if not any(str(c or '').strip() for c in fila):
            continue
        salida.append({encabezado[i]: (fila[i] if i < len(fila) else '') for i in range(len(encabezado))})
    return salida


def leer_cartola(contenido, nombre=''):
    """Cartola del banco (CSV o XLSX) → [(fecha, monto_abono, descripcion)]. Toma la
    columna de abonos si existe (Abonos / Depósitos / Crédito), si no la de monto."""
    movimientos = []
    for fila in _filas_cartola(contenido, nombre):
        claves = list(fila)
        col_fecha = next((k for k in claves if 'FECHA' in k or k == 'DATE'), None)
        col_abono = next((k for k in claves if any(p in k for p in ('ABONO', 'DEPOSITO', 'CREDITO', 'HABER'))), None)
        col_monto = col_abono or next((k for k in claves if 'MONTO' in k or 'IMPORTE' in k or k == 'AMOUNT'), None)
        col_desc = next((k for k in claves if 'DESCRIP' in k or 'GLOSA' in k or 'DETALLE' in k), None)
        if not col_fecha or not col_monto:
            continue
        fecha = _fecha_libre(fila.get(col_fecha))
        monto = _monto(fila.get(col_monto))
        if fecha is None or monto <= 0:
            continue
        movimientos.append((fecha, monto, str(fila.get(col_desc, '') if col_desc else '')))
    return movimientos


TOLERANCIA_APROXIMADA = 1000   # pesos de diferencia que se muestran como «calce aproximado»


def conciliar_cartola(movimientos, tolerancia_dias=5, aplicar=False, configs=None):
    """Marca «visto en cartola» los retiros que calzan con un abono del banco.

    Calce: mismo monto y fecha del abono entre la del retiro y `tolerancia_dias`
    después (el banco abona el mismo día o el hábil siguiente; con feriado, más).
    Cada abono se usa una sola vez; un retiro ya visto no toma abonos nuevos;
    ante empate gana el que menciona Mercado Pago. `configs`: cajas de la cuenta
    MP de esta cartola (cada cuenta retira a su propio banco): sin esto, los
    retiros de la otra empresa salían «sin abono» y podían robarse un abono.
    Devuelve también los calces aproximados (≤ $1.000 de diferencia), que no se
    aplican solos, y los retiros sin abono hasta la fecha de la cartola.
    """
    if not movimientos:
        return {'calzados': [], 'aproximados': [], 'sin_calce_banco': [], 'retiros_sin_abono': []}
    desde = min(m[0] for m in movimientos) - timedelta(days=tolerancia_dias)
    hasta = max(m[0] for m in movimientos)
    qs = (RetiroMercadoPago.objects.filter(fecha__gte=hasta - timedelta(days=60), fecha__lte=hasta)
          .exclude(estado='REVERTIDO').order_by('fecha', 'id'))
    if configs is not None:
        qs = qs.filter(config_id__in=list(configs))
    retiros = [r for r in qs if not (isinstance(r.raw_reporte, dict) and r.raw_reporte.get('revertido'))]
    usados = set()
    calzados, aproximados = [], []
    for ret in retiros:
        if ret.visto_en_cartola or ret.fecha < desde:
            continue
        candidatos = [
            (i, m) for i, m in enumerate(movimientos)
            if i not in usados and m[1] == ret.monto
            and 0 <= (m[0] - ret.fecha).days <= tolerancia_dias
        ]
        if not candidatos:
            cerca = [(i, m) for i, m in enumerate(movimientos)
                     if i not in usados and 0 < abs(m[1] - ret.monto) <= TOLERANCIA_APROXIMADA
                     and 0 <= (m[0] - ret.fecha).days <= tolerancia_dias]
            if cerca:
                i, mov = min(cerca, key=lambda c: abs(c[1][1] - ret.monto))
                aproximados.append({'withdrawal_id': ret.withdrawal_id, 'fecha_retiro': str(ret.fecha),
                                    'fecha_banco': str(mov[0]), 'monto': ret.monto, 'monto_banco': mov[1],
                                    'diferencia': mov[1] - ret.monto, 'glosa': mov[2][:80]})
            continue
        candidatos.sort(key=lambda c: (0 if 'MERCADO' in _sin_tildes(c[1][2]).upper() else 1,
                                       (c[1][0] - ret.fecha).days))
        i, mov = candidatos[0]
        usados.add(i)
        calzados.append({'withdrawal_id': ret.withdrawal_id, 'fecha_retiro': str(ret.fecha),
                         'fecha_banco': str(mov[0]), 'monto': ret.monto, 'glosa': mov[2][:80]})
        if aplicar:
            marcar_abono(ret, True, origen='CARTOLA', fecha=mov[0])
    retiros_calzados = {c['withdrawal_id'] for c in calzados}
    montos_retiros = {r.monto for r in retiros}
    return {
        'calzados': calzados,
        'aproximados': aproximados,
        'sin_calce_banco': [
            {'fecha': str(m[0]), 'monto': m[1], 'glosa': m[2][:80]}
            for i, m in enumerate(movimientos)
            if i not in usados and ('MERCADO' in _sin_tildes(m[2]).upper() or m[1] in montos_retiros)
        ],
        'retiros_sin_abono': [
            {'withdrawal_id': r.withdrawal_id, 'fecha': str(r.fecha), 'monto': r.monto}
            for r in retiros if r.withdrawal_id not in retiros_calzados and not r.visto_en_cartola
        ],
    }


# ─────────────────────────── 4. Cruce contra la API de MP ─────────────────────

def diferencias_contra_mp(desde, hasta, sucursal_id=None):
    """Cruza lo que MP cobró (payments/search, por cuenta y día) con el sistema.

    Devuelve:
      - `cajas`: por caja y día, sistema vs MP (reutiliza conciliar_cierre_mp)
        junto al teórico MP y el cierre físico del arqueo de esa sucursal/día.
      - `sin_registro`: pagos aprobados en MP que el sistema no tiene (ni como
        cobro integrado ni como pago manual con ese N° de operación).
      - `manuales_sin_pago`: pagos "MP manual" del POS cuyo N° no aparece en MP.
      - `errores`: cuentas que no se pudieron consultar.
    Tope: 7 días (cada día es 1+ llamadas por cuenta).
    """
    d, h = rango_fechas(desde, hasta, dias_defecto=1, max_dias=7)
    configs = list(MercadoPagoConfig.objects.select_related('sucursal', 'cuenta')
                   .filter(habilitado=True))
    if sucursal_id:
        configs = [c for c in configs if c.sucursal_id == int(sucursal_id)]

    pagos_por_token_dia = {}
    config_de_token = {}
    atribucion = {}              # N° de pago -> {config_id} donde el cierre lo atribuyó
    errores = []
    cajas = []
    ids_mp_vistos = set()
    sin_registro = []
    dias = [d + timedelta(days=i) for i in range((h - d).days + 1)]
    cfg_por_id = {c.id: c for c in configs}

    for cfg in configs:
        try:
            token = mp._token(cfg)
        except mp.MercadoPagoError as e:
            errores.append({'caja': cfg.nombre, 'sucursal': cfg.sucursal.alias, 'error': e.mensaje})
            continue
        config_de_token.setdefault(token, cfg)
        for dia in dias:
            clave = (token, dia)
            if clave not in pagos_por_token_dia:
                try:
                    pagos_por_token_dia[clave] = mp.buscar_pagos_dia(cfg, dia)
                except mp.MercadoPagoError as e:
                    pagos_por_token_dia[clave] = None
                    errores.append({'caja': cfg.nombre, 'sucursal': cfg.sucursal.alias,
                                    'dia': str(dia), 'error': e.mensaje})
            pagos = pagos_por_token_dia[clave]
            if pagos is None:
                continue
            res = mp.conciliar_cierre_mp(cfg, dia, pagos=pagos)
            # El cierre sabe de qué caja es cada pago sin registro (por el punto de
            # venta que informa MP): se guarda para decir de qué TIENDA es cada fila
            # de abajo, en vez de «la primera caja de la cuenta».
            for item in res.get('sin_registro') or []:
                if item.get('atribuible'):
                    atribucion.setdefault(str(item.get('payment_id') or ''), set()).add(cfg.id)
            arqueo = (ArqueoCaja.objects.filter(sucursal_id=cfg.sucursal_id, fecha_arqueo=dia)
                      .order_by('-id').values('total_mercadopago_pos_teorico', 'cierre_mp_fisico')
                      .first()) or {}
            cajas.append({
                'dia': str(dia),
                'sucursal': cfg.sucursal.alias,
                'caja': cfg.nombre,
                'ok': res.get('ok', False),
                'error': res.get('error', ''),
                'sistema': res.get('sistema_total', 0),
                'mp': res.get('mp_total', 0),
                'diferencia': res.get('diferencia', 0),
                'cuadra': res.get('cuadra', False),
                'sin_registro': len(res.get('sin_registro') or []),
                'sin_confirmar': len(res.get('sin_confirmar') or []),
                'arqueo_teorico_mp': int(arqueo.get('total_mercadopago_pos_teorico') or 0),
                'arqueo_cierre_fisico': int(arqueo.get('cierre_mp_fisico') or 0),
            })

    # Pagos de MP (todas las cuentas, sin duplicar) vs registros locales.
    todos = {}
    cfg_de_pago = {}
    token_de_pago = {}
    for (tok, _dia), pagos in pagos_por_token_dia.items():
        for p in pagos or []:
            todos[str(p.get('id'))] = p
            cfg_de_pago.setdefault(str(p.get('id')), config_de_token.get(tok))
            token_de_pago.setdefault(str(p.get('id')), tok)
    ids_mp_vistos = set(todos)
    refs = {str(p.get('external_reference') or '') for p in todos.values()} - {''}
    refs_locales = set(TransaccionMercadoPago.objects.filter(external_reference__in=refs)
                       .values_list('external_reference', flat=True))
    ids_locales = set(filter(None, TransaccionMercadoPago.objects.filter(
        Q(payment_id_mp__in=ids_mp_vistos) | Q(payment_id__in=ids_mp_vistos))
        .values_list('payment_id_mp', flat=True)))
    vouchers_manuales = {}
    for pago in (TicketDetallePago.objects
                 .filter(metodo_pago__in=METODOS_MP,
                         creado_en__date__gte=d - timedelta(days=1),
                         creado_en__date__lte=h + timedelta(days=1))
                 .select_related('ticket', 'ticket__sucursal')):
        voucher = _solo_digitos(pago.voucher)
        if voucher:
            vouchers_manuales[voucher] = pago
    # Cuentas con UNA sola caja habilitada (todas las tiendas, no solo la filtrada):
    # ahí un pago sin datos de caja es de esa caja sin duda.
    cajas_por_token = {}
    for c in MercadoPagoConfig.objects.select_related('sucursal', 'cuenta').filter(habilitado=True):
        try:
            cajas_por_token.setdefault(mp._token(c), []).append(c)
        except mp.MercadoPagoError:
            continue
    unica_de_cuenta = {tok: cs[0] for tok, cs in cajas_por_token.items() if len(cs) == 1}
    sin_atribuir_ocultos = 0
    # N° de «MP manual» que YA son un pago aprobado de MP por el mismo monto (solo
    # falta registrarlos al aplicar liberaciones): esas ventas tienen su cobro y no
    # se le sugieren a otro pago.
    vouchers_calzados = sorted(
        v for v, pago in vouchers_manuales.items()
        if v in todos and str(todos[v].get('status') or '') == 'approved'
        and _monto(todos[v].get('transaction_amount')) == int(pago.monto or 0))
    for pid, p in todos.items():
        if str(p.get('status') or '') != 'approved':
            continue
        ref = str(p.get('external_reference') or '')
        if ref in refs_locales or pid in ids_locales or pid in vouchers_manuales:
            continue
        cfg_pago = cfg_de_pago.get(pid)
        sucursal_ref = mp._sucursal_de_referencia(ref) if ref else None
        # Tienda del pago: la caja a la que lo atribuyó el cierre; si no, la de la
        # referencia propia; si la cuenta tiene una sola caja, esa. Si no, queda
        # «sin atribuir» (nunca la primera caja de la cuenta a ciegas).
        cajas_pago = [cfg_por_id[i] for i in (atribucion.get(pid) or ()) if i in cfg_por_id]
        cfg_atr = None
        if cajas_pago and len({c.sucursal_id for c in cajas_pago}) == 1:
            # Una o más cajas de la MISMA tienda: la tienda es segura (se usa la principal).
            cfg_atr = sorted(cajas_pago, key=lambda c: (not c.es_principal, c.id))[0]
        if cfg_atr is None and sucursal_ref is not None:
            cfg_atr = next((c for c in configs if c.sucursal_id == sucursal_ref), None)
        if cfg_atr is None:
            cfg_atr = unica_de_cuenta.get(token_de_pago.get(pid))
        if sucursal_id and (cfg_atr is None or cfg_atr.sucursal_id != int(sucursal_id)):
            if cfg_atr is None:
                sin_atribuir_ocultos += 1
            continue   # con una tienda elegida: solo los pagos que son de ella
        instante = _instante(p.get('date_created'))
        sin_registro.append({
            'tipo': 'SIN_REGISTRO',
            'payment_id': pid,
            # Hora de Chile (MP la entrega con su propio huso)
            'fecha': (timezone.localtime(instante).strftime('%Y-%m-%d %H:%M') if instante
                      else str(p.get('date_created') or '')[:16].replace('T', ' ')),
            'instante': instante.isoformat() if instante else '',
            'monto': _monto(p.get('transaction_amount')),
            'medio': mp.etiqueta_medio_mp(p.get('payment_type_id')),
            'payment_type': str(p.get('payment_type_id') or ''),
            'ultimos_4': str((p.get('card') or {}).get('last_four_digits') or '')[:4],
            'external_reference': ref,
            'descripcion': str(p.get('description') or '')[:60],
            'sucursal_ref': sucursal_ref,
            'atribuida': cfg_atr is not None,
            # Caja por la que se lee el pago en la API (cualquiera de la cuenta sirve).
            'config_id': (cfg_atr or cfg_pago).id if (cfg_atr or cfg_pago) else None,
            'caja': f'{cfg_atr.sucursal.alias} · {cfg_atr.nombre}' if cfg_atr else 'Sin atribuir',
            'sucursal_id': cfg_atr.sucursal_id if cfg_atr else None,
            'sucursal': cfg_atr.sucursal.alias if cfg_atr else '',
        })

    manuales_sin_pago = []
    for pago in (TicketDetallePago.objects
                 .filter(metodo_pago__in=METODOS_MP, origen_pago='MANUAL', ticket__estado='PAGADO',
                         transacciones_mercadopago__isnull=True,
                         creado_en__date__gte=d, creado_en__date__lte=h)
                 .select_related('ticket', 'ticket__sucursal')):
        if sucursal_id and pago.ticket.sucursal_id != int(sucursal_id):
            continue
        voucher = _solo_digitos(pago.voucher)
        pago_mp = todos.get(voucher) if voucher else None
        estado_mp = str((pago_mp or {}).get('status') or '')
        monto_mp = _monto((pago_mp or {}).get('transaction_amount'))
        if pago_mp is not None and estado_mp == 'approved' and monto_mp == int(pago.monto or 0):
            continue
        if not voucher:
            motivo = 'sin N° de operación'
        elif pago_mp is not None and estado_mp != 'approved':
            motivo = f'el N° está «{estado_mp}» en Mercado Pago'
        elif pago_mp is not None:
            motivo = f'el N° existe en Mercado Pago pero por ${monto_mp:,}'.replace(',', '.')
        elif len(voucher) < 9:
            motivo = f'N° de {len(voucher)} dígitos: no parece de Mercado Pago (tienen 12)'
        else:
            motivo = 'N° no encontrado en Mercado Pago'
        manuales_sin_pago.append({
            'tipo': 'MANUAL_SIN_PAGO',
            'pago_id': pago.id,
            'ticket_id': pago.ticket_id,
            'sucursal_id': pago.ticket.sucursal_id,
            'metodo_pago': pago.metodo_pago,
            'ticket': pago.ticket.correlativo,
            'sucursal': pago.ticket.sucursal.alias,
            'fecha': timezone.localtime(pago.creado_en).strftime('%Y-%m-%d %H:%M'),
            'monto': pago.monto,
            'voucher': pago.voucher or '',
            'motivo': motivo,
            'sugerencia_n': None,
        })

    # N° mal digitado: un «MP manual» no reconocido y un pago sin registro del
    # mismo monto, día y tienda cuyo N° difiere en UN dígito. Se sugiere solo si
    # el par es único en las dos direcciones (los N° de una cuenta son casi
    # correlativos y los precios se repiten).
    parejas = {}
    for m in manuales_sin_pago:
        cands = [s for s in sin_registro
                 if s['monto'] == int(m['monto'] or 0) and s['fecha'][:10] == m['fecha'][:10]
                 and s['sucursal_id'] in (None, m['sucursal_id'])
                 and n_parecido(m['voucher'], s['payment_id'])]
        if len(cands) == 1:
            parejas.setdefault(cands[0]['payment_id'], []).append((m, cands[0]))
    for lista in parejas.values():
        if len(lista) != 1:
            continue
        m, s = lista[0]
        explicacion = n_parecido(m['voucher'], s['payment_id'])
        m['sugerencia_n'] = {'payment_id': s['payment_id'], 'config_id': s['config_id'],
                             'explicacion': explicacion, 'fecha': s['fecha'],
                             'medio': s['medio'], 'ultimos_4': s['ultimos_4']}
        s['n_mal_digitado'] = {'pago_id': m['pago_id'], 'ticket': m['ticket'],
                               'voucher': m['voucher'], 'explicacion': explicacion}

    return {
        'desde': str(d), 'hasta': str(h),
        'cajas': cajas,
        'sin_registro': sorted(sin_registro, key=lambda x: x['fecha']),
        'manuales_sin_pago': manuales_sin_pago,
        'errores': errores,
        'pagos_mp_consultados': len(todos),
        # con una tienda elegida: pagos de la cuenta que MP no permite atribuir a una tienda
        'sin_atribuir_ocultos': sin_atribuir_ocultos,
        'vouchers_calzados': vouchers_calzados,
    }
