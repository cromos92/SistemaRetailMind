"""
Informe final de una toma de inventario, con el formato del informe antiguo que
usa gerencia («2026 2479 Inventario General» de NICK1 / Matta 2479):

1. «Por marca»: INVENTARIO ANTIGUO (sistema) | DIF | INVENTARIO NUEVO (pistola),
   con Stock, P COSTO y P VENTA por marca y Total general.
2. «Diferencias»: una fila por SKU con diferencia, con las columnas del informe
   antiguo: id, sku, art, marca, costo, costo2, stk, pistola, mov, final, pvp,
   ttcosto1, ttpvp1, ttcosto2, ttpvp2.

Definiciones (las mismas en pantalla, Excel y scripts):
- stk      stock del sistema a la fecha de corte.
- mov      movimientos entre el corte y el momento del conteo (0 si se contó con
           la tienda cerrada). La base comparada es stk + mov.
- pistola  conteo físico. Una línea SIN CONTAR se informa con pistola 0: en un
           inventario completo, lo que no apareció en la pistola es faltante.
- final    ajuste = pistola − (stk + mov). Positivo sobrante, negativo faltante.
- costo2   costo + sobreprecio del producto. Es el «P COSTO» del informe antiguo:
           se verificó contra el informe de NICK1 (venta/costo ADIDAS 1,72; con
           costo solo daría 1,92). ttcosto1/ttpvp1 = base × costo2/pvp (sistema),
           ttcosto2/ttpvp2 = pistola × costo2/pvp (físico).
- El ajuste que entra al kardex se valoriza al costo FIFO de la toma
  (costo_unitario_sistema); va aparte en el Resumen para cuadrar con el módulo.

Las líneas excluidas del análisis (SKUs operativos: VISA, bolsas, envíos…, o
excluidas a mano) no suman en ningún total: se listan en su propia hoja.

El cálculo es puro (iterables de dict) para que lo usen la vista y los scripts
que simulan una toma antes de cargarla.
"""
import re
import unicodedata
from collections import OrderedDict

from django.utils import timezone

SIN_MARCA = 'SIN MARCA'

# Textos que la toma deja en la observación de la línea al resolver lo que no
# apareció en la pistola (views_gestion_inventarios._resolver_no_contados). El
# informe los lee para separar lo mantenido / lo faltante de lo realmente contado.
OBSERVACION_MANTENIDO = 'No se pistoleó: se mantiene el stock del sistema'
OBSERVACIONES_MANTENIDO = (OBSERVACION_MANTENIDO, 'No contado: se asume igual al sistema')
OBSERVACION_NO_APARECIO = 'No apareció en el conteo: faltante'

# Columnas de la hoja «Diferencias» (formato del informe antiguo + talla/estado)
COLUMNAS_DIFERENCIAS = [
    ('id', 'id', 'Id del producto/talla en el sistema'),
    ('sku', 'sku', 'Código SKU'),
    ('art', 'art', 'Artículo'),
    ('talla', 'talla', 'Talla'),
    ('marca', 'marca', 'Marca'),
    ('costo', 'costo', 'Costo del producto'),
    ('costo2', 'costo2', 'Costo + sobreprecio (P COSTO del informe antiguo)'),
    ('stk', 'stk', 'Stock del sistema a la fecha de corte'),
    ('pistola', 'pistola', 'Conteo físico (pistola); vacío = no apareció'),
    ('mov', 'mov', 'Movimientos entre el corte y el conteo'),
    ('final', 'final', 'Ajuste = pistola − (stk + mov)'),
    ('pvp', 'pvp', 'Precio de venta'),
    ('ttcosto1', 'ttcosto1', 'Sistema a costo2: (stk + mov) × costo2'),
    ('ttpvp1', 'ttpvp1', 'Sistema a venta: (stk + mov) × pvp'),
    ('ttcosto2', 'ttcosto2', 'Físico a costo2: pistola × costo2'),
    ('ttpvp2', 'ttpvp2', 'Físico a venta: pistola × pvp'),
    ('estado', 'estado', 'Contado / No apareció en la pistola'),
    ('alerta', 'alerta', 'Qué revisar antes de aprobar (doble lectura, sin stock en sistema, negativo)'),
]


def alerta_diferencia(base, fisico, final, no_aparecio=False):
    """
    Pista de qué revisar en una diferencia. La «doble lectura» es el caso real de
    PAO4 07-10-2026: la curva completa CHALADA 23-MITSU-2 (filas 1255-1260 de la
    pistola) vino exactamente al doble del sistema. Un código pistoleado de más
    varias veces (1 en sistema, 8 en pistola) cae en «lectura repetida».
    """
    if not final:
        return ''
    if no_aparecio:
        return 'No apareció: buscar antes de darlo por perdido' if base >= 2 else 'No apareció en la pistola'
    if base > 0 and fisico == 2 * base:
        return 'Posible doble lectura (pistola = 2 × sistema): recontar'
    if base > 0 and fisico > 2 * base:
        return f'Posible lectura repetida (pistola {fisico} vs sistema {base}): recontar'
    if base == 0 and final > 0:
        return 'Sin stock en sistema: ¿llegó sin ingreso o traspaso?'
    if base < 0:
        return 'Stock negativo en sistema: el conteo lo corrige'
    return ''


def _num(valor):
    try:
        return float(valor or 0)
    except (TypeError, ValueError):
        return 0.0


def calcular_fila(fila):
    """
    Completa una fila con los derivados. Claves de entrada: id, sku, art, talla,
    marca, costo, sobreprecio, pvp, costo_fifo, stk, mov, pistola (None = sin
    contar), excluida, motivo_exclusion. Devuelve la misma fila.
    """
    stk = int(fila.get('stk') or 0)
    mov = int(fila.get('mov') or 0)
    pistola = fila.get('pistola')
    fisico = int(pistola or 0)
    base = stk + mov
    costo = _num(fila.get('costo'))
    costo2 = costo + _num(fila.get('sobreprecio'))
    pvp = _num(fila.get('pvp'))
    final = fisico - base
    fila.update({
        'marca': (fila.get('marca') or '').strip() or SIN_MARCA,
        'stk': stk, 'mov': mov, 'base': base, 'fisico': fisico,
        'costo': costo, 'costo2': costo2, 'pvp': pvp,
        'final': final,
        'ttcosto1': base * costo2, 'ttpvp1': base * pvp,
        'ttcosto2': fisico * costo2, 'ttpvp2': fisico * pvp,
        'ajuste_fifo': final * _num(fila.get('costo_fifo') if fila.get('costo_fifo') is not None else costo),
        'estado': (
            'Excluido' if fila.get('excluida')
            else 'No apareció' if pistola is None or fila.get('no_aparecio')
            else 'Mantenido' if fila.get('mantenido')
            else 'Contado'
        ),
    })
    fila['alerta'] = '' if fila.get('excluida') else alerta_diferencia(
        base, fisico, final, no_aparecio=fila['estado'] == 'No apareció'
    )
    return fila


def _acumulador_marca(marca):
    return {
        'marca': marca, 'skus': 0, 'skus_con_diferencia': 0, 'sin_contar': 0,
        'ant_stock': 0, 'ant_costo': 0.0, 'ant_venta': 0.0,
        'nue_stock': 0, 'nue_costo': 0.0, 'nue_venta': 0.0,
    }


def analizar(filas):
    """
    Una pasada sobre las filas (sirve con un iterator de 300k líneas).
    Devuelve dict con: marcas (lista ordenada como el informe antiguo, sin
    distinguir mayúsculas), total, resumen, diferencias (filas con final != 0)
    y excluidas.
    """
    marcas = {}
    total = _acumulador_marca('Total general')
    diferencias, excluidas, mantenidas = [], [], []
    r = {
        'skus': 0, 'contados': 0, 'sin_contar': 0, 'excluidos': 0, 'coinciden': 0,
        'sobrantes_lineas': 0, 'sobrantes_unidades': 0, 'sobrantes_costo': 0.0, 'sobrantes_venta': 0.0,
        'faltantes_lineas': 0, 'faltantes_unidades': 0, 'faltantes_costo': 0.0, 'faltantes_venta': 0.0,
        'sin_contar_unidades': 0, 'excluidos_unidades': 0,
        # no aparecieron en la pistola: sin resolver + resueltas como «Faltante»
        'no_aparecieron': 0, 'no_aparecieron_unidades': 0,
        'posibles_dobles': 0, 'posibles_dobles_unidades': 0,
        'sin_stock_sistema': 0, 'sin_stock_sistema_unidades': 0,
        # no se pistolearon y conservan el stock del sistema (toman el antiguo)
        'mantenidos': 0, 'mantenidos_unidades': 0,
        'ajuste_fifo_neto': 0.0, 'reconteos': 0,
    }

    for fila in filas:
        calcular_fila(fila)
        if fila.get('excluida'):
            r['excluidos'] += 1
            r['excluidos_unidades'] += fila['base']
            excluidas.append(fila)
            continue

        r['skus'] += 1
        acc = marcas.get(fila['marca'])
        if acc is None:
            acc = marcas[fila['marca']] = _acumulador_marca(fila['marca'])
        for a in (acc, total):
            a['skus'] += 1
            a['ant_stock'] += fila['base']
            a['ant_costo'] += fila['ttcosto1']
            a['ant_venta'] += fila['ttpvp1']
            a['nue_stock'] += fila['fisico']
            a['nue_costo'] += fila['ttcosto2']
            a['nue_venta'] += fila['ttpvp2']
            if fila['final']:
                a['skus_con_diferencia'] += 1
            if fila.get('pistola') is None:
                a['sin_contar'] += 1

        if fila.get('pistola') is None:
            r['sin_contar'] += 1
            r['sin_contar_unidades'] += max(fila['base'], 0)
        else:
            r['contados'] += 1
        if fila['estado'] == 'No apareció':
            r['no_aparecieron'] += 1
            r['no_aparecieron_unidades'] += max(fila['base'], 0)
        elif fila['estado'] == 'Mantenido':
            r['mantenidos'] += 1
            r['mantenidos_unidades'] += max(fila['fisico'], 0)
            mantenidas.append(fila)
        if fila['alerta'].startswith(('Posible doble', 'Posible lectura repetida')):
            r['posibles_dobles'] += 1
            r['posibles_dobles_unidades'] += fila['final']
        elif fila['alerta'].startswith('Sin stock en sistema'):
            r['sin_stock_sistema'] += 1
            r['sin_stock_sistema_unidades'] += fila['final']
        if fila.get('reconteo'):
            r['reconteos'] += 1

        if fila['final'] > 0:
            r['sobrantes_lineas'] += 1
            r['sobrantes_unidades'] += fila['final']
            r['sobrantes_costo'] += fila['final'] * fila['costo2']
            r['sobrantes_venta'] += fila['final'] * fila['pvp']
        elif fila['final'] < 0:
            r['faltantes_lineas'] += 1
            r['faltantes_unidades'] += -fila['final']
            r['faltantes_costo'] += -fila['final'] * fila['costo2']
            r['faltantes_venta'] += -fila['final'] * fila['pvp']
        else:
            r['coinciden'] += 1
        r['ajuste_fifo_neto'] += fila['ajuste_fifo']
        if fila['final']:
            diferencias.append(fila)

    for a in list(marcas.values()) + [total]:
        a['dif'] = a['nue_stock'] - a['ant_stock']
        a['dif_costo'] = a['nue_costo'] - a['ant_costo']
        a['dif_venta'] = a['nue_venta'] - a['ant_venta']

    r.update({
        'unidades_sistema': total['ant_stock'], 'unidades_fisicas': total['nue_stock'],
        'dif_unidades': total['dif'],
        'costo_sistema': total['ant_costo'], 'costo_fisico': total['nue_costo'],
        'venta_sistema': total['ant_venta'], 'venta_fisico': total['nue_venta'],
        'precision': round(r['coinciden'] * 100 / r['skus'], 2) if r['skus'] else 0,
    })
    orden = lambda f: (f['marca'].upper(), str(f.get('art') or ''), str(f.get('talla') or ''), str(f['sku']))  # noqa: E731
    diferencias.sort(key=orden)
    mantenidas.sort(key=orden)
    return {
        'marcas': sorted(marcas.values(), key=lambda a: a['marca'].upper()),
        'total': total,
        'resumen': r,
        'diferencias': diferencias,
        'excluidas': excluidas,
        'mantenidas': mantenidas,
    }


# ---------------------------------------------------------------------------
# Lectura desde la BD
# ---------------------------------------------------------------------------

def filas_desde_toma(inventario):
    """Genera las filas del informe desde los detalles de una toma (streaming)."""
    campos = (
        'producto_talla_id', 'sku', 'producto_nombre', 'talla_nombre', 'marca_nombre',
        'stock_sistema', 'stock_movimientos_post_corte', 'stock_fisico', 'contado',
        'excluir_de_analisis', 'reconteo_requerido', 'costo_unitario_sistema',
        'precio_venta_sistema', 'producto_talla__producto__costo',
        'producto_talla__producto__sobreprecio', 'observaciones',
    )
    qs = inventario.detalles.order_by('id').values(*campos)
    for d in qs.iterator(chunk_size=2000):
        yield {
            'id': d['producto_talla_id'],
            'sku': d['sku'],
            'art': d['producto_nombre'],
            'talla': d['talla_nombre'] or '',
            'marca': d['marca_nombre'],
            'costo': d['producto_talla__producto__costo'] or 0,
            'sobreprecio': d['producto_talla__producto__sobreprecio'] or 0,
            # El precio de la toma es el snapshot al corte (no el de hoy)
            'pvp': d['precio_venta_sistema'] or 0,
            'costo_fifo': d['costo_unitario_sistema'],
            'stk': d['stock_sistema'],
            'mov': d['stock_movimientos_post_corte'] if d['contado'] else 0,
            'pistola': d['stock_fisico'] if d['contado'] else None,
            # resuelta con «Faltante» (no contados en 0): se informa como no apareció
            'no_aparecio': bool(d['contado'] and OBSERVACION_NO_APARECIO in (d['observaciones'] or '')),
            # resuelta con «Mantener stock del sistema»: no se pistoleó, toma el antiguo
            'mantenido': bool(d['contado'] and any(t in (d['observaciones'] or '') for t in OBSERVACIONES_MANTENIDO)),
            'reconteo': d['reconteo_requerido'] and not d['excluir_de_analisis'],
            'excluida': d['excluir_de_analisis'],
            'motivo_exclusion': _motivo_exclusion(d['observaciones']) if d['excluir_de_analisis'] else '',
        }


def _motivo_exclusion(observaciones):
    texto = observaciones or ''
    if 'SKU operativo' in texto:
        return 'SKU operativo (no es mercadería)'
    return 'Excluido manualmente'


def no_cargados_desde_logs(inventario):
    """
    Códigos de la pistola que no existen en la sucursal (no se pudieron cargar),
    con su cantidad si la importación la guardó, y en qué otra sucursal existe el
    SKU (típicamente mercadería que llegó sin traspaso o un traspaso sin recibir).
    Si el mismo código vino en varias importaciones, manda la última (re-importar
    reemplaza el conteo, igual que para los SKUs que sí existen).
    """
    from app.models import Producto_Talla

    cantidades = OrderedDict()
    for log in inventario.logs.filter(tipo_accion='REGISTRO_CONTEO').order_by('created_at'):
        datos = log.datos_adicionales or {}
        detalle = datos.get('no_encontrados_detalle')
        if detalle:
            for item in detalle:
                cantidades[str(item.get('sku'))] = item.get('cantidad')
        else:
            for sku in datos.get('no_encontrados') or []:
                cantidades.setdefault(str(sku), None)
    if not cantidades:
        return []

    numericos = [int(s) for s in cantidades if s.isdigit() and len(s) <= 18]
    donde = {}
    if numericos:
        for pt in (Producto_Talla.objects.filter(sku__in=numericos)
                   .exclude(producto__sucursal_id=inventario.sucursal_id)
                   .values('sku', 'stock', 'producto__sucursal__alias', 'producto__articulo',
                           'producto__atributo1__valor')):
            donde.setdefault(str(pt['sku']), []).append(pt)

    resultado = []
    for sku, cantidad in cantidades.items():
        otras = donde.get(sku, [])
        resultado.append({
            'sku': sku,
            'cantidad': cantidad,
            'articulo': otras[0]['producto__articulo'] if otras else '',
            'marca': (otras[0]['producto__atributo1__valor'] or '') if otras else '',
            'existe_en': ', '.join(f"{o['producto__sucursal__alias']} ({o['stock']})" for o in otras),
            'sugerencia': _sugerencia_no_cargado(sku, otras),
        })
    return resultado


def _sugerencia_no_cargado(sku, otras):
    if otras:
        return 'Existe en otra sucursal: traspasarlo a esta tienda o revisar traspaso sin recibir'
    if sku.isdigit() and len(sku) >= 12:
        return 'Parece código de barras del proveedor (EAN), no el SKU: re-escanear la etiqueta'
    return 'Código desconocido: revisar la etiqueta o crear el producto'


def cabecera_desde_toma(inventario):
    corte = timezone.localtime(inventario.fecha_corte)
    sucursal = inventario.sucursal
    alcance, alcance_detalle = alcance_de_toma(inventario)
    return {
        'anio': corte.year,
        'direccion': (sucursal.direccion or '').strip(),
        'alias': sucursal.alias,
        'numero': inventario.numero_inventario,
        'nombre': inventario.nombre,
        'corte': corte.strftime('%d/%m/%Y %H:%M'),
        'tienda_cerrada': inventario.conteo_tienda_cerrada,
        'estado': inventario.get_estado_display(),
        'tipo': inventario.get_tipo_inventario_display(),
        'tipo_codigo': inventario.tipo_inventario,
        'alcance': alcance,
        'alcance_detalle': alcance_detalle,
    }


# Cómo se nombra cada tipo de toma en el título («2026 2438  Inventario por Categoría · Calzado»). El
# completo conserva el título del informe antiguo de gerencia: «Inventario General».
TIPO_EN_TITULO = {
    'COMPLETO': 'General',
    'POR_MARCA': 'por Marca',
    'POR_CATEGORIA': 'por Categoría',
    'POR_ATRIBUTO': 'por Atributo',
    'SELECTIVO': 'Selectivo',
    'CICLICO': 'Cíclico',
    'ALEATORIO': 'Aleatorio',
}
# En el título van hasta estos nombres; con más se resume («Calzado (10 de 13 categorías)»).
MAX_NOMBRES_TITULO = 3


def _nombres_cortos(nombres):
    if len(nombres) <= MAX_NOMBRES_TITULO:
        return ', '.join(nombres)
    return f"{', '.join(nombres[:MAX_NOMBRES_TITULO])} y {len(nombres) - MAX_NOMBRES_TITULO} más"


def alcance_de_toma(inventario):
    """
    Qué cubre la toma, según sus filtros: (corto para el título, detalle para la hoja Resumen).

    Categorías agrupadas por su padre: todas las hijas de Calzado elegidas → «Calzado»; unas
    pocas → «Calzado: Zapatillas, Botines»; muchas → «Calzado (10 de 13 categorías)». Además
    marcas, color/género y productos sueltos. Un inventario completo sin filtros → ('', '').
    """
    from app.models import AtributoOpcion, Categoria

    filtros = inventario.filtros_aplicados or {}
    cortos, detalles = [], []

    ids_categorias = [int(i) for i in (filtros.get('categorias') or []) if str(i).isdigit()]
    if ids_categorias:
        por_padre = OrderedDict()
        for c in Categoria.objects.filter(id__in=ids_categorias).select_related('padre').order_by('padre__nombre', 'nombre'):
            por_padre.setdefault(c.padre, []).append(c.nombre)
        for padre, nombres in por_padre.items():
            if padre is None:  # categorías viejas, sin padre
                cortos.append(_nombres_cortos(nombres))
                detalles.append(', '.join(nombres))
                continue
            total = Categoria.objects.filter(padre=padre).count()
            if len(nombres) >= total:
                cortos.append(padre.nombre)
            elif len(nombres) <= MAX_NOMBRES_TITULO:
                cortos.append(f"{padre.nombre}: {', '.join(nombres)}")
            else:
                cortos.append(f'{padre.nombre} ({len(nombres)} de {total} categorías)')
            detalles.append(f"{padre.nombre}: {', '.join(nombres)}")

    ids_marcas = [int(i) for i in (filtros.get('marcas') or []) if str(i).isdigit()]
    if ids_marcas:
        nombres = list(AtributoOpcion.objects.filter(id__in=ids_marcas).order_by('valor').values_list('valor', flat=True))
        cortos.append(_nombres_cortos(nombres))
        detalles.append(f"Marcas: {', '.join(nombres)}")

    for clave, etiqueta in (('color', 'Color'), ('genero', 'Género')):
        ids = [int(i) for i in ((filtros.get('atributos') or {}).get(clave) or []) if str(i).isdigit()]
        if ids:
            nombres = list(AtributoOpcion.objects.filter(id__in=ids).order_by('valor').values_list('valor', flat=True))
            cortos.append(f'{etiqueta}: {_nombres_cortos(nombres)}')
            detalles.append(f"{etiqueta}: {', '.join(nombres)}")

    if filtros.get('productos'):
        n = len(filtros['productos'])
        cortos.append(f'{n} producto{"s" if n != 1 else ""}')
        detalles.append(f'{n} producto{"s" if n != 1 else ""} elegidos')

    return ' + '.join(cortos), ' · '.join(detalles)


def nombre_archivo_informe(cabecera):
    """informe_<número>_<tienda>[_<alcance>].xlsx, en ASCII (va en Content-Disposition)."""
    alcance = unicodedata.normalize('NFKD', cabecera.get('alcance') or '').encode('ascii', 'ignore').decode()
    alcance = re.sub(r'[^A-Za-z0-9]+', '_', alcance).strip('_')[:40]
    return f"informe_{cabecera.get('numero', '')}_{cabecera.get('alias', '')}{'_' + alcance if alcance else ''}.xlsx"


# ---------------------------------------------------------------------------
# Excel
# ---------------------------------------------------------------------------

def titulo_informe(cabecera):
    """'2026 2458  Inventario General' como el informe antiguo (número de la dirección). Una toma
    parcial dice su tipo y qué cubre: '2026 2438  Inventario por Categoría · Calzado'."""
    numero = re.search(r'\d+', cabecera.get('direccion') or '')
    tipo = TIPO_EN_TITULO.get(cabecera.get('tipo_codigo') or 'COMPLETO', 'General')
    alcance = cabecera.get('alcance') or ''
    return (f"{cabecera.get('anio', '')} {numero.group(0) if numero else cabecera.get('alias', '')}  "
            f"Inventario {tipo}{f' · {alcance}' if alcance else ''}")


def construir_workbook(cabecera, analisis, no_cargados=()):
    """Workbook con hojas Por marca, Diferencias, Mantenidos, No cargados, Excluidos y Resumen."""
    import openpyxl
    from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
    from openpyxl.utils import get_column_letter

    PESOS = '"$" #,##0;[Red]-"$" #,##0;"$" -'
    UNIDADES = '#,##0;[Red]-#,##0;0'
    fino = Side(style='thin', color='808080')
    borde = Border(left=fino, right=fino, top=fino, bottom=fino)
    negrita = Font(bold=True)
    rojo = Font(bold=True, color='C00000')
    centro = Alignment(horizontal='center', vertical='center')
    cabecera_fill = PatternFill(start_color='DDE3EF', end_color='DDE3EF', fill_type='solid')

    wb = openpyxl.Workbook()

    # ===================== POR MARCA (foto «Inventario General») =====================
    ws = wb.active
    ws.title = 'Por marca'
    subtitulo = f"{cabecera.get('direccion') or ''} · {cabecera.get('alias') or ''}".strip(' ·')
    ws.merge_cells('A1:K1')
    ws['A1'] = titulo_informe(cabecera)
    ws['A1'].font = Font(bold=True, size=16)
    ws['A1'].alignment = centro
    ws.merge_cells('A2:D2'); ws['A2'] = 'INVENTARIO ANTIGUO'
    ws.merge_cells('H2:K2'); ws['H2'] = 'INVENTARIO NUEVO'
    ws.merge_cells('A3:D3'); ws['A3'] = f'{subtitulo} · sistema al corte {cabecera.get("corte", "")}'
    ws.merge_cells('H3:K3'); ws['H3'] = f'{subtitulo} · conteo físico (pistola)'
    for celda in ('A2', 'H2', 'A3', 'H3'):
        ws[celda].font = rojo
        ws[celda].alignment = centro
    encabezado = ['Marca', 'Stock', 'P COSTO', 'P VENTA', None, 'DIF', None, 'Marca', 'Stock', 'P COSTO', 'P VENTA',
                  None, 'DIF P COSTO', 'DIF P VENTA']
    for col, texto in enumerate(encabezado, 1):
        if texto is None:
            continue
        c = ws.cell(row=4, column=col, value=texto)
        c.font = negrita
        c.alignment = centro
        c.border = borde
        c.fill = cabecera_fill

    def _fila_marca(fila_n, a, total=False):
        valores = {
            1: a['marca'], 2: a['ant_stock'], 3: a['ant_costo'], 4: a['ant_venta'],
            6: a['dif'],
            8: a['marca'], 9: a['nue_stock'], 10: a['nue_costo'], 11: a['nue_venta'],
            13: a['dif_costo'], 14: a['dif_venta'],
        }
        for col, valor in valores.items():
            c = ws.cell(row=fila_n, column=col, value=valor)
            c.border = borde
            if col in (2, 6, 9):
                c.number_format = UNIDADES
                c.alignment = centro
            elif col in (3, 4, 10, 11, 13, 14):
                c.number_format = PESOS
            else:
                c.alignment = centro
            if total:
                c.font = negrita

    fila_n = 5
    for a in analisis['marcas']:
        _fila_marca(fila_n, a)
        fila_n += 1
    _fila_marca(fila_n, analisis['total'], total=True)
    ws.cell(row=fila_n + 2, column=1, value=(
        'P COSTO = unidades × (costo + sobreprecio), igual que el informe antiguo. '
        'Inventario nuevo: lo faltante cuenta 0 y lo no pistoleado que se mantuvo toma la cantidad del antiguo. '
        'No incluye los SKUs excluidos (operativos: VISA, bolsas, envíos…), cuyo stock queda igual.'
    )).font = Font(italic=True, size=9, color='595959')
    notas = []  # (texto, color)
    if analisis['resumen']['no_aparecieron']:
        notas.append((
            f"{analisis['resumen']['no_aparecieron']} SKU(s) del sistema no aparecieron en la pistola "
            f"({analisis['resumen']['no_aparecieron_unidades']} u.) y se dieron como faltante "
            f"(hoja «Diferencias», estado «No apareció»).",
            'C00000',
        ))
    if analisis['resumen']['mantenidos']:
        notas.append((
            f"{analisis['resumen']['mantenidos']} SKU(s) no se pistolearon y conservan el stock del sistema "
            f"({analisis['resumen']['mantenidos_unidades']} u.): en el inventario nuevo se toma la cantidad del "
            f"antiguo (hoja «Mantenidos»).",
            '595959',
        ))
    for i, (texto, color) in enumerate(notas):
        ws.cell(row=fila_n + 3 + i, column=1, value=texto).font = Font(italic=True, size=9, color=color)
    anchos = {'A': 24, 'B': 9, 'C': 15, 'D': 15, 'E': 2, 'F': 7, 'G': 2, 'H': 24, 'I': 9, 'J': 15, 'K': 15,
              'L': 2, 'M': 15, 'N': 15}
    for letra, ancho in anchos.items():
        ws.column_dimensions[letra].width = ancho
    ws.freeze_panes = 'A5'

    # ===================== DIFERENCIAS (foto de análisis por SKU) =====================
    wd = wb.create_sheet('Diferencias')
    for col, (_, titulo, ayuda) in enumerate(COLUMNAS_DIFERENCIAS, 1):
        c = wd.cell(row=1, column=col, value=titulo)
        c.font = negrita
        c.alignment = centro
        c.border = borde
        c.fill = cabecera_fill
    columnas_pesos = {'costo', 'costo2', 'pvp', 'ttcosto1', 'ttpvp1', 'ttcosto2', 'ttpvp2'}
    columnas_unidades = {'stk', 'pistola', 'mov', 'final'}
    fila_n = 2
    for f in analisis['diferencias']:
        for col, (clave, _, _) in enumerate(COLUMNAS_DIFERENCIAS, 1):
            valor = f.get(clave)
            if clave == 'sku':
                valor = int(valor) if str(valor).isdigit() and len(str(valor)) < 16 else valor
            if clave == 'mov' and not valor:
                valor = None  # como el informe antiguo: vacío si no hubo movimientos
            if clave == 'pistola' and f.get('estado') == 'No apareció':
                valor = None  # como el informe antiguo: pistola vacía = no se leyó
            c = wd.cell(row=fila_n, column=col, value=valor)
            if clave in columnas_pesos:
                c.number_format = PESOS
            elif clave in columnas_unidades:
                c.number_format = UNIDADES
                c.alignment = centro
            if clave == 'final':
                c.font = Font(bold=True, color='008866' if (valor or 0) > 0 else 'C00000')
        fila_n += 1
    ultima = fila_n - 1
    if ultima >= 2:
        wd.cell(row=fila_n, column=1, value='TOTAL').font = negrita
        for col, (clave, _, _) in enumerate(COLUMNAS_DIFERENCIAS, 1):
            if clave in columnas_unidades or clave in ('ttcosto1', 'ttpvp1', 'ttcosto2', 'ttpvp2'):
                letra = get_column_letter(col)
                c = wd.cell(row=fila_n, column=col, value=f'=SUBTOTAL(9,{letra}2:{letra}{ultima})')
                c.font = negrita
                c.number_format = PESOS if clave.startswith('tt') else UNIDADES
        wd.auto_filter.ref = f'A1:{get_column_letter(len(COLUMNAS_DIFERENCIAS))}{ultima}'
    else:
        wd.cell(row=2, column=1, value='Sin diferencias: el conteo coincide con el sistema en todos los SKUs.')
    wd.freeze_panes = 'C2'
    for col, (clave, _, _) in enumerate(COLUMNAS_DIFERENCIAS, 1):
        wd.column_dimensions[get_column_letter(col)].width = {
            'art': 26, 'marca': 16, 'sku': 13, 'estado': 14, 'talla': 7, 'id': 9, 'alerta': 52,
        }.get(clave, 11 if clave in columnas_unidades else 13)

    # ===================== MANTENIDOS =====================
    wm = wb.create_sheet('Mantenidos')
    wm.append(['SKUs que no se pistolearon y CONSERVAN el stock del sistema (no se ajustan; '
               'en el inventario nuevo se toma la cantidad del antiguo)'])
    wm['A1'].font = negrita
    wm.append(['id', 'sku', 'art', 'talla', 'marca', 'stock que se mantiene', 'P COSTO', 'P VENTA'])
    for c in wm[2]:
        c.font = negrita
        c.fill = cabecera_fill
        c.border = borde
    for f in analisis.get('mantenidas', []):
        wm.append([f['id'], f['sku'], f.get('art'), f.get('talla'), f['marca'], f['fisico'],
                   f['ttcosto2'], f['ttpvp2']])
        wm.cell(row=wm.max_row, column=7).number_format = PESOS
        wm.cell(row=wm.max_row, column=8).number_format = PESOS
    if not analisis.get('mantenidas'):
        wm.append(['(ninguno)'])
    for letra, ancho in {'A': 9, 'B': 13, 'C': 26, 'D': 7, 'E': 16, 'F': 20, 'G': 14, 'H': 14}.items():
        wm.column_dimensions[letra].width = ancho

    # ===================== NO CARGADOS =====================
    wn = wb.create_sheet('No cargados')
    wn.append(['Códigos leídos por la pistola que NO existen en esta sucursal (no se pudieron cargar a la toma)'])
    wn['A1'].font = negrita
    wn.append(['Código', 'Cantidad', 'Existe en', 'Artículo', 'Marca', 'Qué hacer'])
    for c in wn[2]:
        c.font = negrita
        c.fill = cabecera_fill
        c.border = borde
    for n in no_cargados:
        wn.append([n['sku'], n['cantidad'], n['existe_en'], n['articulo'], n['marca'], n['sugerencia']])
    if not no_cargados:
        wn.append(['(ninguno)'])
    for letra, ancho in {'A': 16, 'B': 10, 'C': 26, 'D': 22, 'E': 16, 'F': 70}.items():
        wn.column_dimensions[letra].width = ancho

    # ===================== EXCLUIDOS =====================
    we = wb.create_sheet('Excluidos')
    we.append(['SKUs excluidos del análisis: NO se ajustan (su stock queda igual) y no suman en el informe'])
    we['A1'].font = negrita
    we.append(['id', 'sku', 'art', 'talla', 'marca', 'stk (sistema)', 'pistola', 'motivo'])
    for c in we[2]:
        c.font = negrita
        c.fill = cabecera_fill
        c.border = borde
    for f in analisis['excluidas']:
        we.append([f['id'], f['sku'], f.get('art'), f.get('talla'), f['marca'], f['base'], f.get('pistola'),
                   f.get('motivo_exclusion') or ''])
    if not analisis['excluidas']:
        we.append(['(ninguno)'])
    for letra, ancho in {'A': 9, 'B': 13, 'C': 26, 'D': 7, 'E': 16, 'F': 13, 'G': 9, 'H': 36}.items():
        we.column_dimensions[letra].width = ancho

    # ===================== RESUMEN =====================
    wr = wb.create_sheet('Resumen')
    r = analisis['resumen']
    filas_resumen = [
        ('TOMA DE INVENTARIO', None),
        ('Número', cabecera.get('numero')),
        ('Nombre', cabecera.get('nombre')),
        ('Sucursal', subtitulo),
        ('Tipo', cabecera.get('tipo')),
        ('Qué cubre', cabecera.get('alcance_detalle') or 'Toda la tienda'),
        ('Fecha de corte', cabecera.get('corte') + (' (contado con la tienda cerrada)' if cabecera.get('tienda_cerrada') else '')),
        ('Estado', cabecera.get('estado')),
        ('Generado', timezone.localtime().strftime('%d/%m/%Y %H:%M')),
        (None, None),
        ('UNIDADES', None),
        ('Sistema (inventario antiguo)', r['unidades_sistema']),
        ('Físico (inventario nuevo)', r['unidades_fisicas']),
        ('Diferencia neta', r['dif_unidades']),
        (None, None),
        ('SKUs', None),
        ('SKUs en el análisis', r['skus']),
        ('Coinciden exacto', r['coinciden']),
        ('Precisión (% SKUs sin diferencia)', f"{r['precision']}%"),
        ('Sobrantes (SKUs / unidades)', f"{r['sobrantes_lineas']} / {r['sobrantes_unidades']}"),
        ('Faltantes (SKUs / unidades)', f"{r['faltantes_lineas']} / {r['faltantes_unidades']}"),
        ('No aparecieron en la pistola (SKUs / unidades)', f"{r['no_aparecieron']} / {r['no_aparecieron_unidades']}"),
        ('…de ellos todavía sin resolver en la toma', r['sin_contar']),
        ('No pistoleados que conservan el stock del sistema (SKUs / unidades)',
         f"{r['mantenidos']} / {r['mantenidos_unidades']}"),
        ('Posible lectura doble o repetida: pistola ≥ 2 × sistema (SKUs / unidades de más)',
         f"{r['posibles_dobles']} / {r['posibles_dobles_unidades']}"),
        ('Sobrantes sin stock en sistema (SKUs / unidades)',
         f"{r['sin_stock_sistema']} / {r['sin_stock_sistema_unidades']}"),
        ('Esperan reconteo', r['reconteos']),
        ('Excluidos (SKUs / unidades)', f"{r['excluidos']} / {r['excluidos_unidades']}"),
        ('Códigos de la pistola no cargados', len(no_cargados)),
        (None, None),
        ('VALORIZACIÓN', None),
        ('Sistema a P COSTO', r['costo_sistema']),
        ('Físico a P COSTO', r['costo_fisico']),
        ('Diferencia a P COSTO', r['costo_fisico'] - r['costo_sistema']),
        ('Sistema a P VENTA', r['venta_sistema']),
        ('Físico a P VENTA', r['venta_fisico']),
        ('Diferencia a P VENTA', r['venta_fisico'] - r['venta_sistema']),
        ('Sobrantes a P COSTO', r['sobrantes_costo']),
        ('Faltantes a P COSTO', -r['faltantes_costo']),
        ('Ajuste neto que entra al kardex (costo FIFO)', r['ajuste_fifo_neto']),
        (None, None),
        ('CÓMO LEER LA HOJA «Diferencias»', None),
    ] + [(titulo, ayuda) for _, titulo, ayuda in COLUMNAS_DIFERENCIAS]
    for etiqueta, valor in filas_resumen:
        wr.append([etiqueta, valor])
        fila = wr.max_row
        if etiqueta and valor is None:
            wr.cell(row=fila, column=1).font = negrita
        # los montos se acumulan como float y las cantidades como int
        if isinstance(valor, float):
            wr.cell(row=fila, column=2).number_format = PESOS
        elif isinstance(valor, int):
            wr.cell(row=fila, column=2).number_format = UNIDADES
    wr.column_dimensions['A'].width = 46
    wr.column_dimensions['B'].width = 60

    return wb
